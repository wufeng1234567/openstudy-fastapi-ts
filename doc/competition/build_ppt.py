# -*- coding: utf-8 -*-
"""
生成《OpenStudy · 华为开发者大赛福建赛区》10 页参赛 PPT。
设计系统对齐 OpenStudy 前台的 DSH 同源令牌：近黑主色 + 单一蓝强调 + 发丝描边 + 大留白。
运行：D:\\Python3.11\\python.exe build_ppt.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(os.path.dirname(HERE), "screenshots")

# ---------------- 设计令牌 ----------------
INK = RGBColor(0x0F, 0x11, 0x15)
INK2 = RGBColor(0x35, 0x36, 0x38)
GRAY = RGBColor(0x61, 0x66, 0x6B)
GRAY2 = RGBColor(0x81, 0x85, 0x8C)
GRAY3 = RGBColor(0xAD, 0xB2, 0xB8)
LINE = RGBColor(0xE9, 0xEC, 0xF2)
LINE2 = RGBColor(0xF1, 0xF3, 0xF5)
BG = RGBColor(0xF9, 0xFA, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x41, 0x76, 0xE6)
BLUE_L = RGBColor(0xED, 0xF3, 0xFE)
BLUE_D = RGBColor(0x28, 0x31, 0x42)
AMBER = RGBColor(0xDD, 0x86, 0x29)
AMBER_L = RGBColor(0xFE, 0xF5, 0xE7)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
GREEN_L = RGBColor(0xE6, 0xFA, 0xED)

FONT = "微软雅黑"
FONT_L = "微软雅黑 Light"

W, H = 13.333, 7.5
M_L = 0.85
CW = W - 2 * M_L  # 11.633


def set_spacing(run, spc):
    """字符间距（磅的百分之一）"""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(spc)))


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, radius=None, shape_type=None):
    st = shape_type or (MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    no_shadow(shp)
    shp.text_frame.word_wrap = True
    return shp


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
         font=FONT, line_spacing=None, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0),
         spc=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        if spc:
            set_spacing(r, spc)
    return tb


def rich(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=None, margins=(0, 0, 0, 0)):
    """runs: list of dict(text,size,bold,color,font) — 同一段落内混排"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    for it in runs:
        r = p.add_run()
        r.text = it["text"]
        r.font.name = it.get("font", FONT)
        r.font.size = Pt(it.get("size", 14))
        r.font.bold = it.get("bold", False)
        r.font.color.rgb = it.get("color", INK)
    return tb


def line(slide, x1, y1, x2, y2, color=LINE, width=0.75):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)


def bg(slide, color=WHITE):
    b = rect(slide, 0, 0, W, H, fill=color)
    # 移到最底层
    sp = b._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def header(slide, kicker, title, lede=None):
    text(slide, M_L, 0.52, CW, 0.24, kicker, size=9.5, bold=True, color=BLUE, spc=120)
    text(slide, M_L, 0.80, CW, 0.62, title, size=27, bold=True, color=INK, line_spacing=1.08)
    if lede:
        text(slide, M_L, 1.44, CW, 0.34, lede, size=12, color=GRAY, line_spacing=1.2)


def footer(slide, idx, dark=False):
    c1 = RGBColor(0x43, 0x45, 0x4A) if dark else LINE
    c2 = RGBColor(0x81, 0x85, 0x8C) if not dark else RGBColor(0x61, 0x66, 0x6B)
    line(slide, M_L, 6.98, W - M_L, 6.98, color=c1, width=0.5)
    text(slide, M_L, 7.06, 6.0, 0.22, "OpenStudy · 华为开发者大赛福建赛区", size=8, color=c2)
    text(slide, W - M_L - 1.2, 7.06, 1.2, 0.22, f"{idx:02d} / 10", size=8, color=c2, align=PP_ALIGN.RIGHT)


def add_image(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


_CROP_CACHE = {}


def pic_card(slide, path, x, y, w, h, caption=None, pad=0.0):
    """带发丝描边的图片卡：先用 PIL 裁成目标比例落盘，再按精确尺寸插入（避免溢出）。"""
    from PIL import Image
    key = (path, round(w, 3), round(h, 3))
    if key not in _CROP_CACHE:
        im = Image.open(path).convert("RGB")
        iw, ih = im.size
        target = w / h
        if iw / ih > target:  # 原图更宽 → 裁左右
            nw = int(ih * target)
            left = (iw - nw) // 2
            box = (left, 0, left + nw, ih)
        else:  # 原图更高 → 裁上下
            nh = int(iw / target)
            top = (ih - nh) // 2
            box = (0, top, iw, top + nh)
        out = os.path.join(HERE, "_crop_cache")
        os.makedirs(out, exist_ok=True)
        dst = os.path.join(out, f"crop_{abs(hash(key)) % 10**8}.jpg")
        im.crop(box).save(dst, quality=92)
        _CROP_CACHE[key] = dst
    pic = slide.shapes.add_picture(_CROP_CACHE[key], Inches(x), Inches(y), Inches(w), Inches(h))
    if caption:
        text(slide, x, y + h + 0.08, w, 0.22, caption, size=8, color=GRAY2)
    return pic


def stat(slide, x, y, num, unit, label, num_size=22, num_color=INK, unit_color=GRAY2,
         label_color=GRAY2, w=2.0):
    """数字在上、单位并入说明行——避免大小字号混排时的基线重叠。"""
    text(slide, x, y, w, 0.42, num, size=num_size, bold=True, color=num_color, font=FONT_L)
    cap = f"{unit} · {label}" if label else unit
    if cap:
        text(slide, x, y + 0.42, w, 0.20, cap, size=7.5, color=label_color)


# ---------------- 建 PPT ----------------
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

S1 = os.path.join(SHOTS, "01-agent-chat.png")
S2 = os.path.join(SHOTS, "judge-1-chat.png")
S3 = os.path.join(SHOTS, "03-studio.png")
S4 = os.path.join(SHOTS, "09-stats.png")


# ============ 01 封面 ============
def slide1():
    s = prs.slides.add_slide(BLANK)
    bg(s, INK)
    # 装饰圆
    c = rect(s, 10.2, -2.2, 6.2, 6.2, fill=None, line=RGBColor(0x23, 0x23, 0x24), lw=1.0, radius=0.5)
    c2 = rect(s, 11.4, 3.6, 4.2, 4.2, fill=None, line=RGBColor(0x1B, 0x1B, 0x1C), lw=1.0, radius=0.5)

    text(s, M_L, 0.85, 7.4, 0.24, "HUAWEI DEVELOPER COMPETITION  ·  FUJIAN 2026", size=9, bold=True,
         color=RGBColor(0x56, 0x86, 0xFE), spc=140)
    text(s, M_L, 1.32, 6.4, 1.1, "OpenStudy", size=52, bold=True, color=WHITE, font=FONT_L)
    text(s, M_L, 2.42, 6.4, 0.62, "把不会的，问到会", size=26, bold=True, color=WHITE)
    line(s, M_L, 3.22, M_L + 1.1, 3.22, color=BLUE, width=2.5)
    text(s, M_L, 3.45, 5.9, 1.0,
         "会读你资料的 AI 学习智能体\n从一份 Word，到一套能刷的题、一份下一步计划",
         size=13.5, color=RGBColor(0xAD, 0xB2, 0xB8), line_spacing=1.45)

    # 元信息
    y = 5.05
    meta = [
        ("赛道", "AI + 教育 · 垂直智能体"),
        ("交付", "可演示作品 / 10 页方案 / GitCode 源码 / 演示视频"),
        ("开发", "华为云码道 CodeArts 代码智能体 · 全程留痕"),
    ]
    for k, v in meta:
        text(s, M_L, y, 0.62, 0.24, k, size=8.5, color=RGBColor(0x61, 0x66, 0x6B), spc=80)
        text(s, M_L + 0.72, y - 0.02, 5.6, 0.26, v, size=10, color=RGBColor(0xD3, 0xD6, 0xDB))
        y += 0.42

    text(s, M_L, 6.55, 7.0, 0.24, "OpenStudy  ·  在线学习平台  ·  FastAPI + Vue 3", size=8,
         color=RGBColor(0x43, 0x45, 0x4A))

    # 截图卡
    pic_card(s, S1, 7.62, 1.18, 5.05, 3.18)
    rect(s, 7.62, 1.18, 5.05, 3.18, fill=None, line=RGBColor(0x2C, 0x2C, 0x2E), lw=1.0, radius=0.02)
    text(s, 7.62, 4.52, 5.05, 0.24, "真实运行截图  ·  智能体对话即入口", size=8, color=RGBColor(0x61, 0x66, 0x6B))

    # 三个小指标
    stats = [("31", "个公开题库"), ("14", "个智能体工具"), ("3,166", "次码道工具调用")]
    x = 7.62
    for num, lab in stats:
        text(s, x, 5.15, 1.62, 0.5, num, size=22, bold=True, color=WHITE, font=FONT_L)
        text(s, x, 5.66, 1.62, 0.24, lab, size=8.5, color=RGBColor(0x81, 0x85, 0x8C))
        x += 1.72
    line(s, 7.62, 6.06, 12.67, 6.06, color=RGBColor(0x2C, 0x2C, 0x2E), width=0.5)
    text(s, 7.62, 6.18, 5.05, 0.3, "gitcode.com/qi142410131/openstudy-fastapi-ts", size=9,
         color=RGBColor(0x56, 0x86, 0xFE))


# ============ 02 场景与痛点 ============
def slide2():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "01  /  场景价值 · 25%", "谁的问题：不是没题刷，是「资料变不成题、错题没人管」",
           "面向高校学生与自学者——他们手里有资料，缺的是一条把资料变成练习、再把练习变成下一步的路径。")

    cards = [
        ("01", "资料在硬盘里，题在纸上",
         "课件、真题、教材大多是 Word / PDF / 扫描件；手动录 20 道题约 1 小时，公式和图还会丢。",
         "耗时", AMBER),
        ("02", "错题越积越多，掌握率起不来",
         "演示账号真实数据：累计作答 299 次、正确率 49.5%、待消灭错题 48 道、已掌握 0 道。",
         "无效", AMBER),
        ("03", "通用 AI 只会聊，不会「办事」",
         "能解释知识点，却不能建题库、不能出卷、不能改错题本——聊完还得自己动手。",
         "断层", AMBER),
    ]
    x = M_L
    cw = (CW - 2 * 0.22) / 3
    for num, t, d, tag, col in cards:
        rect(s, x, 2.02, cw, 2.28, fill=BG, line=LINE, lw=0.75, radius=0.05)
        text(s, x + 0.28, 2.24, 0.8, 0.3, num, size=11, bold=True, color=col, spc=60)
        text(s, x + 0.28, 2.60, cw - 0.56, 0.62, t, size=15, bold=True, color=INK, line_spacing=1.25)
        text(s, x + 0.28, 3.28, cw - 0.56, 0.92, d, size=10.5, color=GRAY, line_spacing=1.45)
        x += cw + 0.22

    # 用户原话
    rect(s, M_L, 4.58, CW, 0.86, fill=BLUE_L, line=None, radius=0.06)
    rect(s, M_L, 4.58, 0.055, 0.86, fill=BLUE)
    text(s, M_L + 0.30, 4.74, 1.0, 0.28, "用户原话", size=8.5, bold=True, color=BLUE, spc=60)
    text(s, M_L + 1.14, 4.70, CW - 1.5, 0.5,
         "“我只想把我自己的资料变成能刷的题，并且告诉我下一步练什么。”",
         size=13, bold=True, color=BLUE_D, line_spacing=1.2)

    # 结论
    text(s, M_L, 5.68, 4.4, 0.26, "所以我们把产品做成三件事", size=9.5, bold=True, color=GRAY2, spc=60)
    outs = [("资料即题库", "上传即解析入库，AI 依据原文出题并标注出处"),
            ("练习即诊断", "每一道错题都进错题本，统计到知识点"),
            ("对话即操作", "一句话建库、组卷、排计划，全部落库")]
    x = M_L
    ow = (CW - 2 * 0.2) / 3
    for t, d in outs:
        line(s, x, 5.98, x + ow - 0.2, 5.98, color=INK, width=1.2)
        text(s, x, 6.08, ow - 0.2, 0.28, t, size=11.5, bold=True, color=INK)
        text(s, x, 6.38, ow - 0.2, 0.4, d, size=9, color=GRAY, line_spacing=1.3)
        x += ow + 0.2
    footer(s, 2)


# ============ 03 产品方案 ============
def slide3():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "02  /  产品方案", "我们的答案：对话即入口，一条从「资料」到「下一步」的闭环",
           "不做又一个聊天框，也不做又一个题库——把 AI 的能力接到真实业务表上，每一步都可执行、可回滚。")

    # 流程
    steps = [("①", "上传资料", "Word / PDF / MD\n含公式与内嵌图"),
             ("②", "解析入库", "分块 · 向量化\n混合检索可命中"),
             ("③", "建库出题", "按原文出题\n标注出处"),
             ("④", "刷题模考", "顺序 / 随机 / 错题 / 模拟考"),
             ("⑤", "诊断计划", "薄弱点 TOP5\n下一步练什么")]
    x = M_L
    sw = 2.02
    gap = (CW - 5 * sw) / 4
    for i, (n, t, d) in enumerate(steps):
        fill = INK if i == 2 else BG
        tc = WHITE if i == 2 else INK
        dc = RGBColor(0xAD, 0xB2, 0xB8) if i == 2 else GRAY
        rect(s, x, 1.98, sw, 1.30, fill=fill, line=None if i == 2 else LINE, lw=0.75, radius=0.06)
        text(s, x + 0.20, 2.10, 0.4, 0.22, n, size=9, bold=True, color=BLUE if i == 2 else GRAY3)
        text(s, x + 0.20, 2.36, sw - 0.4, 0.26, t, size=12.5, bold=True, color=tc)
        text(s, x + 0.20, 2.66, sw - 0.4, 0.5, d, size=8.5, color=dc, line_spacing=1.3)
        if i < 4:
            text(s, x + sw + 0.04, 2.48, gap - 0.08, 0.3, "→", size=13, color=GRAY3, align=PP_ALIGN.CENTER)
        x += sw + gap

    # 回环箭头文字
    text(s, M_L, 3.36, CW, 0.24, "↑  诊断结果回到「③ 建库出题」：只练不会的，题越刷越少", size=8.5, color=GRAY2)

    # 左：三个不是
    text(s, M_L, 3.78, 5.4, 0.26, "三个「不是」", size=9.5, bold=True, color=GRAY2, spc=60)
    items = [("不是又一个聊天框", "对话里的每个动作都调用真实工具、写入真实数据表"),
             ("不是又一个题库", "题库只是入口，真正卖点是「资料→题→诊断→计划」的闭环"),
             ("不是演示脚本", "31 个公开题库、299 次作答、24 天学习记录都在库里")]
    y = 4.10
    for t, d in items:
        rect(s, M_L, y + 0.06, 0.10, 0.10, fill=BLUE, radius=0.5)
        text(s, M_L + 0.24, y, 5.0, 0.26, t, size=11, bold=True, color=INK)
        text(s, M_L + 0.24, y + 0.27, 5.1, 0.34, d, size=9, color=GRAY, line_spacing=1.3)
        y += 0.82

    # 右：截图
    pic_card(s, S2, 6.55, 3.72, 5.95, 3.10)
    rect(s, 6.55, 3.72, 5.95, 3.10, fill=None, line=LINE, lw=0.75, radius=0.02)
    text(s, 6.55, 6.86, 5.95, 0.2, "真实截图：一句话让两位专家并行建库 + 写题，写操作先确认后执行", size=7.5, color=GRAY2)
    footer(s, 3)


# ============ 04 产品完成度 ============
def slide4():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "03  /  产品完成度 · 25%", "四条主线，现场都能走通",
           "不是功能清单堆砌——每条主线都闭环到「用户能拿到结果」，并且有探针守着不回归。")

    cards = [
        ("A", "题库与考试", "建库 / 组卷 / 顺序·随机·错题·模拟考；答题卡、自动判分、考试记录、再考一次、跨账号隔离",
         "20+ 接口", BLUE),
        ("B", "知识库 RAG", "Word/PDF/MD 解析 → 分块 → 向量化 → 混合检索；教材原文出题、标注出处、图文并排预览",
         "bge-m3 · 1024 维", BLUE),
        ("C", "学习闭环", "错题本、收藏、斩题、笔记（草稿自动保存）、学习统计、薄弱知识点 TOP5 与难度分布",
         "31 题库 / 299 作答", GREEN),
        ("D", "智能体", "多专家并行、写操作二次确认、一键回退、跨会话长期记忆、对话内导出题库",
         "14 工具 / 5 专家", GREEN),
    ]
    x0, y0 = M_L, 1.96
    cw, ch = 3.92, 2.02
    for i, (tag, t, d, badge, col) in enumerate(cards):
        x = x0 + (i % 2) * (cw + 0.16)
        y = y0 + (i // 2) * (ch + 0.18)
        rect(s, x, y, cw, ch, fill=BG, line=LINE, lw=0.75, radius=0.05)
        text(s, x + 0.24, y + 0.18, 0.32, 0.26, tag, size=9.5, bold=True, color=col, spc=40)
        text(s, x + 0.58, y + 0.16, cw - 0.82, 0.30, t, size=14, bold=True, color=INK)
        text(s, x + 0.24, y + 0.62, cw - 0.48, 1.0, d, size=9.5, color=GRAY, line_spacing=1.45)
        text(s, x + 0.24, y + 1.66, cw - 0.48, 0.22, badge, size=8, bold=True, color=col, spc=30)

    # 右：两张截图
    pic_card(s, S3, 8.98, 1.96, 3.50, 2.02)
    rect(s, 8.98, 1.96, 3.50, 2.02, fill=None, line=LINE, lw=0.75, radius=0.02)
    text(s, 8.98, 4.02, 3.50, 0.2, "题库搭建：AI 生成 / 发布 / 可见性", size=7.5, color=GRAY2)
    pic_card(s, S4, 8.98, 4.28, 3.50, 2.02)
    rect(s, 8.98, 4.28, 3.50, 2.02, fill=None, line=LINE, lw=0.75, radius=0.02)
    text(s, 8.98, 6.34, 3.50, 0.2, "学习统计：正确率 / 趋势 / 薄弱点", size=7.5, color=GRAY2)

    # 底部指标条
    line(s, M_L, 6.42, 8.72, 6.42, color=LINE, width=0.5)
    text(s, M_L, 6.54, CW, 0.24, "32", size=13, bold=True, color=INK)
    text(s, M_L + 0.02, 6.82, 2.0, 0.2, "个 API 模块", size=7.5, color=GRAY2)
    for j, (n, l) in enumerate([("156", "个后端源文件 / 4.9 万行"), ("45", "个自动化探针"), ("200+", "次 Git 提交")]):
        text(s, M_L + (j + 1) * 2.06, 6.54, 2.0, 0.26, n, size=13, bold=True, color=INK)
        text(s, M_L + (j + 1) * 2.06 + 0.02, 6.82, 2.0, 0.2, l, size=7.5, color=GRAY2)
    footer(s, 4)


# ============ 05 技术架构 + 创新 ============
def slide5():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "04  /  技术实现 · 20%   ·   创新性 · 15%", "轻量自研，把「可控」写进架构",
           "不引入重型编排框架：路由零 LLM 成本、工具契约统一、写操作可确认可回退——这三件事决定了它能不能真用。")

    # 左：分层架构
    bx, bw = M_L, 7.28
    bands = [
        ("前端  Vue 3 + TS", "对话页 / 题库 / 知识库抽屉 / 学习统计  ·  单一设计令牌体系", WHITE, LINE),
        ("流式网关", "纯 ASGI 中间件 · SSE 逐字下发（修掉「攒批」导致的假流式）", WHITE, LINE),
        ("Agent Runtime", "Orchestrator（关键词路由，零 LLM） → 5 类专家子代理 → 14 个工具（统一 {success,error,hint} 信封）", BLUE_L, None),
        ("RAG 检索", "解析 → jieba 分块 → bge-m3（1024 维）→ HNSW 余弦 + tsvector GIN → RRF（k=60）", WHITE, LINE),
        ("数据层", "PostgreSQL + pgvector  ·  Redis  ·  本地/对象存储", WHITE, LINE),
    ]
    y = 1.96
    for i, (t, d, fill, ln) in enumerate(bands):
        hgt = 0.80
        rect(s, bx, y, bw, hgt, fill=fill, line=ln, lw=0.75, radius=0.04)
        text(s, bx + 0.22, y + 0.12, 1.42, 0.24, t, size=9.5, bold=True, color=BLUE_D if fill == BLUE_L else INK)
        text(s, bx + 1.70, y + 0.10, bw - 1.92, 0.62, d, size=8.5,
             color=BLUE_D if fill == BLUE_L else GRAY, line_spacing=1.35)
        if i < len(bands) - 1:
            text(s, bx + 0.55, y + hgt - 0.02, 0.2, 0.14, "▾", size=6, color=GRAY3)
        y += hgt + 0.075

    # 右：三个创新
    text(s, 8.48, 1.94, 4.0, 0.26, "三个能讲清楚的创新点", size=9.5, bold=True, color=GRAY2, spc=60)
    inns = [
        ("01", "多专家并行，单专家隐身", "一轮里并发多个 delegate_task；≥2 个专家才展开「工作组」卡片，1 个就融进主答，不打扰用户。"),
        ("02", "让知识库「看得见图」", "文档内嵌图按【图N】归属到分块；喂模型用 data URL，落盘路径只用于回插——扫描件也能出带图的题。"),
        ("03", "写操作可确认、可回退", "confirm_ticket 二次确认 + 题库/题目/分享/智能体四域整树快照；回退不了的操作必须明确告知。"),
    ]
    y = 2.26
    for n, t, d in inns:
        text(s, 8.48, y, 0.36, 0.24, n, size=9, bold=True, color=BLUE, spc=40)
        text(s, 8.84, y - 0.03, 3.6, 0.26, t, size=11, bold=True, color=INK)
        text(s, 8.84, y + 0.27, 3.60, 0.78, d, size=8.5, color=GRAY, line_spacing=1.4)
        y += 1.28

    # 底部一行技术指标
    line(s, M_L, 6.42, W - M_L, 6.42, color=LINE, width=0.5)
    tech = [("0", "次 LLM 用于路由"), ("1,658", "次命令 / 5 天（码道）"), ("45", "个探针 · 零 LLM 成本"), ("14", "个已注册工具")]
    x = M_L
    for n, l in tech:
        text(s, x, 6.54, 2.5, 0.26, n, size=12.5, bold=True, color=INK)
        text(s, x + 0.02, 6.82, 2.5, 0.2, l, size=7.5, color=GRAY2)
        x += 2.86
    footer(s, 5)


# ============ 06 码道开发① ============
def slide6():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "05  /  怎么用码道开发 ①", "把 AI 编码变成一条可复核的流水线",
           "码道 CodeArts 不是「补全插件」，是一个能自己读仓库、改代码、跑命令、提交的编码智能体。我们把它当同事用，不当魔法用。")

    # 左：五步
    text(s, M_L, 1.94, 4.2, 0.24, "我们固定的五步工作法", size=9.5, bold=True, color=GRAY2, spc=60)
    steps = [
        ("01", "需求进对话", "一句话说清目标 + 验收标准。例：「给智能体加知识库写入工具，要处理解析等待」。"),
        ("02", "先读后写", "grep / read / glob 先定位到文件与行号，再动手；禁止凭空猜 API。"),
        ("03", "计划拆解", "todowrite 列步骤，一次只改一个可验证的点，改完立刻验证。"),
        ("04", "自测自修", "跑探针、type-check、接口联调；失败就自己改到过，不把问题丢回给人。"),
        ("05", "规范提交", "一个改动一个 commit，中文 conventional message，提交前扫密钥。"),
    ]
    y = 2.26
    for n, t, d in steps:
        text(s, M_L, y + 0.02, 0.42, 0.24, n, size=8.5, bold=True, color=BLUE, spc=30)
        text(s, M_L + 0.48, y - 0.02, 3.9, 0.26, t, size=10.5, bold=True, color=INK)
        text(s, M_L + 0.48, y + 0.25, 4.05, 0.52, d, size=8, color=GRAY, line_spacing=1.35)
        y += 0.83

    # 右：指标面板
    rect(s, 5.62, 1.94, 6.88, 2.62, fill=INK, line=None, radius=0.04)
    text(s, 5.90, 2.10, 6.3, 0.22, "码道真实开发数据  ·  2026-10-08 → 10-12  ·  本机数据目录可核", size=7.5,
         bold=True, color=RGBColor(0x81, 0x85, 0x8C), spc=40)
    nums = [("5", "天", "连续开发"), ("11", "个", "项目会话"), ("134", "轮", "需求对话"),
            ("3,166", "次", "工具调用"), ("183", "个", "文件被改写"), ("50", "次", "Git 提交")]
    for i, (n, u, l) in enumerate(nums):
        col = i % 3
        row = i // 3
        stat(s, 5.90 + col * 2.16, 2.40 + row * 0.98, n, u, l, num_size=20, num_color=WHITE,
             unit_color=RGBColor(0xAD, 0xB2, 0xB8), label_color=RGBColor(0x61, 0x66, 0x6B), w=1.95)

    # 工具构成条形图
    text(s, 5.62, 4.72, 6.8, 0.24, "工具调用构成：读得比写多，验证比生成多", size=9, bold=True, color=INK)
    bars = [("bash  命令 / 测试 / Git", 1658, INK), ("read  读文件", 629, BLUE),
            ("edit  改代码", 476, RGBColor(0x7A, 0xAA, 0xFF)), ("write  新建文件", 191, RGBColor(0xB7, 0xC8, 0xFE)),
            ("grep  全局检索", 102, GRAY3), ("其他（todo/glob/…）", 110, LINE)]
    y = 5.02
    maxv = 1658
    for lab, v, col in bars:
        text(s, 5.62, y - 0.01, 2.18, 0.2, lab, size=7.5, color=GRAY)
        barw = max(0.05, (v / maxv) * 3.52)
        rect(s, 7.86, y + 0.02, barw, 0.115, fill=col, radius=None)
        text(s, 7.86 + barw + 0.08, y - 0.01, 0.6, 0.2, f"{v:,}", size=7.5, bold=True, color=INK)
        y += 0.245

    # 结论
    rect(s, M_L, 6.18, CW, 0.5, fill=BLUE_L, line=None, radius=0.06)
    text(s, M_L + 0.22, 6.29, CW - 0.4, 0.28,
         "629 次读 + 667 次改 + 1,658 次命令 = 真实的「读—改—验—提交」，不是「一键生成」。",
         size=9.5, bold=True, color=BLUE_D)
    footer(s, 6)


# ============ 07 码道开发② 案例 ============
def slide7():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "06  /  怎么用码道开发 ②", "案例：多智能体协同，5 天从 0 到 1",
           "选一个最难的模块拆开讲——因为它能同时证明：码道读得懂老代码、改得动核心链路、还能自己验证。")

    # 时间线
    nodes = [
        ("10-09 16:49", "提需求", "「加知识库/文档创建工具，并平衡解析等待体验」——一句话 + 验收标准。"),
        ("10-09 16:50", "先读代码", "80 秒内连读 registry.py / loop.py / rag_write_service.py / document_pipeline.py。"),
        ("10-09 20:12", "并行调研", "派 explore 子代理调研「多智能体改造可行点」，产出对比与结论。"),
        ("10-09 20:52", "落地实现", "多智能体协同：官方角色切分 + 用户自定义智能体。"),
        ("10-11 01:49", "修到能跑", "定位并修掉「子代理工具集被用户开关扣空」的真因。"),
    ]
    x = M_L
    nw = (CW - 4 * 0.14) / 5
    line(s, M_L + 0.1, 2.42, W - M_L - 0.1, 2.42, color=LINE, width=1.0)
    for i, (t, h, d) in enumerate(nodes):
        cx = x + nw / 2
        rect(s, cx - 0.075, 2.345, 0.15, 0.15, fill=BLUE if i in (0, 3) else WHITE, line=BLUE, lw=1.2, radius=0.5)
        text(s, x, 1.98, nw, 0.22, t, size=7.5, bold=True, color=BLUE, align=PP_ALIGN.CENTER, spc=20)
        text(s, x, 2.64, nw, 0.24, h, size=10, bold=True, color=INK, align=PP_ALIGN.CENTER)
        text(s, x + 0.04, 2.94, nw - 0.08, 0.78, d, size=7.5, color=GRAY, align=PP_ALIGN.CENTER, line_spacing=1.35)
        x += nw + 0.14

    # 左：真实提示词
    rect(s, M_L, 3.90, 5.62, 2.42, fill=BG, line=LINE, lw=0.75, radius=0.04)
    text(s, M_L + 0.24, 4.04, 4.9, 0.22, "真实提示词（原文，来自码道会话记录）", size=7.5, bold=True, color=GRAY2, spc=30)
    text(s, M_L + 0.24, 4.32, 5.14, 1.70,
         "「在 openstudy-fastapi 这个 FastAPI 项目里，我要评估『多智能体协同』改造的可行性。\n"
         "请调研并回报以下具体事实（要给出文件路径:行号）：\n"
         "1. registry.py 的 definitions(exclude=...) 是否支持按工具名裁剪？\n"
         "2. …能否直接复用成『按 Agent 角色挂不同工具』？」",
         size=8.5, color=INK2, line_spacing=1.5)
    text(s, M_L + 0.24, 5.94, 5.14, 0.24, "—— 会话 ses_edf6bde6effe… · 2026-10-09 20:12", size=7, color=GRAY3)

    # 右：结果
    rect(s, 6.62, 3.90, 5.88, 2.42, fill=INK, line=None, radius=0.04)
    text(s, 6.88, 4.04, 5.3, 0.22, "同样的 5 天，产出的是可运行的代码，不是 demo 片段", size=7.5, bold=True,
         color=RGBColor(0x81, 0x85, 0x8C), spc=20)
    res = [("11", "个文件", "agent 目录改动"), ("+2,666", "行", "agent 核心新增"),
           ("5", "类", "专家角色（+兜底）"), ("14", "个", "在册工具")]
    for i, (n, u, l) in enumerate(res):
        stat(s, 6.88 + i * 1.38, 4.36, n, u, l, num_size=15, num_color=WHITE,
             unit_color=RGBColor(0xAD, 0xB2, 0xB8), label_color=RGBColor(0x61, 0x66, 0x6B), w=1.32)
    line(s, 6.88, 5.10, 12.26, 5.10, color=RGBColor(0x2C, 0x2C, 0x2E), width=0.5)
    text(s, 6.88, 5.22, 5.3, 0.92,
         "新增 roles / orchestrator / sub_agent / delegate / knowledge_ingest；\n"
         "一轮内并发多专家、单专家自动折叠、写操作先确认——这些是用户能立刻感知的差异。",
         size=8.5, color=RGBColor(0xD3, 0xD6, 0xDB), line_spacing=1.45)

    # 底部：为什么这能证明
    line(s, M_L, 6.52, W - M_L, 6.52, color=LINE, width=0.5)
    text(s, M_L, 6.62, CW, 0.22,
         "可复核：会话 ID、时间戳、工具调用、提交记录四者可对齐——评审可现场按 ID 抽查任意一轮。",
         size=8.5, color=GRAY)
    footer(s, 7)


# ============ 08 码道开发③ 沉淀 ============
def slide8():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "07  /  怎么用码道开发 ③", "让 AI 写得快，也留得下：四类工程沉淀",
           "真正拉开差距的不是「生成了多少行」，而是「下一次还能不能复用、改坏了能不能立刻发现」。")

    cards = [
        ("110", "条", "项目记忆", "本机码道记忆库共 110 份，已脱敏归档 91 份进仓库。每条带 Why / How to apply：为什么这么做。"),
        ("4", "个", "自建 Skill", "GitCode 操作、仓库 wiki、跨仓流水线——把团队流程固化成码道可执行的技能。"),
        ("45", "个", "自动化探针", "改完就跑，零 LLM 成本；本轮 agent 相关探针 32/32、timeline 10/10 全过才提交。"),
        ("11", "个", "项目会话", "codearts session / export 随时导出原始对话；5 天连续运行日志同目录留存。"),
    ]
    x = M_L
    cw = (CW - 3 * 0.18) / 4
    for n, u, t, d in cards:
        rect(s, x, 1.96, cw, 2.10, fill=BG, line=LINE, lw=0.75, radius=0.05)
        rich(s, x + 0.22, 2.14, cw - 0.44, 0.52, [{"text": n, "size": 26, "bold": True, "color": INK, "font": FONT_L},
                                                  {"text": " " + u, "size": 8.5, "color": GRAY2}])
        text(s, x + 0.22, 2.68, cw - 0.44, 0.26, t, size=10.5, bold=True, color=INK)
        text(s, x + 0.22, 3.00, cw - 0.44, 0.90, d, size=8, color=GRAY, line_spacing=1.45)
        x += cw + 0.18

    # 人机分工
    text(s, M_L, 4.38, 5.0, 0.24, "人机分工：我们负责「对不对」，码道负责「快不快」", size=9.5, bold=True, color=GRAY2, spc=40)

    # 左栏
    rect(s, M_L, 4.70, 5.60, 1.62, fill=WHITE, line=LINE, lw=0.75, radius=0.04)
    text(s, M_L + 0.24, 4.84, 4.9, 0.24, "人  ·  产品与验收", size=9, bold=True, color=INK, spc=40)
    for i, t in enumerate(["定义真实场景与验收标准（什么算「能用」）",
                           "产品决策：专家怎么命名、什么时候折叠、回退怎么提示",
                           "真机验收：预置数据、网络、演示节奏、合规底线"]):
        text(s, M_L + 0.24, 5.14 + i * 0.36, 5.1, 0.3, "· " + t, size=8, color=GRAY)

    # 右栏
    rect(s, 6.66, 4.70, 5.84, 1.62, fill=INK, line=None, radius=0.04)
    text(s, 6.90, 4.84, 5.2, 0.24, "码道  ·  工程执行", size=9, bold=True, color=WHITE, spc=40)
    for i, t in enumerate(["读码定位：跨 200+ 文件找到真正的调用点与数据流",
                           "方案对比：给出 2–3 种做法 + 取舍，再按确认的方案改",
                           "批量改写、自测修复、规范提交，一次到位"]):
        text(s, 6.90, 5.14 + i * 0.36, 5.3, 0.3, "· " + t, size=8, color=RGBColor(0xAD, 0xB2, 0xB8))

    # 证据位置
    line(s, M_L, 6.52, W - M_L, 6.52, color=LINE, width=0.5)
    rich(s, M_L, 6.62, CW, 0.24, [
        {"text": "证据位置：", "size": 8, "bold": True, "color": GRAY2},
        {"text": "doc/dev-process/", "size": 8, "color": BLUE},
        {"text": "   logs（每日日志） · sessions（6 份完整会话） · memory（91 份决策笔记） · codebase（索引记录）", "size": 8, "color": GRAY},
    ])
    footer(s, 8)


# ============ 09 现场演示 ============
def slide9():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "08  /  现场演示", "3 分钟，四个「看得见」",
           "演示不追求功能全，只追求一件事：让评委在 3 分钟内看到「AI 真的动了我的数据」。")

    stages = [
        ("0:00 – 0:40", "看得见处理", "上传一份含公式与图片的 Word/PDF",
         "知识库面板实时显示解析进度；完成后点开分块，公式与图片都在。", BLUE),
        ("0:40 – 1:30", "看得见结果", "一句话：「用这份资料建个题库，出 10 道题」",
         "两位专家并行：一个建库、一个出题；写操作先弹确认卡，确认后题目落库。", GREEN),
        ("1:30 – 2:20", "看得见闭环", "开考 → 交卷 → 错题本 → 学习统计",
         "自动判分；错题进本；薄弱点 TOP5 与难度分布立刻更新。", AMBER),
        ("2:20 – 3:00", "看得见智能", "「按我的薄弱点排一周复习计划」",
         "智能体读取真实统计，给出按天、按知识点的计划，并可一键写进错题本。", INK),
    ]
    x = M_L
    cw = (CW - 3 * 0.16) / 4
    for time, t, act, res, col in stages:
        rect(s, x, 1.96, cw, 2.62, fill=BG, line=LINE, lw=0.75, radius=0.05)
        rect(s, x, 1.96, cw, 0.045, fill=col)
        text(s, x + 0.22, 2.12, cw - 0.44, 0.22, time, size=8, bold=True, color=col, spc=40)
        text(s, x + 0.22, 2.38, cw - 0.44, 0.26, t, size=12, bold=True, color=INK)
        text(s, x + 0.22, 2.72, cw - 0.44, 0.60, act, size=8.5, color=INK2, line_spacing=1.4)
        line(s, x + 0.22, 3.40, x + cw - 0.22, 3.40, color=LINE, width=0.5)
        text(s, x + 0.22, 3.50, cw - 0.44, 0.96, res, size=8, color=GRAY, line_spacing=1.45)
        x += cw + 0.16

    # 稳定性保障
    text(s, M_L, 4.94, 6.0, 0.24, "演示稳定性：把「万一」提前处理掉", size=9.5, bold=True, color=GRAY2, spc=40)
    guards = [
        ("预置数据", "演示账号预置 31 个题库 / 299 条作答，开场即有内容可看"),
        ("提前热身", "演示前 2 小时跑 daily-start.sh，预热模型与向量检索（首字 45s → 2.8s）"),
        ("双保险", "现场网络异常时切本地录屏；关键链路均有探针，演示前 10 分钟全跑一遍"),
    ]
    x = M_L
    gw = (CW - 2 * 0.18) / 3
    for t, d in guards:
        text(s, x, 5.28, gw, 0.24, "· " + t, size=9, bold=True, color=INK)
        text(s, x + 0.14, 5.58, gw - 0.14, 0.62, d, size=8, color=GRAY, line_spacing=1.4)
        x += gw + 0.18

    rect(s, M_L, 6.34, CW, 0.34, fill=BLUE_L, line=None, radius=0.08)
    text(s, M_L + 0.20, 6.40, CW - 0.4, 0.22,
         "一句话记住：先给结果，再讲技术；每按一次按钮，都要有一处数据变化被看见。", size=8.5, bold=True, color=BLUE_D)
    footer(s, 9)


# ============ 10 迭代 + 团队 ============
def slide10():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, "09  /  迭代计划  ·  路演与团队 · 15%", "下一步：把「能用」做成「好用」",
           "赛后 30 天与 90 天的路线已经排好——每一条都来自真实使用中发现的边界，不是拍脑袋的功能清单。")

    cols = [
        ("已交付 · 本轮", "DONE", [
            "多智能体协同（指挥家 + 5 类专家）",
            "知识库图文解析与原文出题",
            "模拟考试、错题本、学情诊断",
            "写操作确认 + 四域回退",
        ], INK),
        ("30 天 · 补能力", "NEXT", [
            "扫描版 PDF 接入 OCR，解决教材形态",
            "多模态向量：以图搜图、文字搜图",
            "移动端适配与消息中心",
            "评测期限流策略与稳定性加固",
        ], BLUE),
        ("90 天 · 上规模", "LATER", [
            "校园版多租户与教师端组卷分析",
            "对象存储直传（OBS），摆脱本地目录",
            "题库/笔记开放 API 与插件机制",
            "接入更多华为云能力（OCR / 搜索）",
        ], GRAY2),
    ]
    x = M_L
    cw = (CW - 2 * 0.20) / 3
    for t, tag, items, col in cols:
        rect(s, x, 1.96, cw, 2.92, fill=BG if tag != "NEXT" else BLUE_L, line=LINE if tag != "NEXT" else None,
             lw=0.75, radius=0.05)
        text(s, x + 0.24, 2.12, cw - 0.48, 0.22, tag, size=7.5, bold=True, color=col, spc=80)
        text(s, x + 0.24, 2.38, cw - 0.48, 0.26, t, size=12.5, bold=True, color=INK)
        y = 2.80
        for it in items:
            text(s, x + 0.24, y, 0.14, 0.22, "·", size=9, color=col)
            text(s, x + 0.40, y, cw - 0.64, 0.44, it, size=8, color=GRAY, line_spacing=1.35)
            y += 0.50
        x += cw + 0.20

    # 团队
    text(s, M_L, 5.10, 5.6, 0.24, "团队分工（按真实职责，可一人多岗）", size=9.5, bold=True, color=GRAY2, spc=40)
    roles = [("产品 / 后端", "场景定义、Agent 与 RAG 后端、数据与部署"),
             ("前端 / 交互", "对话体验、题库与知识库面板、设计令牌"),
             ("测试 / 演示", "探针与回归、演示脚本、材料与答辩")]
    x = M_L
    rw = (CW - 2 * 0.18) / 3
    for t, d in roles:
        line(s, x, 5.42, x + rw - 0.18, 5.42, color=INK, width=1.0)
        text(s, x, 5.50, rw - 0.18, 0.24, t, size=9.5, bold=True, color=INK)
        text(s, x, 5.76, rw - 0.18, 0.40, d, size=7.5, color=GRAY, line_spacing=1.35)
        x += rw + 0.18

    # 收尾
    line(s, M_L, 6.28, W - M_L, 6.28, color=LINE, width=0.5)
    text(s, M_L, 6.40, 6.4, 0.30, "谢谢。欢迎现场提问，也欢迎直接上手体验。", size=11, bold=True, color=INK)
    text(s, M_L, 6.74, 7.2, 0.22, "作品源码：gitcode.com/qi142410131/openstudy-fastapi-ts", size=8, color=GRAY2)
    rect(s, 9.60, 6.36, 3.0, 0.52, fill=INK, line=None, radius=0.10)
    text(s, 9.60, 6.48, 3.0, 0.28, "OpenStudy · 把不会的问到会", size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(s, 10)


for fn in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10):
    fn()

out = os.path.join(HERE, "OpenStudy-华为开发者大赛福建赛区-10页.pptx")
prs.save(out)
print("saved:", out, os.path.getsize(out))
