"""
Excel 导出工具
- 优先使用 openpyxl 生成 .xlsx
- openpyxl 不可用时回退到 csv 格式
- 返回 StreamingResponse 供 FastAPI 直接 return
"""
from fastapi.responses import StreamingResponse
from typing import List, Dict, Any, Optional
import io
import re


# openpyxl 禁止在单元格中写入的控制字符（与 openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE 一致）。
# 题库题干/解析来自历史导入数据时可能残留 0x0C（换页符）等字符，
# 直接写入会抛 IllegalCharacterError 导致整个导出接口 500，因此统一在写值前清洗。
ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def sanitize_cell_value(value: Any) -> Any:
    """剔除 Excel 不接受的控制字符；非字符串原样返回"""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def export_to_excel(
    rows: List[Dict[str, Any]],
    headers: List[Dict[str, str]],  # [{key: "fieldName", label: "列标题"}, ...]
    filename: str = "export",
) -> StreamingResponse:
    """导出数据为 Excel/CSV"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # 表头样式
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # 写表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=sanitize_cell_value(header["label"]))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写数据
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header["key"], "")
                # 处理特殊类型
                if value is None:
                    value = ""
                elif isinstance(value, (list, dict)):
                    value = str(value)
                value = sanitize_cell_value(value)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # 自动列宽
        for col_idx, header in enumerate(headers, 1):
            max_length = len(header["label"])
            for row_idx in range(2, len(rows) + 2):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_length + 4, 50)

        # 冻结首行
        ws.freeze_panes = "A2"

        # 保存到内存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"}
        )

    except ImportError:
        # openpyxl 不可用,回退到 CSV
        import csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([sanitize_cell_value(h["label"]) for h in headers])
        for row_data in rows:
            writer.writerow([sanitize_cell_value(row_data.get(h["key"], "")) for h in headers])

        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue().encode("utf-8-sig")]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}.csv"}
        )
