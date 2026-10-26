"""
文件解析服务 - 支持 TXT/MD/DOCX/PDF/DOC/RTF

- DOCX: python-docx 提取段落 + 图片（含**表格单元格内**的插图）+ 表格文本，OMML 公式转 LaTeX
- PDF : pypdf 提取文本 + **页内插图**（按「先文本、后该页图片」排块）
- DOC : 按**文件魔数**分流 —— ZIP 头的（改名成 .doc 的 docx）走 DOCX；
        真 OLE2 容器则按容器内的流名判断（`WordDocument`/`Workbook`/`PowerPoint Document`），
        老版 .doc 交给 `doc_converter` 借本机 Word 转成 docx 再走 DOCX 解析
- TXT/MD: 直接读取

返回结构化块: `{textContent, blocks, images, totalTextLength, fileName, error?}`

⚠️ 本模块是**全平台唯一的文档解析入口**。聊天 / Agent / ds-free / AI 搭建题库
四条链路最终都调 `parse_file()`，改进这里 = 四条链路同时生效（见 docs/openstudy-docs/v2的OS助手重构目标.md §6 Bug②）。

⚠️ `blocks` 是**有序**的：图片块插在它原本所在的文字位置之间，
`images` 则是同一批图片的平铺列表（顺序与 blocks 中出现的顺序一致）。
"""

import base64
import io
import re
import xml.etree.ElementTree as ET
from typing import Optional

# OMML / WordprocessingML 命名空间
_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

W_NS = _NS["w"]
M_NS = _NS["m"]


def _qn(ns: str, tag: str) -> str:
    """构造带命名空间的限定标签名 {ns}tag"""
    return f"{{{ns}}}{tag}"


# ============================================================
# OMML -> LaTeX 转换
# ============================================================
def _omml_to_latex(elem: ET.Element) -> str:
    """将 <m:oMath> 元素转换为 LaTeX 字符串(递归处理)。

    只实现常见结构:分数、根号、上下标、求和/积分、定界符、run。
    未识别元素回退为提取其纯文本。
    """
    return _convert_element(elem).strip()


def _convert_element(elem: ET.Element) -> str:
    """递归转换单个 OMML 元素"""
    tag = elem.tag
    # 去掉命名空间,只保留本地名
    if "}" in tag:
        tag = tag.split("}", 1)[1]

    if tag == "oMath":
        # 公式块:拼接所有子元素
        return "".join(_convert_element(child) for child in elem)

    if tag == "oMathPara":
        # 展示型公式块,内部含 oMath
        return "".join(_convert_element(child) for child in elem)

    if tag == "r":
        # run:取 <m:t> 文本
        return _run_text(elem)

    if tag == "t":
        return elem.text or ""

    if tag == "f":
        # 分数 \frac{num}{den}
        num = _find_child_local(elem, "num")
        den = _find_child_local(elem, "den")
        num_latex = _convert_element(num) if num is not None else ""
        den_latex = _convert_element(den) if den is not None else ""
        return f"\\frac{{{num_latex}}}{{{den_latex}}}"

    if tag == "rad":
        # 根号 \sqrt{content} 或 \sqrt[deg]{content}
        deg = _find_child_local(elem, "deg")
        e = _find_child_local(elem, "e")
        e_latex = _convert_element(e) if e is not None else ""
        if deg is not None and _has_content(deg):
            deg_latex = _convert_element(deg)
            return f"\\sqrt[{deg_latex}]{{{e_latex}}}"
        return f"\\sqrt{{{e_latex}}}"

    if tag == "sSup":
        # 上标 base^{sup}
        e = _find_child_local(elem, "e")
        sup = _find_child_local(elem, "sup")
        base = _convert_element(e) if e is not None else ""
        sup_latex = _convert_element(sup) if sup is not None else ""
        return f"{{{base}}}^{{{sup_latex}}}"

    if tag == "sSub":
        # 下标 base_{sub}
        e = _find_child_local(elem, "e")
        sub = _find_child_local(elem, "sub")
        base = _convert_element(e) if e is not None else ""
        sub_latex = _convert_element(sub) if sub is not None else ""
        return f"{{{base}}}_{{{sub_latex}}}"

    if tag == "sSubSup":
        # 上下标同时 base_{sub}^{sup}
        e = _find_child_local(elem, "e")
        sub = _find_child_local(elem, "sub")
        sup = _find_child_local(elem, "sup")
        base = _convert_element(e) if e is not None else ""
        sub_latex = _convert_element(sub) if sub is not None else ""
        sup_latex = _convert_element(sup) if sup is not None else ""
        return f"{{{base}}}_{{{sub_latex}}}^{{{sup_latex}}}"

    if tag == "nary":
        # n 元运算符(求和/积分等),通过 m:chr 属性判断符号
        chr_val = _get_nary_char(elem)
        sub = _find_child_local(elem, "sub")
        sup = _find_child_local(elem, "sup")
        e = _find_child_local(elem, "e")
        op = _nary_char_to_latex(chr_val)
        sub_latex = _convert_element(sub) if sub is not None and _has_content(sub) else ""
        sup_latex = _convert_element(sup) if sup is not None and _has_content(sup) else ""
        e_latex = _convert_element(e) if e is not None else ""
        result = op
        if sub_latex:
            result += f"_{{{sub_latex}}}"
        if sup_latex:
            result += f"^{{{sup_latex}}}"
        result += f" {{{e_latex}}}"
        return result

    if tag == "d":
        # 定界符 \left( content \right)
        e = _find_child_local(elem, "e")
        e_latex = _convert_element(e) if e is not None else ""
        beg = _get_delimiter(elem, "begChr", "(")
        end = _get_delimiter(elem, "endChr", ")")
        return f"\\left{beg} {e_latex} \\right{end}"

    if tag == "e" or tag == "num" or tag == "den" or tag == "deg" \
            or tag == "sub" or tag == "sup":
        # 容器元素:递归拼接子元素
        return "".join(_convert_element(child) for child in elem)

    if tag == "fPr" or tag == "radPr" or tag == "sSupPr" or tag == "sSubPr" \
            or tag == "sSubSupPr" or tag == "naryPr" or tag == "dPr" or tag == "rPr":
        # 属性节点:忽略
        return ""

    # 默认:递归拼接子元素
    return "".join(_convert_element(child) for child in elem)


def _run_text(elem: ET.Element) -> str:
    """提取 <m:r> 内 <m:t> 的文本"""
    texts = []
    for t in elem.iter(_qn(M_NS, "t")):
        if t.text:
            texts.append(t.text)
    return "".join(texts)


def _find_child_local(elem: ET.Element, local_name: str) -> Optional[ET.Element]:
    """查找直接子元素中本地名为 local_name 的第一个元素(任意命名空间)"""
    for child in elem:
        ctag = child.tag
        if "}" in ctag:
            ctag = ctag.split("}", 1)[1]
        if ctag == local_name:
            return child
    return None


def _has_content(elem: ET.Element) -> bool:
    """判断元素是否含有实际内容(文本或子元素)"""
    if (elem.text or "").strip():
        return True
    return len(list(elem)) > 0


def _get_nary_char(elem: ET.Element) -> str:
    """从 <m:nary> 的 <m:naryPr><m:chr m:val="..."/> 取运算符字符"""
    nary_pr = _find_child_local(elem, "naryPr")
    if nary_pr is None:
        return ""
    chr_elem = _find_child_local(nary_pr, "chr")
    if chr_elem is None:
        return ""
    return chr_elem.get(_qn(M_NS, "val"), "") or chr_elem.get("val", "") or ""


def _nary_char_to_latex(chr_val: str) -> str:
    """将 nary 运算符字符映射为 LaTeX 命令"""
    mapping = {
        "∑": "\\sum",
        "∫": "\\int",
        "∏": "\\prod",
        "∪": "\\bigcup",
        "∩": "\\bigcap",
        "∮": "\\oint",
        "∬": "\\iint",
        "∭": "\\iiint",
    }
    return mapping.get(chr_val, "\\sum")


def _get_delimiter(elem: ET.Element, attr_local: str, default: str) -> str:
    """从 <m:dPr> 取定界符字符(begChr/endChr)"""
    d_pr = _find_child_local(elem, "dPr")
    if d_pr is None:
        return default
    node = _find_child_local(d_pr, attr_local)
    if node is None:
        return default
    val = node.get(_qn(M_NS, "val"), "") or node.get("val", "")
    return val if val else default


# ============================================================
# 图片归一化
# ============================================================

#: 小于这个边长的图多半是图标 / 分隔线 / 项目符号，没有信息量（口径对齐 sp3）
_MIN_IMAGE_SIDE = 50
#: 插图长边上限：超过就等比缩小（请求体要能过上游，data URL 会膨胀 4/3）
_MAX_IMAGE_SIDE = 1600
#: 原格式已可直接渲染时的体积上限；超过就重编码压缩
_RAW_KEEP_BYTES = 2 * 1024 * 1024
#: 浏览器认得的图片格式 —— 这些可以原样透传，不重编码（保清晰、省 CPU）
_PASSTHROUGH_MIME = {
    "image/png",
    "image/jpeg",
    "image/jpg",
    "image/gif",
    "image/webp",
    "image/bmp",
}


def _mime_ext(mime: str) -> str:
    """image/jpeg -> jpg（文件名用）"""
    ext = (mime or "").split("/", 1)[-1].split(";")[0].strip().lower()
    return "jpg" if ext == "jpeg" else (ext or "png")


def _prepare_pil_image(image) -> Optional[tuple]:
    """PIL.Image → (bytes, mime, ext)；不合适的（太小 / 编码失败）返回 None

    统一重编码策略：有透明通道走 PNG（保住透明），其余走 JPEG（体积小得多）。
    """
    from PIL import Image

    try:
        # 先看尺寸再 load()：超大图光是解码就要几百 MB，能缩就尽早缩
        if image.width < _MIN_IMAGE_SIDE or image.height < _MIN_IMAGE_SIDE:
            return None
        if max(image.size) > _MAX_IMAGE_SIDE:
            image.thumbnail((_MAX_IMAGE_SIDE, _MAX_IMAGE_SIDE), Image.LANCZOS)
        else:
            image.load()
        buf = io.BytesIO()
        has_alpha = image.mode in ("RGBA", "LA") or (
            image.mode == "P" and "transparency" in image.info
        )
        if has_alpha:
            image.convert("RGBA").save(buf, format="PNG", optimize=True)
            return buf.getvalue(), "image/png", "png"
        image.convert("RGB").save(buf, format="JPEG", quality=88, optimize=True)
        return buf.getvalue(), "image/jpeg", "jpg"
    except Exception:
        return None


def _prepare_doc_image(blob: bytes, content_type: str = "") -> Optional[tuple]:
    """文档内嵌图片字节 → (bytes, mime, ext)；不可渲染 / 太小 / 解不开都返回 None

    三种情况：
      · 格式浏览器认、尺寸体积都 OK → **原样透传**（不重编码，保清晰）
      · 尺寸过大 / 体积过大 → 用 Pillow 等比缩小后重编码
      · EMF / WMF / TIFF 等浏览器渲染不了的 → 用 Pillow 转成 PNG/JPEG；转不了就丢弃
    """
    if not blob:
        return None
    mime = (content_type or "").lower().split(";")[0].strip()
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as probe:
            width, height = probe.size
            if (
                mime in _PASSTHROUGH_MIME
                and max(width, height) <= _MAX_IMAGE_SIDE
                and len(blob) <= _RAW_KEEP_BYTES
                and width >= _MIN_IMAGE_SIDE
                and height >= _MIN_IMAGE_SIDE
            ):
                return blob, mime, _mime_ext(mime)
            return _prepare_pil_image(probe)
    except Exception:
        return None


def _payload_to_data_url(blob: bytes, mime: str) -> str:
    return f"data:{mime};base64,{base64.b64encode(blob).decode('ascii')}"


# ============================================================
# 块收集器（docx / pdf 共用）
# ============================================================
#: Markdown 表格的一行：`| a | b | c |`
_RE_MD_TABLE_LINE = re.compile(r"^\s*\|.*\|\s*$")

#: Markdown 表格的对齐/分隔行：`| --- | :--: | --: |`
_RE_MD_TABLE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _split_markdown_tables(text: str) -> list:
    """
    把一段纯文本按 **Markdown 表格** 切开

    Returns:
        `[("text", str) | ("table", {"headers": [...], "rows": [[...]], "markdown": str})]`

    判定条件（两条都满足才算表格）：
      ① 连续 `|...|` 行
      ② 其中有**且仅有一行**是分隔行（全是 `-` / `:` / `|` / 空格）

    为什么需要：`.txt` / `.md` 输入走 `add_text` 整段写入，
    纯文本里的表格从此与 docx 表格一样被识别成 table 块，
    导出时才能渲染成真表格（否则竖线原样印在纸上）。
    """
    if "|" not in text or "\n" not in text:
        return []

    lines = text.split("\n")
    out: list = []
    buf: list = []
    i = 0

    def _flush() -> None:
        if buf:
            body = "\n".join(buf).strip()
            if body:
                out.append(("text", body))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        if _RE_MD_TABLE_LINE.match(line):
            # 收集连续的表格行
            j = i
            block = []
            while j < len(lines) and _RE_MD_TABLE_LINE.match(lines[j]):
                block.append(lines[j])
                j += 1
            # 至少要有 表头 + 分隔行（+ 可选数据行）
            if len(block) >= 2 and "|" in block[1]:
                sep_idx = [k for k, ln in enumerate(block) if _RE_MD_TABLE_SEP.match(ln)]
                if len(sep_idx) == 1:
                    sep = sep_idx[0]
                    header_cells = _md_cells(block[sep - 1]) if sep >= 1 else []
                    body_rows = [_md_cells(ln) for ln in block[sep + 1 :]]
                    if header_cells:
                        _flush()
                        out.append(
                            (
                                "table",
                                {
                                    "headers": header_cells,
                                    "rows": body_rows,
                                    "markdown": "\n".join(block),
                                },
                            )
                        )
                        i = j
                        continue
            # 不是表格 → 当作普通行
            buf.extend(block)
            i = j
            continue

        buf.append(line)
        i += 1

    _flush()
    return out


def _md_cells(line: str) -> list:
    """`| a | b |` → `["a", "b"]`（`\\|` 是转义竖线，不切）"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.replace("\\|", "|").strip() for c in re.split(r"(?<!\\)\|", s)]


class _BlockBuilder:
    """按出现顺序收集 text / image 块，并同步维护 `textContent` 与 `images`

    「图片落在文字之间」这件事只在 blocks 里体现；textContent 仍是纯文本
    （老调用方 / 没有 blocks 的降级路径要靠它）。
    """

    def __init__(self) -> None:
        self.blocks: list = []
        self.images: list = []
        self._text_parts: list = []

    def add_text(self, text: str) -> None:
        text = (text or "").strip()
        if not text:
            return

        # ── 2026-10-26 v3：整段文字里若有 Markdown 表格，拆成 text/table 块 ──
        #
        # 为什么在这里做：`parse_file` 处理 **.txt / .md** 输入时，
        # 走的是"整段 `add_text`"，从来不经 `_emit_table`（那是 docx 专用路径）。
        # 于是纯文本里的 `| X | 0 | 1 |` 依然是普通文字，导出时竖线原样印出。
        #
        # 只处理**含 `|` 且跨行**的情况，普通正文零额外开销。
        if "|" in text and "\n" in text:
            pieces = _split_markdown_tables(text)
            if len(pieces) > 1 or (pieces and pieces[0][0] == "table"):
                for kind, payload in pieces:
                    if kind == "table":
                        self.add_table(
                            payload["headers"],
                            payload["rows"],
                            markdown=payload.get("markdown", ""),
                        )
                    else:
                        body = str(payload).strip()
                        if body:
                            self.blocks.append(
                                {"type": "text", "content": body, "orderIndex": len(self.blocks)}
                            )
                            self._text_parts.append(body)
                return

        self.blocks.append(
            {"type": "text", "content": text, "orderIndex": len(self.blocks)}
        )
        self._text_parts.append(text)

    def add_table(
        self,
        headers: list,
        rows: list,
        *,
        markdown: str = "",
    ) -> None:
        """
        加入一个**表格块**（2026-10-26 v3 新增）

        在这之前，Word 表格被 `_emit_table` 压成一段 Markdown 文本塞进 `text` 块，
        下游（导出 Word/PDF）**完全不知道"这里原本是一张表"**，
        于是竖线原样印在纸上（实测题库 432 第 33 题）。

        现在同时产出两种表示：

        · `blocks` 里的 `{"type":"table", ...}` —— **结构化数据**，供导出层渲染真表格
        · `textContent` / 紧随其后的纯文本 —— **Markdown 表示**，供
          `knowledge_ingest` / `chat_engine` / 前端 marked 这些老调用方使用

        这样 6 个 `parse_file` 调用点**一行都不用改**（向后兼容），
        而导出层拿到的就是能直接渲染成 Word/PDF 原生表格的数据。
        """
        rows = [list(r) for r in (rows or [])]
        headers = list(headers or [])
        if not headers and not rows:
            return

        ncol = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
        if ncol <= 0:
            return
        headers = headers + [""] * (ncol - len(headers))
        rows = [r + [""] * (ncol - len(r)) for r in rows]

        index = len(self.blocks)
        self.blocks.append(
            {
                "type": "table",
                "content": markdown or "",
                "orderIndex": index,
                "headers": headers,
                "rows": rows,
            }
        )
        # 老调用方靠 textContent 取正文 → 表格仍以 Markdown 形式进纯文本
        if markdown:
            self._text_parts.append(markdown)

    def add_image(self, blob: bytes, content_type: str = "") -> bool:
        """加入一张**未归一化**的图片（docx 的 rel blob 走这里）"""
        prepared = _prepare_doc_image(blob, content_type)
        if not prepared:
            return False
        return self.add_payload(*prepared)

    def add_payload(self, payload: bytes, mime: str, ext: str = "") -> bool:
        """加入一张**已归一化**的图片（PDF 链路里 pypdf 已解好码，别再压一遍）"""
        if not payload:
            return False
        index = len(self.blocks)
        data_url = _payload_to_data_url(payload, mime)
        image_name = f"image{len(self.images) + 1}.{ext or _mime_ext(mime)}"
        self.blocks.append(
            {
                "type": "image",
                "content": "",
                "orderIndex": index,
                "imageUrl": data_url,
                "imageName": image_name,
            }
        )
        self.images.append(
            {"orderIndex": index, "imageUrl": data_url, "imageName": image_name}
        )
        return True

    @property
    def text_length(self) -> int:
        """与 `textContent` 的长度一致（含块之间的换行符），别改成各块长度之和"""
        return len("\n".join(self._text_parts))

    def result(self) -> dict:
        return {
            "textContent": "\n".join(self._text_parts),
            "blocks": self.blocks,
            "images": self.images,
            "totalTextLength": self.text_length,
        }


# ============================================================
# 主解析逻辑
# ============================================================
#: 各格式的文件头魔数（用于纠正「扩展名撒谎」的情况，最常见的是 docx 被改名成 .doc）
_MAGIC_ZIP = b"PK\x03\x04"  # docx/xlsx/pptx 都是 zip 容器
_MAGIC_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # 老版 .doc/.xls/.ppt 的 OLE2 复合文档
_MAGIC_PDF = b"%PDF-"


def _sniff_kind(content: bytes) -> str:
    """按文件头判断真实格式：'zip' / 'ole2' / 'pdf' / ''（认不出来）"""
    head = (content or b"")[:1024].lstrip()
    if head.startswith(_MAGIC_ZIP):
        return "zip"
    if head.startswith(_MAGIC_OLE2):
        return "ole2"
    if head.startswith(_MAGIC_PDF):
        return "pdf"
    return ""


_OLE2_DOC_ERROR = (
    "这是老版 .doc（OLE2 二进制格式），当前环境没有可用的转换器（需要 Windows + 装 MS Word + 装 pywin32）。"
    "请在 Word 里「另存为 .docx」后重新上传。"
)


def parse_file(filename: str, content: bytes) -> dict:
    """主入口：按**扩展名 + 文件头**分发解析。

    ⚠️ 只看扩展名会踩坑：用户常把 .docx 直接改名成 .doc（甚至 .pdf），
    只认扩展名会把明明能解析的文件判成「不支持」。所以这里先嗅探文件头，
    **文件头说了算**，扩展名只作兜底。

    Args:
        filename: 文件名（用于判断扩展名）
        content: 文件二进制内容

    Returns:
        {textContent, blocks, images, totalTextLength, fileName, error?}
    """
    ext = filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    kind = _sniff_kind(content)

    try:
        if kind == "zip":
            if ext in ("doc", "docx", "docm", ""):
                # 覆盖「用户把 docx 改名成 .doc」这个最常见的场景
                result = _parse_docx(content)
            else:
                # xlsx / pptx 也是 zip 容器，但本解析器只做 Word / PDF / 文本
                result = _empty_result(
                    f"这个文件是压缩包容器（.docx/.xlsx/.pptx 都是），扩展名是 .{ext}，"
                    "本解析器只支持 Word / PDF / 文本文件"
                )
        elif kind == "ole2":
            result = _parse_ole2(content)
        elif ext == "pdf" or kind == "pdf":
            result = _parse_pdf(content)
        elif ext == "docx":
            result = _parse_docx(content)
        elif ext in ("txt", "md", "markdown", "doc"):
            # .doc 走到这里说明文件头既不是 ZIP 也不是 OLE2 —— 多半是纯文本改的名
            result = _parse_txt(content)
        else:
            # 未知类型：尝试按文本读取
            result = _parse_txt(content)
    except Exception as e:  # noqa: BLE001 - 解析失败不该把整个请求打挂
        result = _empty_result(f"{type(e).__name__}: {e}")

    result["fileName"] = filename
    return result


def _empty_result(error: str = "") -> dict:
    result = {"textContent": "", "blocks": [], "images": [], "totalTextLength": 0}
    if error:
        result["error"] = error
    return result


def _parse_ole2(content: bytes) -> dict:
    """OLE2 复合文档：老版 .doc 借 Word COM 转成 docx 再解析；其余（.xls/.ppt）直接报错

    `.doc` / `.xls` / `.ppt` 文件头完全一样，所以要看容器里的流名来判断真实格式
    （见 `doc_converter.ole2_flavor`）。

    ⚠️ 转换是**阻塞**的（1~3s，最坏 `CONVERT_TIMEOUT_SECONDS`）。见
    `doc_converter.convert_to_docx` 的说明 —— 异步调用方要用 `asyncio.to_thread` 兜住。
    """
    from app.services.doc_converter import (
        convert_to_docx,
        is_available,
        ole2_flavor,
    )

    flavor = ole2_flavor(content)
    if flavor and flavor != "doc":
        return _empty_result(
            f"这是老版 .{flavor} 文件（OLE2 二进制格式），本解析器只支持 Word / PDF / 文本文件"
        )

    if not is_available():
        return _empty_result(_OLE2_DOC_ERROR)

    docx_bytes = convert_to_docx(content)
    if not docx_bytes:
        return _empty_result(
            "老版 .doc 自动转换失败（Word 报错或超时），请在 Word 里「另存为 .docx」后重新上传"
        )

    converted = _parse_docx(docx_bytes)
    if not converted["blocks"] and not converted["images"]:
        return _empty_result(
            "老版 .doc 转成 docx 后没有解析出可用的文字或图片，请确认文件内容是否正常"
        )
    return converted


# ============================================================
# 把解析结果渲染成「带图片标记的正文」
# ============================================================
#: 图片标记模板（必须与 `tools/question_manage._IMAGE_MARKER` 一致）
#:
#: ⚠️ 这是**数据标记**，不是提示词文案 —— 即使系统提示词全英文，它也保持中文：
#: 模型把它原样写进题干 → 工具用 `【图(\d+)】` 正则替换成 `<img>` → 缺一不可。
#: 别为了「全英文」把它改成 `[IMAGE_N]`，那会让整条图片链路失效。
IMAGE_MARKER_TMPL = "【图{}】"


def build_marked_document(parsed: dict, offset: int = 0):
    """把解析结果渲染成**带图片标记**的正文

    为什么要标记：文档里的「如下图 / 见下表」必须能定位到具体是哪张图。
    标记编号与返回的图片列表下标一一对应（`offset` 用于给「用户另外上传的图片」让位）。
    渲染结果与 `blocks` 的先后顺序完全一致 —— 图片就落在它原本所在的文字位置之间。

    Args:
        parsed: `parse_file()` 的返回值
        offset: 图片编号起始值（用户另外上传了图片时，文档插图要顺延编号）

    Returns:
        `(正文, 图片 data URL 列表 —— 顺序即 【图N】 的 N)`
    """
    blocks = parsed.get("blocks") or []
    image_urls: list = []
    parts: list = []

    if blocks:
        for block in blocks:
            if not isinstance(block, dict):
                continue
            if block.get("type") == "image" and block.get("imageUrl"):
                parts.append(IMAGE_MARKER_TMPL.format(offset + len(image_urls)))
                image_urls.append(block["imageUrl"])
            else:
                content = (block.get("content") or "").strip()
                if content:
                    parts.append(content)
    else:
        # 没有块信息（老数据 / 解析器降级）：正文原样，图片统一附在末尾
        text = (parsed.get("textContent") or "").strip()
        if text:
            parts.append(text)
        for image in parsed.get("images") or []:
            if isinstance(image, dict) and image.get("imageUrl"):
                image_urls.append(image["imageUrl"])
        if image_urls:
            parts.append(
                "".join(
                    IMAGE_MARKER_TMPL.format(offset + i) for i in range(len(image_urls))
                )
            )

    return "\n".join(parts), image_urls


def _parse_txt(content: bytes) -> dict:
    """解析纯文本文件(TXT/MD)"""
    text = content.decode("utf-8", errors="ignore")
    builder = _BlockBuilder()
    builder.add_text(text)
    return builder.result()


def _local(tag: str) -> str:
    """去掉 XML 命名空间前缀，只留标签本地名"""
    return tag.split("}", 1)[1] if "}" in tag else tag


def _iter_paragraph_parts(node):
    """按**文档顺序**产出 `('text', 文本)` / `('image', rId)`

    为什么不用 `_extract_paragraph_text`：那个把整段的文本拼在一起、图片再统一抽，
    于是「文字在前、图在后」的段落会把图排到文字前面。这里走一遍原始 XML，
    保证图片落在它真正所在的位置（对应需求里的「图片导入到对应位置」）。

    覆盖的引用形式：
      · DrawingML：`w:drawing` → … → `a:blip/@r:embed`
      · VML 老式图片（Word 97 兼容模式）：`w:pict` → `v:imagedata/@r:id`
    """
    for child in node:
        tag = _local(child.tag)
        if tag == "oMath":
            latex = _omml_to_latex(child)
            if latex:
                yield ("text", f"${latex}$")
        elif tag == "oMathPara":
            for omath in child.iter(_qn(M_NS, "oMath")):
                latex = _omml_to_latex(omath)
                if latex:
                    yield ("text", f"$${latex}$$")
        elif tag == "t":
            if child.text:
                yield ("text", child.text)
        elif tag == "tab":
            yield ("text", "\t")
        elif tag in ("br", "cr"):
            yield ("text", "\n")
        elif tag == "blip":
            rid = child.get(_qn(_NS["r"], "embed")) or child.get(_qn(_NS["r"], "link"))
            if rid:
                yield ("image", rid)
        elif tag == "imagedata":
            rid = child.get(_qn(_NS["r"], "id")) or child.get(_qn(_NS["r"], "embed"))
            if rid:
                yield ("image", rid)
        else:
            # 其余层级（w:r / w:ins / w:sdt / w:drawing / wp:inline / pic:pic …）
            # 一律递归下去 —— Word 的嵌套方式太多，穷举容器标签不划算
            yield from _iter_paragraph_parts(child)


def _collect_image_rels(part) -> dict:
    """收集 OPC part 里的图片关系：rId → (blob, content_type)"""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    blobs: dict = {}
    try:
        for rel in part.rels.values():
            if rel.reltype != RT.IMAGE:
                continue
            try:
                target = rel.target_part
                blobs[rel.rId] = (
                    target.blob,
                    getattr(target, "content_type", "") or "image/png",
                )
            except Exception:  # noqa: BLE001 - 坏的关系项跳过即可
                continue
    except Exception:  # noqa: BLE001
        pass
    return blobs


def _emit_paragraph(builder: _BlockBuilder, p_elem, image_blobs: dict) -> None:
    """输出一个段落：文本与插图**按出现顺序**交替落块"""
    buffer: list = []

    def flush():
        if buffer:
            builder.add_text("".join(buffer))
            buffer.clear()

    for kind, value in _iter_paragraph_parts(p_elem):
        if kind == "text":
            buffer.append(value)
            continue
        flush()
        blobs = image_blobs.get(value)
        if blobs:
            builder.add_image(*blobs)
    flush()


def _md_cell(text: str) -> str:
    """把单元格文本规范成安全的 Markdown 表格单元格

    - 折叠换行/多余空白为单个空格；
    - 转义单元格里的 `|`（否则会被当成列分隔符，把表格切错）。
      Markdown 规范里单元格内的竖线要用 `\\|` 表示。
    """
    s = (text or "").replace("\r", " ").replace("\n", " ").strip()
    s = " ".join(s.split())
    return s.replace("|", "\\|")


def _emit_table(builder: _BlockBuilder, tbl_elem, image_blobs: dict) -> None:
    """输出一个表格为**标准 Markdown 表格**（用户 2026-10-11 方案 A），并紧跟各行的插图

    为什么改：原先每行只输出 `a | b | c` 纯文本行 —— **没有表头分隔行**，不是合法 Markdown
    表格，前端 `marked` 认不出 → 用户看到的就是一串竖线文本（Word 里明明是整齐的表格）。
    现在输出：
        | X\\Y | 0 | 1 | 2 |
        | --- | --- | --- | --- |
        | 0   | 0.1 | 0.3 | 0.2 |
    前端 `renderMarkdownWithLatex` 会把它渲染成真表格。
    （以前只取 `cell.text` 会丢掉表格里的图，故每行仍按顺序补发该行单元格内的插图。）
    """
    table_rows: list = []  # 每项 = {"cells": [str,...], "images": [rId,...]}
    for row in tbl_elem.findall(_qn(W_NS, "tr")):
        row_texts: list = []
        row_images: list = []
        for cell in row.findall(_qn(W_NS, "tc")):
            cell_texts: list = []
            for child in cell:
                tag = _local(child.tag)
                if tag == "p":
                    for kind, value in _iter_paragraph_parts(child):
                        if kind == "text":
                            cell_texts.append(value)
                        else:
                            row_images.append(value)
                elif tag == "tbl":
                    # 嵌套表格：它的行列由外层 findall 管不到，直接抽其中的文本与图
                    for kind, value in _iter_paragraph_parts(child):
                        if kind == "text":
                            cell_texts.append(value)
                        else:
                            row_images.append(value)
            row_texts.append(_md_cell("".join(cell_texts)))
        if any(c for c in row_texts) or row_images:
            table_rows.append({"cells": row_texts, "images": row_images})

    if not table_rows:
        return

    # 列数取各行最大值（Markdown 表格要求表头与分隔行列数一致）
    ncol = max(len(r["cells"]) for r in table_rows)
    lines: list = []
    for i, r in enumerate(table_rows):
        cells = list(r["cells"]) + [""] * (ncol - len(r["cells"]))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            # 表头分隔行：让 Markdown 识别为表格（首行会被当成表头）
            lines.append("| " + " | ".join(["---"] * ncol) + " |")

    markdown = "\n".join(lines)

    # ── 2026-10-26 v3：除 Markdown 外，**再产出一个结构化 table 块** ──
    #
    # 只给 Markdown 是不够的：导出 Word/PDF 时它会被当普通文字，竖线原样印出
    # （实测题库 432 第 33 题，`| X | 0 | 1 | 2 | 3 |` 整行印在 PDF 上）。
    # table 块让导出层能渲染成**原生表格**，同时 markdown 仍进 textContent，
    # 老调用方（知识库入库 / 前端 marked）行为完全不变。
    headers = [c for c in table_rows[0]["cells"]]
    body_rows = [[c for c in r["cells"]] for r in table_rows[1:]]
    builder.add_table(headers, body_rows, markdown=markdown)

    # 表格各行的插图仍按原顺序补发（放在表格块之后）
    for r in table_rows:
        for rid in r["images"]:
            blobs = image_blobs.get(rid)
            if blobs:
                builder.add_image(*blobs)


def _walk_docx_body(builder: _BlockBuilder, node, image_blobs: dict) -> None:
    """按文档顺序遍历 body（含 `w:sdt` 内容控件、嵌套容器）"""
    for child in node:
        tag = _local(child.tag)
        if tag == "p":
            _emit_paragraph(builder, child, image_blobs)
        elif tag == "tbl":
            _emit_table(builder, child, image_blobs)
        elif tag in ("sdt", "sdtContent", "body"):
            _walk_docx_body(builder, child, image_blobs)


def _parse_docx(content: bytes) -> dict:
    """解析 DOCX：段落（含 OMML 公式）+ 正文插图 + 表格（含单元格内插图）"""
    from docx import Document

    doc = Document(io.BytesIO(content))
    image_blobs = _collect_image_rels(doc.part)

    builder = _BlockBuilder()
    _walk_docx_body(builder, doc.element.body, image_blobs)
    return builder.result()


def _parse_pdf(content: bytes) -> dict:
    """解析 PDF：逐页「**先该页文本、后该页插图**」

    插图用 pypdf 的 `page.images` 抽（内部借 Pillow 解码，不需要额外依赖）。

    ⚠️ 位置只能精确到「页」：PDF 是排版指令流，没有段落结构，
    要精确到「行间」得解析内容流的坐标 —— sp3 也是按页近似的，这里对齐。
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    builder = _BlockBuilder()

    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - 单页取文本失败不影响整份文件
            page_text = ""
        builder.add_text(page_text)
        _collect_pdf_page_images(builder, page)

    return builder.result()


def _collect_pdf_page_images(builder: _BlockBuilder, page) -> None:
    """把一页里的插图逐张喂给 builder —— **单张失败只跳过那一张**"""
    try:
        total = len(page.images)
    except Exception:  # noqa: BLE001 - 整页取不出就当作没有图
        return

    for index in range(total):
        try:
            image_file = page.images[index]
        except Exception:  # noqa: BLE001
            continue

        prepared = None
        try:
            pil = image_file.image  # pypdf 已解码好的 Pillow 图（最稳，能顺手过滤小图标）
        except Exception:  # noqa: BLE001
            pil = None
        if pil is not None:
            prepared = _prepare_pil_image(pil)
        if prepared is None:
            try:
                prepared = _prepare_doc_image(image_file.data, "")
            except Exception:  # noqa: BLE001
                prepared = None
        if prepared:
            builder.add_payload(*prepared)

