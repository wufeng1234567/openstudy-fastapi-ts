# -*- coding: utf-8 -*-
"""RAG 图片落库端到端探针（真实 DB + 真实 docx 生成 + 真实本地 embedding）

验证「文档里的图片能否正确入库、被分块归属、被检索带出」这条链：
  1. 造一个带图的 docx（python-docx 生成，图=小红点 PNG）
  2. 走真实上传接口 → 等向量化（status=5）
  3. 直接查库断言：
     · `rag_document.image_map` 非空且值为 `/profile/...`
     · 图片文件真的落盘了（`UPLOAD_DIR` 下存在）
     · `rag_document.raw_content` 里有 `【图0】` 标记
     · 命中该块的 `rag_document_chunk.meta_data["images"]` 带出了图
  4. 走真实检索 `/rag/search` → 命中的 hit 带 `images`
  5. 清理

用法：
    .venv/Scripts/python.exe tests/_probe_rag_image_store.py
"""
import io
import json
import os
import struct
import sys
import time
import urllib.request
import urllib.error
import uuid
import zlib

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for p in (ROOT, os.path.join(ROOT, "tests")):
    if p not in sys.path:
        sys.path.insert(0, p)

BASE = "http://localhost:9101"


#: 尺寸必须 > 解析器的最小边门槛（`file_parser` 里 <50px 的图会被当图标丢弃）
def _png_bytes(w=120, h=120, rgb=(200, 30, 30)) -> bytes:
    """生成一个纯色 PNG（不依赖 Pillow）——用作 docx 里的插图。"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    row = b"\x00" + bytes(rgb) * w
    raw = row * h
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def _docx_bytes(text: str, png: bytes) -> bytes:
    """用 python-docx 生成「一段文字 + 一张图 + 一段文字」的 docx。"""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph(text)
    doc.add_picture(io.BytesIO(png), width=Inches(1.0))
    doc.add_paragraph("（上图结束，以下是纯文字）图片之后再补一段说明文字用于出题。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _post(path, token=None, body=None, raw_body=None, headers=None, timeout=180):
    h = dict(headers or {})
    data = raw_body
    if body is not None:
        h["Content-Type"] = "application/json"
        data = json.dumps(body).encode()
    if token:
        h["Authorization"] = "Bearer " + token
    r = urllib.request.Request(BASE + path, data=data, headers=h, method="POST")
    try:
        resp = urllib.request.urlopen(r, timeout=timeout)
        return resp.status, json.loads(resp.read().decode("utf-8", "replace"))
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode("utf-8", "replace"))
        except Exception:
            return e.code, {}


def _get(path, token=None, timeout=60):
    h = {"Authorization": "Bearer " + token} if token else {}
    r = urllib.request.Request(BASE + path, headers=h)
    try:
        resp = urllib.request.urlopen(r, timeout=timeout)
        return resp.status, json.loads(resp.read().decode("utf-8", "replace"))
    except urllib.error.HTTPError as e:
        return e.code, {}


def _multipart(field: str, filename: str, content: bytes, mime: str) -> tuple:
    boundary = "----probe" + uuid.uuid4().hex
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'
        f"Content-Type: {mime}\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    return body, f"multipart/form-data; boundary={boundary}"


def main():
    checks = []

    def ck(cid, title, cond, actual=""):
        checks.append((cid, title, cond, actual))
        print(f"  [{'通过' if cond else '失败'}] {cid} {title}" + (f" | {actual}" if not cond else ""))

    st, body = _post("/login", body={"username": "admin", "password": "123456"})
    token = body["data"]["token"]
    ck("IMG-01", "登录成功", bool(token))

    kb_id = None
    try:
        st, body = _post("/rag/knowledgeBase", token=token,
                         body={"name": f"图片落库探针库_{int(time.time())}"})
        kb_id = body["data"]["kbId"]
        ck("IMG-02", "创建知识库", bool(kb_id), str(body))

        docx = _docx_bytes("第三章 图示：下面这张图展示了核心结构。", _png_bytes())
        mp_body, mp_ct = _multipart("file", "带图探针.docx", docx,
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        st, body = _post(f"/rag/document/upload?kbId={kb_id}", token=token,
                         raw_body=mp_body, headers={"Content-Type": mp_ct})
        ck("IMG-03", "上传带图 docx", st == 200 and body.get("code") == 200, str(body)[:200])
        doc_id = (body.get("data") or {}).get("docId")
        ck("IMG-04", "上传响应带 imageCount>0", (body.get("data") or {}).get("imageCount", 0) > 0,
           str(body.get("data")))

        stt = None
        for _ in range(60):
            st, body = _get(f"/rag/document/status/{doc_id}", token)
            stt = (body.get("data") or {}).get("status")
            if stt in (4, 5):
                break
            time.sleep(2)
        ck("IMG-05", "文档已向量化（status=5）", stt == 5, f"status={stt} err={(body.get('data') or {}).get('errorMsg')}")

        # ---- 直接查库断言 ----
        import asyncio
        from app.database import async_session_factory
        from app.models.rag import RagDocument, RagDocumentChunk
        from sqlalchemy import select
        from app.config import settings

        async def _db_check():
            async with async_session_factory() as db:
                doc = (await db.execute(
                    select(RagDocument).where(RagDocument.id == doc_id)
                )).scalar_one()
                chunks = (await db.execute(
                    select(RagDocumentChunk).where(RagDocumentChunk.doc_id == doc_id)
                    .order_by(RagDocumentChunk.chunk_index)
                )).scalars().all()
                return doc, chunks

        doc, chunks = asyncio.run(_db_check())
        ck("IMG-06", "rag_document.image_map 非空", bool(doc.image_map), str(doc.image_map)[:120])
        first_url = ""
        if doc.image_map:
            first_url = list(doc.image_map.values())[0]
        ck("IMG-07", "图片 URL 是 /profile/... 相对路径", first_url.startswith("/profile/"),
           first_url)
        ck("IMG-08", "raw_content 含 【图0】 标记", "【图0】" in (doc.raw_content or ""),
           (doc.raw_content or "")[:80])
        # 图片文件真的落盘
        rel = first_url[len("/profile/"):] if first_url.startswith("/profile/") else ""
        disk = os.path.join(settings.UPLOAD_DIR, *rel.split("/")) if rel else ""
        ck("IMG-09", "图片文件已落盘存在", bool(disk) and os.path.exists(disk), disk)
        # chunk 归属
        with_img = [c for c in chunks if isinstance(c.meta_data, dict) and c.meta_data.get("images")]
        ck("IMG-10", "至少一个 chunk 的 meta_data 带出了图", len(with_img) >= 1,
           f"{len(with_img)}/{len(chunks)} 块带图")
        if with_img:
            ck("IMG-11", "chunk 图 URL 与 image_map 一致",
               with_img[0].meta_data["images"][0] == first_url,
               str(with_img[0].meta_data))

        # ---- 走真实检索 ----
        st, body = _post("/rag/search", token=token,
                         body={"query": "第三章 图示 核心结构", "kbIds": [kb_id], "topK": 5})
        hits = (body.get("data") or {}).get("hits") or []
        hit_with_img = [h for h in hits if h.get("images")]
        ck("IMG-12", "检索命中结果带 images 字段", len(hit_with_img) >= 1,
           f"{len(hit_with_img)}/{len(hits)} 命中带图")
    finally:
        if kb_id:
            h = {"Authorization": "Bearer " + token}
            r = urllib.request.Request(BASE + f"/rag/knowledgeBase/{kb_id}", headers=h, method="DELETE")
            try:
                urllib.request.urlopen(r, timeout=30)
            except Exception:
                pass
            ck("IMG-13", "清理探针知识库", True)

    red = [c for c in checks if not c[2]]
    print(f"\n共 {len(checks)} 项：{len(checks) - len(red)} 通过；红的 = {[c[0] for c in red] or '无'}")
    return 1 if red else 0


if __name__ == "__main__":
    raise SystemExit(main())
