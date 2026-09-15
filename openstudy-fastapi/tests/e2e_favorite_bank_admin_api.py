# -*- coding: utf-8 -*-
"""后台管理页面「题库收藏」模块 E2E 接口测试（真实 HTTP 调用）

被测页面：openstudy-vue3-ts/src/views/favoriteBank/favoriteBank/index.vue
被测接口：app/api/v1/questions.py 的 /favoriteBank/favoriteBank/*
关联链路：bankShare 分享领取 → 写入题库收藏（source=share / canEdit=False / shareStatus）
账号角色：admin(1) 管理员 / user1(105) 普通用户 / user2(106) 普通用户
用例编号：P*（P=题库收藏）

判定口径（与同项目「斩题管理/错题管理」及 SpringBoot 参考版一致）：
  - list 支持 userId/bankId/bankName/notes/tags/isStarred/favoriteStatus/studyCount/lastStudyTime + 分页
  - 管理员不传 userId → 可看全部用户；普通用户 → 只能看自己
  - 新增：管理员可代任意用户录入（payload.userId），并落库 sortOrder/isStarred/favoriteStatus/
    lastStudyTime/studyCount 等全字段；重复收藏 400；题库/用户不存在要有明确提示
  - 详情/修改/删除：管理员可操作任意用户记录，普通用户越权必须 403
  - 页面「导出」按钮对应的 POST /favoriteBank/favoriteBank/export 必须存在且支持筛选
  - 分享领取产生的收藏：source='share'、canEdit=False；分享撤销后 shareStatus='invalid'

用法（后端需先跑起来，默认 127.0.0.1:9101）：
    cd openstudy-fastapi
    .venv/Scripts/python.exe tests/e2e_favorite_bank_admin_api.py

产出：脚本同目录 e2e_favorite_bank_admin_api_result.json
说明：全部测试数据建在临时题库上，结束时删除题库（级联清理收藏/分享），不留残留。
"""
import json
import os
import sys
import time
import urllib.error
import urllib.request
from urllib.parse import urlencode

BASE = "http://127.0.0.1:9101"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "e2e_favorite_bank_admin_api_result.json")

ROWS = []


def call(method, path, token=None, body=None, params=None, timeout=25):
    url = BASE + path
    if params:
        clean = {k: v for k, v in params.items() if v is not None}
        if clean:
            url += "?" + urlencode(clean, doseq=True)
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read()
            ct = r.headers.get("Content-Type", "")
            if "json" in ct:
                try:
                    return r.status, json.loads(raw.decode("utf-8"))
                except Exception:
                    return r.status, raw
            return r.status, raw
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode("utf-8"))
        except Exception:
            return e.code, raw


def login(account):
    st, res = call("POST", "/quickLogin", body={"account": account})
    return res["data"]["token"]


def T(cid, module, title, pre, steps, expected, ok, actual):
    ROWS.append({
        "cid": cid, "module": module, "title": title, "pre": pre,
        "steps": steps, "expected": expected, "actual": actual,
        "status": "通过" if ok else "失败",
    })
    print(("  [PASS] " if ok else "  [FAIL] ") + cid + " " + title + " -> " + str(actual))


def code_of(res):
    return res.get("code") if isinstance(res, dict) else "-"


def raw_data(res):
    return res.get("data") if isinstance(res, dict) else None


def dd(res):
    d = raw_data(res)
    return d if isinstance(d, dict) else {}


def dl(res):
    d = raw_data(res)
    return d if isinstance(d, list) else []


def rows_of(res):
    return (res.get("rows") or []) if isinstance(res, dict) else []


def is_file(res):
    return isinstance(res, (bytes, bytearray)) and len(res) > 0


def is_xlsx(res):
    return is_file(res) and res[:2] == b"PK"


def desc(res, st=200):
    if is_file(res):
        return "HTTP %s 文件流 大小=%d 字节" % (st, len(res))
    msg = (res.get("msg") or "")[:80] if isinstance(res, dict) else ""
    return "HTTP %s code=%s %s" % (st, code_of(res), msg)


# ---------------------------------------------------------------- DB 直连
def _db(sql, params=None, write=False):
    try:
        import asyncio

        from sqlalchemy import text

        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        if root not in sys.path:
            sys.path.insert(0, root)

        from app.config import settings
        from sqlalchemy.ext.asyncio import create_async_engine
        from sqlalchemy.pool import NullPool

        async def _run():
            engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
            try:
                if write:
                    async with engine.begin() as conn:
                        return (await conn.execute(text(sql), params or {})).rowcount
                async with engine.connect() as conn:
                    return (await conn.execute(text(sql), params or {})).all()
            finally:
                await engine.dispose()

        return asyncio.run(_run())
    except Exception as exc:  # noqa: BLE001
        print("    (DB %s 不可用：%s)" % ("写入" if write else "直读", exc))
        return None


def db_query(sql, params=None):
    return _db(sql, params, write=False)


def db_exec(sql, params=None):
    return _db(sql, params, write=True)


def db_one(sql, params=None):
    rows = db_query(sql, params)
    if not rows:
        return None
    return dict(rows[0]._mapping)


def new_bank(tok, name):
    st, res = call("POST", "/questionBank/questionBank", tok, {"bankName": name})
    d = dd(res)
    return d.get("id"), d.get("userId")


def sweep_leftovers(created):
    bank_ids = created.get("bank_ids") or []
    if not bank_ids:
        return
    bank_t = tuple(bank_ids) or (0,)
    db_exec("DELETE FROM user_bank_favorite WHERE bank_id IN :i", {"i": bank_t})
    db_exec("DELETE FROM bank_share_receive WHERE bank_id IN :i", {"i": bank_t})
    db_exec("DELETE FROM bank_share WHERE bank_id IN :i", {"i": bank_t})
    db_exec("DELETE FROM question_sub WHERE main_id IN (SELECT id FROM question_main WHERE bank_id IN :b)",
            {"b": bank_t})
    db_exec("DELETE FROM question_main WHERE bank_id IN :b", {"b": bank_t})
    db_exec("DELETE FROM question_bank WHERE id IN :i", {"i": bank_t})
    db_exec("DELETE FROM bank_share_receive WHERE share_id NOT IN (SELECT id FROM bank_share)")
    db_exec("DELETE FROM user_notification WHERE link LIKE '/front/myQuestion/bankCollect%'")
    db_exec("DELETE FROM user_notification WHERE content LIKE '%题库收藏%'")
    print("  [自检] 已清理上次运行遗留：bank=%s" % bank_ids)


def main():
    if os.path.exists(OUT):
        try:
            with open(OUT, encoding="utf-8") as f:
                sweep_leftovers(json.load(f).get("cleanup") or {})
        except Exception as exc:  # noqa: BLE001
            print("  [自检] 读取上次结果失败：%s" % exc)

    P = "题库收藏"
    BANK_A = P + "A-测试题库"
    BANK_B = P + "B-分享题库"
    BANK_C = P + "C-用户2题库"

    print("=" * 78)
    print("后台管理「题库收藏」模块 E2E 测试   前缀 =", P)
    print("=" * 78)

    ta = login("admin")
    t1 = login("user1")
    t2 = login("user2")

    bank_a, uid1 = new_bank(t1, BANK_A)
    bank_b, _ = new_bank(t1, BANK_B)
    bank_c, uid2 = new_bank(t2, BANK_C)
    probe, uid_admin = new_bank(ta, P + "管理员探针")
    call("DELETE", "/questionBank/questionBank/%s" % probe, ta)
    print("bankA=%s bankB=%s bankC=%s uid_admin=%s uid1=%s uid2=%s"
          % (bank_a, bank_b, bank_c, uid_admin, uid1, uid2))

    # ---- 基线收藏数据（走接口，保证与业务一致）----
    call("POST", "/favoriteBank/favoriteBank", t1, {"bankId": bank_a})
    call("POST", "/favoriteBank/favoriteBank", t1, {"bankId": bank_b})
    call("POST", "/favoriteBank/favoriteBank", t2, {"bankId": bank_c})

    # ---- 用 DB 造「有区分度」的属性，供筛选/回显用例使用（不依赖待测写接口）----
    db_exec("""UPDATE user_bank_favorite SET notes='收藏备注A-重点整理', tags='重点,高频',
               sort_order=5, is_starred=1, study_count=3, last_study_time='2026-09-01 09:00:00'
               WHERE user_id=:u AND bank_id=:b""", {"u": uid1, "b": bank_a})
    db_exec("""UPDATE user_bank_favorite SET notes='收藏备注B', tags='复习', sort_order=0,
               is_starred=0, study_count=0, last_study_time=NULL
               WHERE user_id=:u AND bank_id=:b""", {"u": uid1, "b": bank_b})
    db_exec("""UPDATE user_bank_favorite SET notes='用户2自己的收藏', tags='重点', sort_order=2,
               is_starred=1, study_count=7, last_study_time='2026-09-03 10:00:00'
               WHERE user_id=:u AND bank_id=:b""", {"u": uid2, "b": bank_c})

    def cnt(sql, params):
        return (db_one(sql, params) or {}).get("c")

    exp_all = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE favorite_status=1", {})
    exp_u1 = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE user_id=:u AND favorite_status=1", {"u": uid1})
    exp_u2 = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE user_id=:u AND favorite_status=1", {"u": uid2})
    exp_star = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE is_starred=1 AND favorite_status=1", {})
    exp_sc3 = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE study_count=3 AND favorite_status=1", {})
    print("基线：全部活跃=%s user1=%s user2=%s 已标星=%s studyCount=3=%s"
          % (exp_all, exp_u1, exp_u2, exp_star, exp_sc3))

    M = "题库收藏(后台)"

    # ==================== P1 列表查询与筛选（对应后台列表页 9 个查询条件） ====================
    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta, params={"pageNum": 1, "pageSize": 10})
    T("P01", M, "列表-管理员查询返回分页结构",
      "admin 已登录", "GET /favoriteBank/favoriteBank/list?pageNum=1&pageSize=10",
      "code=200；顶层含 total 与 rows",
      code_of(res) == 200 and "total" in res and "rows" in res,
      "code=%s total=%s rows=%d" % (code_of(res), res.get("total"), len(rows_of(res))))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"userId": uid1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P02", M, "列表-按用户ID筛选",
      "user1 有 %s 条活跃收藏" % exp_u1, "admin 请求 list?userId=%s" % uid1,
      "total=%s 且 rows 的 userId 全为 %s" % (exp_u1, uid1),
      code_of(res) == 200 and res.get("total") == exp_u1 and len(rows) > 0
      and all(r.get("userId") == uid1 for r in rows),
      "total=%s userId集合=%s" % (res.get("total"), sorted({r.get("userId") for r in rows})[:5]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"bankId": bank_a, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P03", M, "列表-按题库ID筛选",
      "bankA=%s 被 user1 收藏" % bank_a, "admin 请求 list?bankId=%s" % bank_a,
      "rows 全部 bankId=bankA 且非空",
      code_of(res) == 200 and len(rows) > 0 and all(r.get("bankId") == bank_a for r in rows),
      "命中 %d 条 bankId集合=%s" % (len(rows), sorted({r.get("bankId") for r in rows})[:5]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"bankName": "收藏A", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P04", M, "列表-按题库名称模糊筛选",
      "题库《%s》" % BANK_A, "admin 请求 list?bankName=收藏A",
      "命中且 rows 的 bankName 均含「收藏A」",
      code_of(res) == 200 and len(rows) > 0 and all("收藏A" in (r.get("bankName") or "") for r in rows),
      "命中 %d 条 bankName=%s" % (len(rows), [r.get("bankName") for r in rows][:3]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"notes": "备注A", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P05", M, "列表-按收藏备注模糊筛选",
      "存在备注含「备注A」的收藏", "admin 请求 list?notes=备注A",
      "命中且 rows 的 notes 均含关键字",
      code_of(res) == 200 and len(rows) > 0 and all("备注A" in (r.get("notes") or "") for r in rows),
      "命中 %d 条 notes=%s" % (len(rows), [r.get("notes") for r in rows][:3]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"tags": "高频", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P06", M, "列表-按标签模糊筛选",
      "存在标签含「高频」的收藏", "admin 请求 list?tags=高频",
      "命中且 rows 的 tags 均含关键字",
      code_of(res) == 200 and len(rows) > 0 and all("高频" in (r.get("tags") or "") for r in rows),
      "命中 %d 条 tags=%s" % (len(rows), [r.get("tags") for r in rows][:3]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"isStarred": 1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P07", M, "列表-按是否标星筛选",
      "全库已标星 %s 条" % exp_star, "admin 请求 list?isStarred=1",
      "total=%s 且 rows 的 isStarred 全为 1" % exp_star,
      code_of(res) == 200 and res.get("total") == exp_star and len(rows) > 0
      and all(r.get("isStarred") == 1 for r in rows),
      "total=%s isStarred集合=%s" % (res.get("total"), sorted({r.get("isStarred") for r in rows})[:5]))

    # 造一条「失效」收藏，验证 favoriteStatus 语义（默认只查有效）
    db_exec("UPDATE user_bank_favorite SET favorite_status=0 WHERE user_id=:u AND bank_id=:b",
            {"u": uid1, "b": bank_b})
    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"pageNum": 1, "pageSize": 100})
    default_ids = {r.get("favoriteId") for r in rows_of(res)}
    st2, res2 = call("GET", "/favoriteBank/favoriteBank/list", ta,
                     params={"favoriteStatus": 0, "pageNum": 1, "pageSize": 100})
    off_rows = rows_of(res2)
    off_id = db_one("SELECT favorite_id FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                    {"u": uid1, "b": bank_b})
    off_id = (off_id or {}).get("favorite_id")
    T("P08", M, "列表-按收藏状态筛选（默认只返回有效）",
      "user1 对 bankB 的收藏已置为 favorite_status=0",
      "① list 不传 favoriteStatus ② list?favoriteStatus=0",
      "① 默认结果不含失效记录 ② favoriteStatus=0 精确命中该条",
      code_of(res) == 200 and code_of(res2) == 200 and off_id not in default_ids
      and len(off_rows) == 1 and off_rows[0].get("favoriteId") == off_id
      and off_rows[0].get("favoriteStatus") == 0,
      "默认含失效=%s；favoriteStatus=0 命中 %d 条" % (off_id in default_ids, len(off_rows)))
    db_exec("UPDATE user_bank_favorite SET favorite_status=1 WHERE favorite_id=:f", {"f": off_id})

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"studyCount": 3, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P09", M, "列表-按学习次数筛选",
      "全库 studyCount=3 共 %s 条" % exp_sc3, "admin 请求 list?studyCount=3",
      "total=%s 且 rows 的 studyCount 全为 3" % exp_sc3,
      code_of(res) == 200 and res.get("total") == exp_sc3 and len(rows) > 0
      and all(r.get("studyCount") == 3 for r in rows),
      "total=%s studyCount集合=%s" % (res.get("total"), sorted({r.get("studyCount") for r in rows})[:5]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"lastStudyTime": "2026-09-01", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P10", M, "列表-按最后学习时间筛选（当天）",
      "存在 last_study_time=2026-09-01 的收藏", "admin 请求 list?lastStudyTime=2026-09-01",
      "命中且 rows 的 lastStudyTime 落在 2026-09-01",
      code_of(res) == 200 and len(rows) > 0
      and all((r.get("lastStudyTime") or "").startswith("2026-09-01") for r in rows),
      "命中 %d 条 %s" % (len(rows), [r.get("lastStudyTime") for r in rows][:3]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"pageNum": 1, "pageSize": 2})
    T("P11", M, "列表-分页返回条数受 pageSize 约束",
      "全库活跃收藏 > 2 条", "admin 请求 list?pageSize=2",
      "rows<=2 且 total 为全量",
      code_of(res) == 200 and len(rows_of(res)) <= 2 and (res.get("total") or 0) > 2,
      "total=%s rows=%d" % (res.get("total"), len(rows_of(res))))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"pageNum": 1, "pageSize": 100})
    owners = sorted({r.get("userId") for r in rows_of(res)})
    T("P12", M, "列表-管理员不传 userId 时可看到全部用户",
      "admin/user1/user2 均有收藏", "admin 请求 list（不传 userId）",
      "rows 中出现 >=2 个不同 userId",
      code_of(res) == 200 and len(owners) >= 2, "total=%s 涉及用户=%s" % (res.get("total"), owners[:6]))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"userId": uid1, "bankId": bank_a, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P13", M, "列表-组合条件筛选（userId + bankId）",
      "user1 对 bankA 有 1 条收藏", "admin 请求 list?userId=%s&bankId=%s" % (uid1, bank_a),
      "命中 1 条且两者都匹配",
      code_of(res) == 200 and len(rows) == 1
      and rows[0].get("userId") == uid1 and rows[0].get("bankId") == bank_a,
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", t2,
                   params={"pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P14", M, "列表-普通用户只能看到自己的收藏",
      "user2 有 %s 条活跃收藏" % exp_u2, "user2 请求 list?pageSize=100",
      "rows 的 userId 全为 %s" % uid2,
      code_of(res) == 200 and len(rows) > 0 and all(r.get("userId") == uid2 for r in rows),
      "total=%s userId集合=%s" % (res.get("total"), sorted({r.get("userId") for r in rows})[:5]))

    # ==================== P2 新增（后台对话框全字段） ====================
    add_body = {
        "userId": uid1, "bankId": bank_c,
        "notes": P + "-后台新增备注", "tags": "必考,重点", "sortOrder": 9,
        "isStarred": 1, "favoriteStatus": 1,
        "lastStudyTime": "2026-09-05 08:00:00", "studyCount": 4,
    }
    st, res = call("POST", "/favoriteBank/favoriteBank", ta, add_body)
    T("P15", M, "新增-管理员代 user1 收藏题库（提交后台表单全字段）",
      "bankC=%s 未被 user1 收藏" % bank_c,
      "admin POST /favoriteBank/favoriteBank  body 含 userId=%s 与全部表单字段" % uid1,
      "code=200", code_of(res) == 200, desc(res, st))

    row = db_one("""SELECT user_id, bank_id, notes, tags, sort_order, is_starred,
                           favorite_status, last_study_time, study_count
                    FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b""",
                 {"u": uid1, "b": bank_c})
    T("P16", M, "新增-表单字段全部落库（归属=被代录用户）",
      "P15 已提交全字段",
      "直读 user_bank_favorite 逐字段比对",
      "user_id=%s；备注/标签一致；sortOrder=9 标星=1 状态=1；lastStudyTime 非空；studyCount=4" % uid1,
      bool(row) and row.get("user_id") == uid1 and row.get("bank_id") == bank_c
      and row.get("notes") == P + "-后台新增备注" and row.get("tags") == "必考,重点"
      and row.get("sort_order") == 9 and row.get("is_starred") == 1
      and row.get("favorite_status") == 1 and row.get("last_study_time") is not None
      and row.get("study_count") == 4,
      "库中=%s" % row)

    fid_new = (db_one("SELECT favorite_id FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                      {"u": uid1, "b": bank_c}) or {}).get("favorite_id")
    st, res = call("GET", "/favoriteBank/favoriteBank/%s" % fid_new, ta)
    d = dd(res)
    need = ["favoriteId", "userId", "bankId", "bankName", "notes", "tags", "sortOrder",
            "isStarred", "lastStudyTime", "studyCount", "favoriteStatus", "createTime", "updateTime"]
    missing = [f for f in need if f not in d]
    T("P17", M, "详情-后台修改对话框所需字段完整回显",
      "favoriteId=%s" % fid_new, "admin GET /favoriteBank/favoriteBank/{favoriteId}",
      "返回体包含后台表格/表单的 13 个字段（缺一即页面显示空白）",
      code_of(res) == 200 and not missing,
      "code=%s 缺失字段=%s" % (code_of(res), missing or "无"))

    st, res = call("POST", "/favoriteBank/favoriteBank", ta, add_body)
    T("P18", M, "新增-同一用户重复收藏同一题库给出业务提示",
      "user1 已收藏 bankC", "admin 再次 POST 同一 body",
      "code=400 且提示「已收藏」，不应 500",
      code_of(res) == 400, "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("POST", "/favoriteBank/favoriteBank", ta,
                   {"userId": uid1, "bankId": 99999999})
    T("P19", M, "新增-题库不存在时的异常处理",
      "题库 99999999 不存在", "admin POST body={bankId:99999999}",
      "code!=200 且提示「题库不存在」",
      code_of(res) != 200 and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("POST", "/favoriteBank/favoriteBank", ta,
                   {"userId": 99999999, "bankId": bank_a})
    T("P20", M, "新增-用户不存在时的异常处理",
      "userId=99999999 不存在", "admin POST body={userId:99999999}",
      "code!=200 且提示「用户不存在」",
      code_of(res) != 200 and "用户" in (res.get("msg") or "") and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    # ==================== P3 修改 ====================
    fid_a = (db_one("SELECT favorite_id FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                    {"u": uid1, "b": bank_a}) or {}).get("favorite_id")
    upd = {
        "favoriteId": fid_a, "notes": P + "-本人修改备注", "tags": "重做",
        "sortOrder": 6, "isStarred": 0, "favoriteStatus": 1,
        "lastStudyTime": "2026-09-08 20:00:00", "studyCount": 12,
    }
    st, res = call("PUT", "/favoriteBank/favoriteBank", t1, upd)
    row_a = db_one("""SELECT notes, tags, sort_order, is_starred, favorite_status,
                             last_study_time, study_count FROM user_bank_favorite
                      WHERE favorite_id=:f""", {"f": fid_a})
    T("P21", M, "修改-本人提交后台表单全字段，字段全部更新",
      "favoriteId=%s（user1 本人记录）" % fid_a,
      "user1 PUT body 含后台表单全部字段",
      "code=200；库中备注/标签/sortOrder=6/标星=0/状态=1/最后学习时间非空/studyCount=12",
      code_of(res) == 200 and bool(row_a)
      and row_a.get("notes") == P + "-本人修改备注" and row_a.get("tags") == "重做"
      and row_a.get("sort_order") == 6 and row_a.get("is_starred") == 0
      and row_a.get("favorite_status") == 1 and row_a.get("last_study_time") is not None
      and row_a.get("study_count") == 12,
      "code=%s 库中=%s" % (code_of(res), row_a))

    fid_c = (db_one("SELECT favorite_id FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                    {"u": uid1, "b": bank_c}) or {}).get("favorite_id")
    upd_admin = {"favoriteId": fid_c, "notes": P + "-管理员修改备注", "studyCount": 20, "sortOrder": 3}
    st, res = call("PUT", "/favoriteBank/favoriteBank", ta, upd_admin)
    row_c = db_one("SELECT notes, study_count, sort_order FROM user_bank_favorite WHERE favorite_id=:f",
                   {"f": fid_c})
    T("P22", M, "修改-管理员可修改他人（user1）的收藏记录",
      "favoriteId=%s 属于 user1" % fid_c,
      "admin PUT /favoriteBank/favoriteBank  body={favoriteId:%s,...}" % fid_c,
      "code=200；库中备注/studyCount=20/sortOrder=3 均已更新",
      code_of(res) == 200 and bool(row_c)
      and row_c.get("notes") == P + "-管理员修改备注"
      and row_c.get("study_count") == 20 and row_c.get("sort_order") == 3,
      "code=%s 库中=%s" % (code_of(res), row_c))

    st, res = call("GET", "/favoriteBank/favoriteBank/%s" % fid_c, ta)
    d = dd(res)
    T("P23", M, "详情-管理员可查看他人收藏记录",
      "favoriteId=%s 属于 user1" % fid_c, "admin GET /favoriteBank/favoriteBank/{favoriteId}",
      "code=200 且 userId=%s、studyCount=20" % uid1,
      code_of(res) == 200 and d.get("userId") == uid1 and d.get("studyCount") == 20,
      "code=%s userId=%s studyCount=%s" % (code_of(res), d.get("userId"), d.get("studyCount")))

    st, res = call("PUT", "/favoriteBank/favoriteBank", ta, {"favoriteId": 99999999, "notes": "x"})
    T("P24", M, "修改-记录不存在时的异常处理",
      "favoriteId=99999999 不存在", "admin PUT body={favoriteId:99999999}",
      "code!=200 且提示「收藏记录不存在」",
      code_of(res) != 200 and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("PUT", "/favoriteBank/favoriteBank", t2,
                   {"favoriteId": fid_a, "notes": P + "-越权篡改"})
    still = db_one("SELECT notes FROM user_bank_favorite WHERE favorite_id=:f", {"f": fid_a})
    T("P25", M, "修改-普通用户不能修改他人收藏（越权拦截）",
      "favoriteId=%s 属于 user1；user2 登录" % fid_a,
      "user2 PUT body={favoriteId:%s}" % fid_a,
      "code=403；user1 的记录未被篡改",
      code_of(res) == 403 and (still or {}).get("notes") == P + "-本人修改备注",
      "code=%s msg=%s 现值=%s" % (code_of(res), res.get("msg"), (still or {}).get("notes")))

    # ==================== P4 删除 ====================
    st, res = call("DELETE", "/favoriteBank/favoriteBank/%s" % fid_c, ta)
    gone = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE favorite_id=:f", {"f": fid_c})
    T("P26", M, "删除-管理员删除他人（user1）单条收藏",
      "favoriteId=%s 属于 user1" % fid_c, "admin DELETE /favoriteBank/favoriteBank/{favoriteId}",
      "code=200；库中该记录已物理删除",
      code_of(res) == 200 and (gone or {}).get("c") == 0,
      "code=%s msg=%s 剩余=%s" % (code_of(res), res.get("msg"), (gone or {}).get("c")))

    # 造两条给批量删除
    call("POST", "/favoriteBank/favoriteBank", ta, {"userId": uid1, "bankId": bank_c})
    call("POST", "/favoriteBank/favoriteBank", ta, {"userId": uid2, "bankId": bank_c})
    ids_rows = db_query("SELECT favorite_id FROM user_bank_favorite WHERE bank_id=:b", {"b": bank_c}) or []
    id_join = ",".join(str(r[0]) for r in ids_rows)
    st, res = call("DELETE", "/favoriteBank/favoriteBank/%s" % id_join, ta)
    left = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE favorite_id IN :i",
                  {"i": tuple(r[0] for r in ids_rows)})
    T("P27", M, "删除-管理员批量删除多人收藏（逗号分隔 id）",
      "bankC 下有 %d 条收藏，ids=%s" % (len(ids_rows), id_join),
      "admin DELETE /favoriteBank/favoriteBank/{ids}",
      "code=200；这些记录均已删除",
      code_of(res) == 200 and (left or {}).get("c") == 0,
      "code=%s 剩余=%s" % (code_of(res), (left or {}).get("c")))

    st, res = call("DELETE", "/favoriteBank/favoriteBank/%s" % fid_a, t2)
    still_a = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE favorite_id=:f", {"f": fid_a})
    T("P28", M, "删除-普通用户不能删除他人收藏（越权拦截）",
      "favoriteId=%s 属于 user1；user2 登录" % fid_a,
      "user2 DELETE /favoriteBank/favoriteBank/{favoriteId}",
      "code=403；记录仍存在",
      code_of(res) == 403 and (still_a or {}).get("c") == 1,
      "code=%s msg=%s 仍存在=%s" % (code_of(res), res.get("msg"), (still_a or {}).get("c")))

    st, res = call("DELETE", "/favoriteBank/favoriteBank/99999999", ta)
    T("P29", M, "删除-不存在的 id 不应 500",
      "favoriteId=99999999 不存在", "admin DELETE /favoriteBank/favoriteBank/99999999",
      "返回业务响应（200 幂等 或 明确提示），不得 500",
      st != 500 and code_of(res) != 500, desc(res, st))

    # ==================== P5 导出（页面「导出」按钮对应的接口） ====================
    st, res = call("POST", "/favoriteBank/favoriteBank/export", ta)
    T("P30", M, "导出-后台「导出」按钮对应的接口存在并返回 Excel",
      "admin 已登录", "POST /favoriteBank/favoriteBank/export",
      "HTTP 200 且返回 xlsx 文件流（当前该路由缺失会落到 catch_all 返回 JSON）",
      st == 200 and is_xlsx(res), desc(res, st) + ("，是xlsx" if is_xlsx(res) else "，不是xlsx"))

    navail = None
    if is_xlsx(res):
        try:
            import io

            from openpyxl import load_workbook

            navail = load_workbook(io.BytesIO(res), read_only=True).active.max_row - 1
        except Exception:  # noqa: BLE001
            pass
    st, res = call("POST", "/favoriteBank/favoriteBank/export", ta, params={"userId": uid1})
    n_u1 = None
    if is_xlsx(res):
        try:
            import io

            from openpyxl import load_workbook

            n_u1 = load_workbook(io.BytesIO(res), read_only=True).active.max_row - 1
        except Exception:  # noqa: BLE001
            pass
    cur_u1 = cnt("SELECT COUNT(*) c FROM user_bank_favorite WHERE user_id=:u AND favorite_status=1", {"u": uid1})
    T("P31", M, "导出-按 userId 筛选生效",
      "user1 当前活跃收藏 %s 条" % cur_u1, "admin POST export?userId=%s" % uid1,
      "xlsx 数据行数 = user1 活跃收藏数",
      st == 200 and is_xlsx(res) and n_u1 == cur_u1,
      "全量行=%s user1行=%s 期望=%s" % (navail, n_u1, cur_u1))

    # ==================== P6 分享来源（收藏的特殊来源） ====================
    st, res = call("POST", "/bankShare/bankShare/create", t1, {"bankId": bank_b})
    share_token = dd(res).get("shareToken")
    share_id = dd(res).get("shareId")
    st2, res2 = call("POST", "/bankShare/bankShare/receive", t2, {"token": share_token})
    T("P32", M, "分享领取-写入收藏并标记来源（source=share/canEdit=false）",
      "user1 分享 bankB；user2 领取（token=%s）" % str(share_token)[:12],
      "POST /bankShare/bankShare/receive 后 user2 请求收藏列表",
      "code=200；user2 收藏列表出现 bankB，且 source=share、canEdit=false、shareStatus=valid",
      code_of(res2) == 200,
      desc(res2, st2))

    st, res = call("GET", "/favoriteBank/favoriteBank/all", t2)
    mine = [x for x in dl(res) if x.get("bankId") == bank_b]
    T("P33", M, "分享领取-收藏列表回显分享来源字段",
      "user2 已领取 bankB 分享", "user2 GET /favoriteBank/favoriteBank/all 查找 bankB",
      "source=share、canEdit=False、shareStatus=valid、shareOwnerName 非空",
      len(mine) == 1 and mine[0].get("source") == "share" and mine[0].get("canEdit") is False
      and mine[0].get("shareStatus") == "valid" and bool(mine[0].get("shareOwnerName")),
      "命中=%d 行=%s" % (len(mine), {k: mine[0].get(k) for k in
                                     ("source", "canEdit", "shareStatus", "shareOwnerName")} if mine else {}))

    call("POST", "/bankShare/bankShare/revoke/%s" % share_id, t1)
    st, res = call("GET", "/favoriteBank/favoriteBank/all", t2)
    mine = [x for x in dl(res) if x.get("bankId") == bank_b]
    T("P34", M, "分享撤销-收藏标记为已失效（shareStatus=invalid）",
      "user1 已撤销该分享", "user2 GET /favoriteBank/favoriteBank/all 再查 bankB",
      "shareStatus=invalid",
      len(mine) == 1 and mine[0].get("shareStatus") == "invalid",
      "shareStatus=%s" % (mine[0].get("shareStatus") if mine else "未命中"))

    # ==================== P7 前端契约 ====================
    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta,
                   params={"userId": uid1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    sample = rows[0] if rows else {}
    table_cols = ["favoriteId", "userId", "bankId", "bankName", "notes", "tags", "sortOrder",
                  "isStarred", "lastStudyTime", "studyCount", "favoriteStatus"]
    missing_cols = [c for c in table_cols if c not in sample]
    T("P35", M, "前端契约-列表返回字段覆盖后台表格所有列",
      "user1 有收藏记录", "比对 index.vue 表格 el-table-column 的 prop 与返回字段",
      "返回字段包含表格全部 11 个 prop（缺 sortOrder 等会显示空白）",
      code_of(res) == 200 and bool(rows) and not missing_cols,
      "缺失字段=%s" % (missing_cols or "无"))

    st, res = call("GET", "/favoriteBank/favoriteBank/list", ta, params={
        "userId": uid1, "bankId": bank_a, "bankName": "收藏A", "notes": "本人修改备注",
        "tags": "重做", "isStarred": 0, "favoriteStatus": 1,
        "studyCount": 12, "lastStudyTime": "2026-09-08", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("P36", M, "前端契约-后台列表页 9 个查询条件同时生效",
      "user1 对 bankA 的收藏（备注/标签/排序/学习次数/时间均已设置）",
      "按 index.vue 的 queryParams 原样请求 9 个条件",
      "命中 1 条且各字段与条件一致",
      code_of(res) == 200 and len(rows) == 1
      and rows[0].get("userId") == uid1 and rows[0].get("bankId") == bank_a
      and rows[0].get("studyCount") == 12 and rows[0].get("isStarred") == 0,
      "命中 %d 条" % len(rows))

    # ==================== P8 题库ID校验（后台对话框自动填充依赖） ====================
    st, res = call("GET", "/favoriteBank/favoriteBank/check/%s" % bank_a, ta)
    d = dd(res)
    T("P37", M, "校验-题库ID存在时返回对象体 {exists, bankId, bankName}",
      "bankA=%s（《%s》）存在" % (bank_a, BANK_A),
      "GET /favoriteBank/favoriteBank/check/{bank_a}",
      "data 为对象且 exists=true、bankId 正确、bankName=题库名（前端靠它自动填充）",
      code_of(res) == 200 and d.get("exists") is True and d.get("bankId") == bank_a
      and d.get("bankName") == BANK_A,
      "data=%s" % d)

    st, res = call("GET", "/favoriteBank/favoriteBank/check/99999999", ta)
    d = dd(res)
    T("P38", M, "校验-不存在的题库ID返回 exists=false（而非报错）",
      "题库 99999999 不存在", "GET /favoriteBank/favoriteBank/check/99999999",
      "code=200 且 data.exists=false",
      code_of(res) == 200 and d.get("exists") is False,
      "data=%s" % d)

    # 回归用户反馈：真实存在的题库（如 175）不应被判成「题库ID不存在」
    any_bank = db_one("SELECT id, bank_name FROM question_bank ORDER BY id LIMIT 1", {})
    st, res = call("GET", "/favoriteBank/favoriteBank/check/%s" % (any_bank or {}).get("id"), ta)
    d = dd(res)
    T("P39", M, "校验-库中真实存在的题库不会被误判为不存在",
      "取库中任意题库 id=%s（含用户反馈的 175 这类记录）" % (any_bank or {}).get("id"),
      "GET /favoriteBank/favoriteBank/check/{id}",
      "exists=true 且 bankName 非空（历史实现返回裸布尔，前端读 data.exists 必为 undefined）",
      code_of(res) == 200 and d.get("exists") is True and bool(d.get("bankName")),
      "data=%s" % d)

    # ==================== P9 权限补充 ====================
    st, res = call("GET", "/favoriteBank/favoriteBank/checkFavorite/%s/%s" % (uid1, bank_a), t2)
    st2, res2 = call("GET", "/favoriteBank/favoriteBank/checkFavorite/%s/%s" % (uid2, bank_c), t2)
    T("P40", M, "权限-普通用户不能查询他人的收藏状态（路径 uid 不可信）",
      "user2 登录；bankA 的收藏属于 user1",
      "user2 请求 checkFavorite/{uid1}/{bankA}（他人），再请求 checkFavorite/{uid2}/{bankC}（自己）",
      "查他人 code=403；查自己 code=200",
      code_of(res) == 403 and code_of(res2) == 200,
      "查他人 code=%s msg=%s；查自己 code=%s" % (code_of(res), res.get("msg"), code_of(res2)))

    st, res = call("POST", "/favoriteBank/favoriteBank", t1,
                   {"userId": uid2, "bankId": bank_a})
    p2 = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                {"u": uid2, "b": bank_a})
    p1 = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE user_id=:u AND bank_id=:b",
                {"u": uid1, "b": bank_a})
    T("P41", M, "权限-普通用户不能代他人新增收藏记录",
      "user1 登录，body 里 userId=%s" % uid2,
      "user1 请求 POST /favoriteBank/favoriteBank  body={userId:user2,bankId:bankA}",
      "code=403；不得写进 user2 名下，也不应悄悄写到自己名下（user1 本就已收藏 bankA）",
      code_of(res) == 403 and (p2 or {}).get("c") == 0 and (p1 or {}).get("c") == 1,
      "code=%s msg=%s；user2记录=%s user1记录=%s"
      % (code_of(res), res.get("msg"), (p2 or {}).get("c"), (p1 or {}).get("c")))

    # ==================== P10 清理 ====================
    for bid in (bank_a, bank_b, bank_c):
        call("DELETE", "/questionBank/questionBank/%s" % bid, t1 if bid != bank_c else t2)
    st, res = call("GET", "/questionBank/questionBank/%s" % bank_a, t1)
    T("P42", M, "清理-删除临时题库（级联清理收藏/分享）",
      "测试题库 bankA=%s" % bank_a, "DELETE 后查详情",
      "code=200；删除后详情不可见",
      code_of(res) != 200, "详情code=%s" % code_of(res))

    ids = (bank_a, bank_b, bank_c)
    db_exec("DELETE FROM bank_share_receive WHERE bank_id IN :i", {"i": ids})
    db_exec("DELETE FROM bank_share WHERE bank_id IN :i", {"i": ids})
    db_exec("DELETE FROM user_bank_favorite WHERE bank_id IN :i", {"i": ids})
    db_exec("DELETE FROM question_sub WHERE main_id IN (SELECT id FROM question_main WHERE bank_id IN :b)",
            {"b": ids})
    db_exec("DELETE FROM question_main WHERE bank_id IN :b", {"b": ids})
    db_exec("DELETE FROM question_bank WHERE id IN :i", {"i": ids})
    db_exec("DELETE FROM user_notification WHERE content LIKE '%题库收藏B-分享题库%'")
    db_exec("DELETE FROM bank_share_receive WHERE share_id NOT IN (SELECT id FROM bank_share)")

    left_bank = db_one("SELECT COUNT(*) c FROM question_bank WHERE id IN :i", {"i": ids})
    left_fav = db_one("SELECT COUNT(*) c FROM user_bank_favorite WHERE bank_id IN :i", {"i": ids})
    left_share = db_one("SELECT COUNT(*) c FROM bank_share WHERE bank_id IN :i", {"i": ids})
    left_recv = db_one("SELECT COUNT(*) c FROM bank_share_receive WHERE bank_id IN :i", {"i": ids})
    T("P43", M, "清理-无测试数据残留（不污染演示数据）",
      "清理完成", "直连库核对测试题库/收藏/分享/领取记录",
      "题库=0 收藏=0 分享=0 领取记录=0",
      all((x or {}).get("c") == 0 for x in (left_bank, left_fav, left_share, left_recv)),
      "题库=%s 收藏=%s 分享=%s 领取=%s"
      % ((left_bank or {}).get("c"), (left_fav or {}).get("c"),
         (left_share or {}).get("c"), (left_recv or {}).get("c")))

    total = len(ROWS)
    passed = sum(1 for r in ROWS if r["status"] == "通过")
    print("=" * 78)
    print("用例总数 %d，通过 %d，失败 %d" % (total, passed, total - passed))
    print("=" * 78)
    for r in ROWS:
        if r["status"] == "失败":
            print("  [FAIL] %s %s" % (r["cid"], r["title"]))
            print("         预期: %s" % r["expected"])
            print("         实际: %s" % r["actual"])

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump({
            "prefix": P, "uid_admin": uid_admin, "uid1": uid1, "uid2": uid2, "rows": ROWS,
            "cleanup": {"bank_ids": [x for x in (bank_a, bank_b, bank_c) if x]},
        }, f, ensure_ascii=False, indent=2)
    print("结果已写入", OUT)


if __name__ == "__main__":
    main()
