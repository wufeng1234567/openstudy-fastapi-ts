# -*- coding: utf-8 -*-
"""后台管理页面「题目收藏」模块 E2E 接口测试（真实 HTTP 调用）

被测页面：openstudy-vue3-ts/src/views/favoriteQuestion/favoriteQuestion/index.vue
被测接口：app/api/v1/questions.py 的 /favoriteQuestion/favoriteQuestion/*
账号角色：admin(1) 管理员 / user1(105) 普通用户 / user2(106) 普通用户
用例编号：Q*（Q=题目收藏）

判定口径（与「题库收藏/斩题管理/错题管理」及 SpringBoot 参考版一致）：
  - list 支持 userId/questionId/notes/tags/difficultyRating/errorTimes/correctTimes/
    reviewCount/isStarred/favoriteStatus/lastReviewTime/nextReviewTime + 分页
  - 管理员不传 userId → 可看全部用户；普通用户 → 只能看自己
  - 新增：管理员可代任意用户录入（payload.userId），全字段落库；重复收藏 400；
    题目/用户不存在要有明确提示
  - 详情/修改/删除：管理员可操作任意用户记录，普通用户越权必须 403
  - 返回体必须带 bankName（后台「所属题库」列）、difficultyRating（难度评级列）
  - 导出 /export 支持与列表一致的筛选

用法（后端需先跑起来，默认 127.0.0.1:9101）：
    cd openstudy-fastapi
    .venv/Scripts/python.exe tests/e2e_favorite_question_admin_api.py

产出：脚本同目录 e2e_favorite_question_admin_api_result.json
说明：全部测试数据建在临时题库/题目上，结束时删除题库（级联清理），不留残留。
"""
import json
import os
import sys
import time
import urllib.error
import urllib.request
from urllib.parse import urlencode

BASE = "http://127.0.0.1:9101"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "e2e_favorite_question_admin_api_result.json")

ROWS = []


def call(method, path, token=None, body=None, params=None, timeout=25):
    url = BASE + path
    if params:
        clean = {k: v for k, v in params.items() if v is not None}
        if clean:
            url += "?" + urlencode(clean, doseq=True)
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read()
            ct = r.headers.get("Content-Type", "")
            if "json" in ct:
                try:
                    return r.status, json.loads(raw.decode("utf-8"))
                except Exception:
                    return r.status, raw
            return r.status, raw
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode("utf-8"))
        except Exception:
            return e.code, raw


def login(account):
    st, res = call("POST", "/quickLogin", body={"account": account})
    return res["data"]["token"]


def T(cid, module, title, pre, steps, expected, ok, actual):
    ROWS.append({
        "cid": cid, "module": module, "title": title, "pre": pre,
        "steps": steps, "expected": expected, "actual": actual,
        "status": "通过" if ok else "失败",
    })
    print(("  [PASS] " if ok else "  [FAIL] ") + cid + " " + title + " -> " + str(actual))


def code_of(res):
    return res.get("code") if isinstance(res, dict) else "-"


def raw_data(res):
    return res.get("data") if isinstance(res, dict) else None


def dd(res):
    d = raw_data(res)
    return d if isinstance(d, dict) else {}


def dl(res):
    d = raw_data(res)
    return d if isinstance(d, list) else []


def rows_of(res):
    return (res.get("rows") or []) if isinstance(res, dict) else []


def is_file(res):
    return isinstance(res, (bytes, bytearray)) and len(res) > 0


def is_xlsx(res):
    return is_file(res) and res[:2] == b"PK"


def desc(res, st=200):
    if is_file(res):
        return "HTTP %s 文件流 大小=%d 字节" % (st, len(res))
    msg = (res.get("msg") or "")[:80] if isinstance(res, dict) else ""
    return "HTTP %s code=%s %s" % (st, code_of(res), msg)


def _db(sql, params=None, write=False):
    try:
        import asyncio

        from sqlalchemy import text

        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        if root not in sys.path:
            sys.path.insert(0, root)

        from app.config import settings
        from sqlalchemy.ext.asyncio import create_async_engine
        from sqlalchemy.pool import NullPool

        async def _run():
            engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
            try:
                if write:
                    async with engine.begin() as conn:
                        return (await conn.execute(text(sql), params or {})).rowcount
                async with engine.connect() as conn:
                    return (await conn.execute(text(sql), params or {})).all()
            finally:
                await engine.dispose()

        return asyncio.run(_run())
    except Exception as exc:  # noqa: BLE001
        print("    (DB %s 不可用：%s)" % ("写入" if write else "直读", exc))
        return None


def db_query(sql, params=None):
    return _db(sql, params, write=False)


def db_exec(sql, params=None):
    return _db(sql, params, write=True)


def db_one(sql, params=None):
    rows = db_query(sql, params)
    if not rows:
        return None
    return dict(rows[0]._mapping)


def new_bank(tok, name):
    st, res = call("POST", "/questionBank/questionBank", tok, {"bankName": name})
    d = dd(res)
    return d.get("id"), d.get("userId")


def new_question(tok, bank_id, text):
    st, res = call("POST", "/questionMain/questionMain", tok, {
        "bankId": bank_id, "questionType": 1, "questionText": text,
        "options": [{"key": "A", "value": "1"}, {"key": "B", "value": "2"}], "answer": "A",
    })
    return dd(res).get("questionId")


def sweep_leftovers(created):
    bank_ids = created.get("bank_ids") or []
    if not bank_ids:
        return
    bank_t = tuple(bank_ids) or (0,)
    db_exec("DELETE FROM user_question_favorite WHERE bank_id IN :i", {"i": bank_t})
    db_exec("DELETE FROM question_sub WHERE main_id IN (SELECT id FROM question_main WHERE bank_id IN :b)",
            {"b": bank_t})
    db_exec("DELETE FROM question_main WHERE bank_id IN :b", {"b": bank_t})
    db_exec("DELETE FROM question_bank WHERE id IN :i", {"i": bank_t})
    print("  [自检] 已清理上次运行遗留：bank=%s" % bank_ids)


def main():
    if os.path.exists(OUT):
        try:
            with open(OUT, encoding="utf-8") as f:
                sweep_leftovers(json.load(f).get("cleanup") or {})
        except Exception as exc:  # noqa: BLE001
            print("  [自检] 读取上次结果失败：%s" % exc)

    P = "题目收藏"
    BANK_A = P + "A-测试题库"
    BANK_C = P + "C-用户2题库"

    print("=" * 78)
    print("后台管理「题目收藏」模块 E2E 测试   前缀 =", P)
    print("=" * 78)

    ta = login("admin")
    t1 = login("user1")
    t2 = login("user2")

    bank_a, uid1 = new_bank(t1, BANK_A)
    bank_c, uid2 = new_bank(t2, BANK_C)
    probe, uid_admin = new_bank(ta, P + "管理员探针")
    call("DELETE", "/questionBank/questionBank/%s" % probe, ta)

    q1 = new_question(t1, bank_a, P + "-主力题1")
    q2 = new_question(t1, bank_a, P + "-状态校验题2")
    q3 = new_question(t1, bank_a, P + "-后台新增题3")
    q4 = new_question(t2, bank_c, P + "-用户2题4")
    print("bankA=%s bankC=%s uid_admin=%s uid1=%s uid2=%s" % (bank_a, bank_c, uid_admin, uid1, uid2))
    print("q1..q4 =", q1, q2, q3, q4)

    # ---- 基线收藏（走接口）----
    call("POST", "/favoriteQuestion/favoriteQuestion", t1, {"questionId": q1, "bankId": bank_a})
    call("POST", "/favoriteQuestion/favoriteQuestion", t1, {"questionId": q2, "bankId": bank_a})
    call("POST", "/favoriteQuestion/favoriteQuestion", t2, {"questionId": q4, "bankId": bank_c})

    # ---- DB 造有区分度的属性（供筛选/回显用例使用，不依赖待测写接口）----
    db_exec("""UPDATE user_question_favorite SET notes='收藏备注Q1-重点整理', tags='重点,错题',
               difficulty_rating=4, error_times=3, correct_times=5, review_count=6, is_starred=1,
               last_review_time='2026-09-02 10:00:00', next_review_time='2026-09-18 10:00:00'
               WHERE user_id=:u AND question_id=:q""", {"u": uid1, "q": q1})
    db_exec("""UPDATE user_question_favorite SET notes='收藏备注Q2', tags='复习',
               difficulty_rating=2, error_times=0, correct_times=1, review_count=1, is_starred=0
               WHERE user_id=:u AND question_id=:q""", {"u": uid1, "q": q2})
    db_exec("""UPDATE user_question_favorite SET notes='用户2自己的收藏', tags='重点',
               difficulty_rating=5, error_times=1, correct_times=2, review_count=3, is_starred=1
               WHERE user_id=:u AND question_id=:q""", {"u": uid2, "q": q4})

    def cnt(sql, params):
        return (db_one(sql, params) or {}).get("c")

    exp_u1 = cnt("SELECT COUNT(*) c FROM user_question_favorite WHERE user_id=:u AND favorite_status=1", {"u": uid1})
    exp_u2 = cnt("SELECT COUNT(*) c FROM user_question_favorite WHERE user_id=:u AND favorite_status=1", {"u": uid2})
    exp_star = cnt("SELECT COUNT(*) c FROM user_question_favorite WHERE is_starred=1 AND favorite_status=1", {})
    print("基线：user1=%s user2=%s 已标星=%s" % (exp_u1, exp_u2, exp_star))

    M = "题目收藏(后台)"

    # ==================== Q1 列表查询与筛选（对应后台列表页的查询条件） ====================
    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta, params={"pageNum": 1, "pageSize": 10})
    T("Q01", M, "列表-管理员查询返回分页结构",
      "admin 已登录", "GET /favoriteQuestion/favoriteQuestion/list",
      "code=200；顶层含 total 与 rows",
      code_of(res) == 200 and "total" in res and "rows" in res,
      "code=%s total=%s rows=%d" % (code_of(res), res.get("total"), len(rows_of(res))))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"userId": uid1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q02", M, "列表-按用户ID筛选",
      "user1 有 %s 条活跃收藏" % exp_u1, "admin 请求 list?userId=%s" % uid1,
      "total=%s 且 rows 的 userId 全为 %s" % (exp_u1, uid1),
      code_of(res) == 200 and res.get("total") == exp_u1 and len(rows) > 0
      and all(r.get("userId") == uid1 for r in rows),
      "total=%s userId集合=%s" % (res.get("total"), sorted({r.get("userId") for r in rows})[:5]))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"questionId": q1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q03", M, "列表-按题目ID精确筛选",
      "user1 收藏了 q1=%s" % q1, "admin 请求 list?questionId=%s" % q1,
      "命中 1 条且 questionId=q1",
      code_of(res) == 200 and len(rows) == 1 and rows[0].get("questionId") == q1,
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"notes": "备注Q1", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q04", M, "列表-按收藏备注模糊筛选",
      "存在备注含「备注Q1」的收藏", "admin 请求 list?notes=备注Q1",
      "命中且 rows 的 notes 均含关键字",
      code_of(res) == 200 and len(rows) > 0 and all("备注Q1" in (r.get("notes") or "") for r in rows),
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"tags": "错题", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q05", M, "列表-按标签模糊筛选",
      "存在标签含「错题」的收藏", "admin 请求 list?tags=错题",
      "命中且 rows 的 tags 均含关键字",
      code_of(res) == 200 and len(rows) > 0 and all("错题" in (r.get("tags") or "") for r in rows),
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"difficultyRating": 4, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q06", M, "列表-按难度评级筛选",
      "存在难度评级=4 的收藏", "admin 请求 list?difficultyRating=4",
      "命中且 rows 的 difficultyRating 全为 4",
      code_of(res) == 200 and len(rows) > 0 and all(r.get("difficultyRating") == 4 for r in rows),
      "命中 %d 条 评级=%s" % (len(rows), sorted({r.get("difficultyRating") for r in rows})[:5]))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"errorTimes": 3, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q07", M, "列表-按错误次数筛选",
      "存在错误次数=3 的收藏", "admin 请求 list?errorTimes=3",
      "命中且 rows 的 errorTimes 全为 3",
      code_of(res) == 200 and len(rows) > 0 and all(r.get("errorTimes") == 3 for r in rows),
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"correctTimes": 5, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q08", M, "列表-按正确次数筛选",
      "存在正确次数=5 的收藏", "admin 请求 list?correctTimes=5",
      "命中且 rows 的 correctTimes 全为 5",
      code_of(res) == 200 and len(rows) > 0 and all(r.get("correctTimes") == 5 for r in rows),
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"reviewCount": 6, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q09", M, "列表-按复习次数筛选",
      "存在复习次数=6 的收藏", "admin 请求 list?reviewCount=6",
      "命中且 rows 的 reviewCount 全为 6",
      code_of(res) == 200 and len(rows) > 0 and all(r.get("reviewCount") == 6 for r in rows),
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"isStarred": 1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q10", M, "列表-按是否标星筛选",
      "全库已标星 %s 条" % exp_star, "admin 请求 list?isStarred=1",
      "total=%s 且 rows 的 isStarred 全为 1" % exp_star,
      code_of(res) == 200 and res.get("total") == exp_star and len(rows) > 0
      and all(r.get("isStarred") == 1 for r in rows),
      "total=%s isStarred集合=%s" % (res.get("total"), sorted({r.get("isStarred") for r in rows})[:5]))

    # favoriteStatus：默认只查"正常收藏(1)"，显式传 2 才看"已掌握"
    db_exec("UPDATE user_question_favorite SET favorite_status=2 WHERE user_id=:u AND question_id=:q",
            {"u": uid1, "q": q2})
    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"pageNum": 1, "pageSize": 100})
    default_ids = {r.get("favoriteId") for r in rows_of(res)}
    st2, res2 = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                     params={"favoriteStatus": 2, "pageNum": 1, "pageSize": 100})
    st_rows = rows_of(res2)
    st_id = (db_one("SELECT favorite_id FROM user_question_favorite WHERE user_id=:u AND question_id=:q",
                    {"u": uid1, "q": q2}) or {}).get("favorite_id")
    T("Q11", M, "列表-按收藏状态筛选（默认只返回正常收藏）",
      "user1 对 q2 的收藏已置为 favorite_status=2（已掌握）",
      "① list 不传 favoriteStatus ② list?favoriteStatus=2",
      "① 默认结果不含已掌握 ② favoriteStatus=2 精确命中该条",
      code_of(res) == 200 and code_of(res2) == 200 and st_id not in default_ids
      and len(st_rows) == 1 and st_rows[0].get("favoriteId") == st_id
      and st_rows[0].get("favoriteStatus") == 2,
      "默认含已掌握=%s；favoriteStatus=2 命中 %d 条" % (st_id in default_ids, len(st_rows)))
    db_exec("UPDATE user_question_favorite SET favorite_status=1 WHERE favorite_id=:f", {"f": st_id})

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"lastReviewTime": "2026-09-02", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q12", M, "列表-按最后复习时间筛选（当天）",
      "存在最后复习时间=2026-09-02 的收藏", "admin 请求 list?lastReviewTime=2026-09-02",
      "命中且 rows 的 lastReviewTime 落在该日",
      code_of(res) == 200 and len(rows) > 0
      and all((r.get("lastReviewTime") or "").startswith("2026-09-02") for r in rows),
      "命中 %d 条 %s" % (len(rows), [r.get("lastReviewTime") for r in rows][:3]))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"nextReviewTime": "2026-09-18", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q13", M, "列表-按下次复习时间筛选（当天）",
      "存在下次复习时间=2026-09-18 的收藏", "admin 请求 list?nextReviewTime=2026-09-18",
      "命中且 rows 的 nextReviewTime 落在该日",
      code_of(res) == 200 and len(rows) > 0
      and all((r.get("nextReviewTime") or "").startswith("2026-09-18") for r in rows),
      "命中 %d 条 %s" % (len(rows), [r.get("nextReviewTime") for r in rows][:3]))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"pageNum": 1, "pageSize": 2})
    T("Q14", M, "列表-分页返回条数受 pageSize 约束",
      "全库活跃收藏 > 2 条", "admin 请求 list?pageSize=2",
      "rows<=2 且 total 为全量",
      code_of(res) == 200 and len(rows_of(res)) <= 2 and (res.get("total") or 0) > 2,
      "total=%s rows=%d" % (res.get("total"), len(rows_of(res))))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"pageNum": 1, "pageSize": 100})
    owners = sorted({r.get("userId") for r in rows_of(res)})
    T("Q15", M, "列表-管理员不传 userId 时可看到全部用户",
      "admin/user1/user2 均有收藏", "admin 请求 list（不传 userId）",
      "rows 中出现 >=2 个不同 userId",
      code_of(res) == 200 and len(owners) >= 2, "total=%s 涉及用户=%s" % (res.get("total"), owners[:6]))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"userId": uid1, "questionId": q1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q16", M, "列表-组合条件筛选（userId + questionId）",
      "user1 收藏了 q1", "admin 请求 list?userId=%s&questionId=%s" % (uid1, q1),
      "命中 1 条且两者都匹配",
      code_of(res) == 200 and len(rows) == 1
      and rows[0].get("userId") == uid1 and rows[0].get("questionId") == q1,
      "命中 %d 条" % len(rows))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", t2,
                   params={"pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q17", M, "列表-普通用户只能看到自己的收藏",
      "user2 有 %s 条活跃收藏" % exp_u2, "user2 请求 list?pageSize=100",
      "rows 的 userId 全为 %s" % uid2,
      code_of(res) == 200 and len(rows) > 0 and all(r.get("userId") == uid2 for r in rows),
      "total=%s userId集合=%s" % (res.get("total"), sorted({r.get("userId") for r in rows})[:5]))

    # ==================== Q2 新增（后台对话框全字段） ====================
    add_body = {
        "userId": uid1, "questionId": q3, "bankId": bank_a,
        "notes": P + "-后台新增备注", "tags": "必考", "difficultyRating": 3,
        "errorTimes": 2, "correctTimes": 4, "reviewCount": 7, "isStarred": 1, "favoriteStatus": 1,
        "lastReviewTime": "2026-09-06 08:00:00", "nextReviewTime": "2026-09-21 08:00:00",
    }
    st, res = call("POST", "/favoriteQuestion/favoriteQuestion", ta, add_body)
    T("Q18", M, "新增-管理员代 user1 收藏题目（提交后台表单全字段）",
      "q3=%s 未被 user1 收藏" % q3,
      "admin POST /favoriteQuestion/favoriteQuestion  body 含 userId=%s 与全部表单字段" % uid1,
      "code=200", code_of(res) == 200, desc(res, st))

    row = db_one("""SELECT user_id, bank_id, notes, tags, difficulty_rating, error_times, correct_times,
                           review_count, is_starred, favorite_status, last_review_time, next_review_time
                    FROM user_question_favorite WHERE user_id=:u AND question_id=:q""",
                 {"u": uid1, "q": q3})
    T("Q19", M, "新增-表单字段全部落库（归属=被代录用户）",
      "Q18 已提交全字段", "直读 user_question_favorite 逐字段比对",
      "user_id=%s；备注/标签一致；评级=3 错=2 对=4 复习=7 标星=1 状态=1；复习时间非空" % uid1,
      bool(row) and row.get("user_id") == uid1 and row.get("bank_id") == bank_a
      and row.get("notes") == P + "-后台新增备注" and row.get("tags") == "必考"
      and row.get("difficulty_rating") == 3 and row.get("error_times") == 2
      and row.get("correct_times") == 4 and row.get("review_count") == 7
      and row.get("is_starred") == 1 and row.get("favorite_status") == 1
      and row.get("last_review_time") is not None and row.get("next_review_time") is not None,
      "库中=%s" % row)

    fid_new = (db_one("SELECT favorite_id FROM user_question_favorite WHERE user_id=:u AND question_id=:q",
                      {"u": uid1, "q": q3}) or {}).get("favorite_id")
    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/%s" % fid_new, ta)
    d = dd(res)
    need = ["favoriteId", "userId", "questionId", "bankId", "bankName", "notes", "tags",
            "difficultyRating", "errorTimes", "correctTimes", "reviewCount", "isStarred",
            "favoriteStatus", "lastReviewTime", "nextReviewTime", "createTime", "updateTime"]
    missing = [f for f in need if f not in d]
    T("Q20", M, "详情-后台修改对话框所需字段完整回显",
      "favoriteId=%s" % fid_new, "admin GET /favoriteQuestion/favoriteQuestion/{favoriteId}",
      "返回体包含后台表格/表单的 17 个字段（缺 bankName/difficultyRating 会导致页面空白）",
      code_of(res) == 200 and not missing,
      "code=%s 缺失字段=%s" % (code_of(res), missing or "无"))

    st, res = call("POST", "/favoriteQuestion/favoriteQuestion", ta, add_body)
    T("Q21", M, "新增-同一用户重复收藏同一题给出业务提示",
      "user1 已收藏 q3", "admin 再次 POST 同一 body",
      "code=400 且提示「已收藏」，不应 500",
      code_of(res) == 400, "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("POST", "/favoriteQuestion/favoriteQuestion", ta,
                   {"userId": uid1, "questionId": 99999999})
    T("Q22", M, "新增-题目不存在时的异常处理",
      "题目 99999999 不存在", "admin POST body={questionId:99999999}",
      "code!=200 且提示「题目」+「不存在」",
      code_of(res) != 200 and "题目" in (res.get("msg") or "") and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("POST", "/favoriteQuestion/favoriteQuestion", ta,
                   {"userId": 99999999, "questionId": q1})
    T("Q23", M, "新增-用户不存在时的异常处理",
      "userId=99999999 不存在", "admin POST body={userId:99999999}",
      "code!=200 且提示「用户」+「不存在」",
      code_of(res) != 200 and "用户" in (res.get("msg") or "") and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    # ==================== Q3 修改 ====================
    fid1 = (db_one("SELECT favorite_id FROM user_question_favorite WHERE user_id=:u AND question_id=:q",
                   {"u": uid1, "q": q1}) or {}).get("favorite_id")
    upd = {
        "favoriteId": fid1, "notes": P + "-本人修改备注", "tags": "重做",
        "difficultyRating": 5, "errorTimes": 8, "correctTimes": 9, "reviewCount": 11,
        "isStarred": 0, "favoriteStatus": 1,
        "lastReviewTime": "2026-09-09 20:00:00", "nextReviewTime": "2026-09-25 20:00:00",
    }
    st, res = call("PUT", "/favoriteQuestion/favoriteQuestion", t1, upd)
    row1 = db_one("""SELECT notes, tags, difficulty_rating, error_times, correct_times, review_count,
                            is_starred, favorite_status, last_review_time, next_review_time
                     FROM user_question_favorite WHERE favorite_id=:f""", {"f": fid1})
    T("Q24", M, "修改-本人提交后台表单全字段，字段全部更新",
      "favoriteId=%s（user1 本人记录）" % fid1,
      "user1 PUT body 含后台表单全部字段",
      "code=200；库中备注/标签/评级=5 错=8 对=9 复习=11 标星=0 状态=1 复习时间非空",
      code_of(res) == 200 and bool(row1)
      and row1.get("notes") == P + "-本人修改备注" and row1.get("tags") == "重做"
      and row1.get("difficulty_rating") == 5 and row1.get("error_times") == 8
      and row1.get("correct_times") == 9 and row1.get("review_count") == 11
      and row1.get("is_starred") == 0 and row1.get("favorite_status") == 1
      and row1.get("last_review_time") is not None and row1.get("next_review_time") is not None,
      "code=%s 库中=%s" % (code_of(res), row1))

    fid3 = (db_one("SELECT favorite_id FROM user_question_favorite WHERE user_id=:u AND question_id=:q",
                   {"u": uid1, "q": q3}) or {}).get("favorite_id")
    upd_admin = {"favoriteId": fid3, "notes": P + "-管理员修改备注", "reviewCount": 20,
                 "difficultyRating": 2, "errorTimes": 6}
    st, res = call("PUT", "/favoriteQuestion/favoriteQuestion", ta, upd_admin)
    row3 = db_one("""SELECT notes, review_count, difficulty_rating, error_times
                     FROM user_question_favorite WHERE favorite_id=:f""", {"f": fid3})
    T("Q25", M, "修改-管理员可修改他人（user1）的收藏记录",
      "favoriteId=%s 属于 user1" % fid3,
      "admin PUT /favoriteQuestion/favoriteQuestion  body={favoriteId:%s,...}" % fid3,
      "code=200；库中备注/reviewCount=20/评级=2/错误次数=6 均已更新",
      code_of(res) == 200 and bool(row3)
      and row3.get("notes") == P + "-管理员修改备注" and row3.get("review_count") == 20
      and row3.get("difficulty_rating") == 2 and row3.get("error_times") == 6,
      "code=%s 库中=%s" % (code_of(res), row3))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/%s" % fid3, ta)
    d = dd(res)
    T("Q26", M, "详情-管理员可查看他人收藏记录",
      "favoriteId=%s 属于 user1" % fid3, "admin GET /favoriteQuestion/favoriteQuestion/{favoriteId}",
      "code=200 且 userId=%s、reviewCount=20、bankName 非空" % uid1,
      code_of(res) == 200 and d.get("userId") == uid1 and d.get("reviewCount") == 20
      and bool(d.get("bankName")),
      "code=%s userId=%s reviewCount=%s bankName=%s" % (code_of(res), d.get("userId"),
                                                        d.get("reviewCount"), d.get("bankName")))

    st, res = call("PUT", "/favoriteQuestion/favoriteQuestion", ta, {"favoriteId": 99999999, "notes": "x"})
    T("Q27", M, "修改-记录不存在时的异常处理",
      "favoriteId=99999999 不存在", "admin PUT body={favoriteId:99999999}",
      "code!=200 且提示「收藏记录不存在」",
      code_of(res) != 200 and "不存在" in (res.get("msg") or ""),
      "code=%s msg=%s" % (code_of(res), res.get("msg")))

    st, res = call("PUT", "/favoriteQuestion/favoriteQuestion", t2,
                   {"favoriteId": fid1, "notes": P + "-越权篡改"})
    still = db_one("SELECT notes FROM user_question_favorite WHERE favorite_id=:f", {"f": fid1})
    T("Q28", M, "修改-普通用户不能修改他人收藏（越权拦截）",
      "favoriteId=%s 属于 user1；user2 登录" % fid1,
      "user2 PUT body={favoriteId:%s}" % fid1,
      "code=403；user1 的记录未被篡改",
      code_of(res) == 403 and (still or {}).get("notes") == P + "-本人修改备注",
      "code=%s msg=%s 现值=%s" % (code_of(res), res.get("msg"), (still or {}).get("notes")))

    # ==================== Q4 删除 ====================
    st, res = call("DELETE", "/favoriteQuestion/favoriteQuestion/%s" % fid3, ta)
    gone = db_one("SELECT COUNT(*) c FROM user_question_favorite WHERE favorite_id=:f", {"f": fid3})
    T("Q29", M, "删除-管理员删除他人（user1）单条收藏",
      "favoriteId=%s 属于 user1" % fid3, "admin DELETE /favoriteQuestion/favoriteQuestion/{favoriteId}",
      "code=200；库中该记录已物理删除",
      code_of(res) == 200 and (gone or {}).get("c") == 0,
      "code=%s msg=%s 剩余=%s" % (code_of(res), res.get("msg"), (gone or {}).get("c")))

    # 造两条给批量删除
    call("POST", "/favoriteQuestion/favoriteQuestion", ta, {"userId": uid1, "questionId": q3})
    call("POST", "/favoriteQuestion/favoriteQuestion", ta, {"userId": uid2, "questionId": q3})
    ids_rows = db_query("SELECT favorite_id FROM user_question_favorite WHERE question_id=:q", {"q": q3}) or []
    id_join = ",".join(str(r[0]) for r in ids_rows)
    st, res = call("DELETE", "/favoriteQuestion/favoriteQuestion/%s" % id_join, ta)
    left = db_one("SELECT COUNT(*) c FROM user_question_favorite WHERE favorite_id IN :i",
                  {"i": tuple(r[0] for r in ids_rows)})
    T("Q30", M, "删除-管理员批量删除多人收藏（逗号分隔 id）",
      "q3 下有 %d 条收藏，ids=%s" % (len(ids_rows), id_join),
      "admin DELETE /favoriteQuestion/favoriteQuestion/{ids}",
      "code=200；这些记录均已删除",
      code_of(res) == 200 and (left or {}).get("c") == 0,
      "code=%s 剩余=%s" % (code_of(res), (left or {}).get("c")))

    st, res = call("DELETE", "/favoriteQuestion/favoriteQuestion/%s" % fid1, t2)
    still1 = db_one("SELECT COUNT(*) c FROM user_question_favorite WHERE favorite_id=:f", {"f": fid1})
    T("Q31", M, "删除-普通用户不能删除他人收藏（越权拦截）",
      "favoriteId=%s 属于 user1；user2 登录" % fid1,
      "user2 DELETE /favoriteQuestion/favoriteQuestion/{favoriteId}",
      "code=403；记录仍存在",
      code_of(res) == 403 and (still1 or {}).get("c") == 1,
      "code=%s msg=%s 仍存在=%s" % (code_of(res), res.get("msg"), (still1 or {}).get("c")))

    st, res = call("DELETE", "/favoriteQuestion/favoriteQuestion/99999999", ta)
    T("Q32", M, "删除-不存在的 id 不应 500",
      "favoriteId=99999999 不存在", "admin DELETE /favoriteQuestion/favoriteQuestion/99999999",
      "返回业务响应（200 幂等 或 明确提示），不得 500",
      st != 500 and code_of(res) != 500, desc(res, st))

    # ==================== Q5 导出 ====================
    st, res = call("POST", "/favoriteQuestion/favoriteQuestion/export", ta)
    T("Q33", M, "导出-后台导出接口返回 Excel",
      "admin 已登录", "POST /favoriteQuestion/favoriteQuestion/export",
      "HTTP 200 且返回 xlsx 文件流",
      st == 200 and is_xlsx(res), desc(res, st) + ("，是xlsx" if is_xlsx(res) else "，不是xlsx"))

    cur_u1 = cnt("SELECT COUNT(*) c FROM user_question_favorite WHERE user_id=:u AND favorite_status=1",
                 {"u": uid1})
    st, res = call("POST", "/favoriteQuestion/favoriteQuestion/export", ta, params={"userId": uid1})
    n_u1 = None
    if is_xlsx(res):
        try:
            import io

            from openpyxl import load_workbook

            n_u1 = load_workbook(io.BytesIO(res), read_only=True).active.max_row - 1
        except Exception:  # noqa: BLE001
            pass
    T("Q34", M, "导出-按 userId 筛选生效",
      "user1 当前活跃收藏 %s 条" % cur_u1, "admin POST export?userId=%s" % uid1,
      "xlsx 数据行数 = user1 活跃收藏数",
      st == 200 and is_xlsx(res) and n_u1 == cur_u1,
      "user1行=%s 期望=%s" % (n_u1, cur_u1))

    # ==================== Q6 题目ID校验 / 权限 ====================
    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/check/%s" % q1, ta)
    d = dd(res)
    T("Q35", M, "校验-题目ID存在时返回对象体 {exists, title, bankId, bankName}",
      "q1=%s 存在" % q1, "GET /favoriteQuestion/favoriteQuestion/check/{q1}",
      "data 为对象且 exists=true、bankId=bank_a、bankName 非空（前端靠它自动填充）",
      code_of(res) == 200 and d.get("exists") is True and d.get("bankId") == bank_a
      and bool(d.get("bankName")),
      "data=%s" % {k: d.get(k) for k in ("exists", "questionId", "bankId", "bankName", "difficulty")})

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/check/99999999", ta)
    d = dd(res)
    T("Q36", M, "校验-不存在的题目ID返回 exists=false（而非报错）",
      "题目 99999999 不存在", "GET /favoriteQuestion/favoriteQuestion/check/99999999",
      "code=200 且 data.exists=false",
      code_of(res) == 200 and d.get("exists") is False, "data=%s" % d)

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/checkFavorite/%s/%s" % (uid1, q1), t2)
    st2, res2 = call("GET", "/favoriteQuestion/favoriteQuestion/checkFavorite/%s/%s" % (uid2, q4), t2)
    T("Q37", M, "权限-普通用户不能查询他人的收藏状态（路径 uid 不可信）",
      "user2 登录；q1 的收藏属于 user1",
      "user2 请求 checkFavorite/{uid1}/{q1}（他人），再请求 checkFavorite/{uid2}/{q4}（自己）",
      "查他人 code=403；查自己 code=200",
      code_of(res) == 403 and code_of(res2) == 200,
      "查他人 code=%s msg=%s；查自己 code=%s" % (code_of(res), res.get("msg"), code_of(res2)))

    st, res = call("POST", "/favoriteQuestion/favoriteQuestion", t1,
                   {"userId": uid2, "questionId": q2})
    p2 = db_one("SELECT COUNT(*) c FROM user_question_favorite WHERE user_id=:u AND question_id=:q",
                {"u": uid2, "q": q2})
    T("Q38", M, "权限-普通用户不能代他人新增收藏记录",
      "user1 登录，body 里 userId=%s（q2 user2 未收藏）" % uid2,
      "user1 请求 POST  body={userId:user2,questionId:q2}",
      "code=403；不得写进 user2 名下",
      code_of(res) == 403 and (p2 or {}).get("c") == 0,
      "code=%s msg=%s；user2记录=%s" % (code_of(res), res.get("msg"), (p2 or {}).get("c")))

    # ==================== Q7 前端契约 ====================
    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta,
                   params={"userId": uid1, "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    sample = rows[0] if rows else {}
    table_cols = ["favoriteId", "userId", "questionId", "bankName", "notes", "difficultyRating",
                  "errorTimes", "correctTimes", "lastReviewTime", "nextReviewTime", "reviewCount",
                  "isStarred", "favoriteStatus"]
    missing_cols = [c for c in table_cols if c not in sample]
    T("Q39", M, "前端契约-列表返回字段覆盖后台表格所有列",
      "user1 有收藏记录", "比对 index.vue 表格 el-table-column 的 prop 与返回字段",
      "返回字段包含表格全部 13 个 prop（缺 bankName/difficultyRating 会显示空白）",
      code_of(res) == 200 and bool(rows) and not missing_cols,
      "缺失字段=%s" % (missing_cols or "无"))

    st, res = call("GET", "/favoriteQuestion/favoriteQuestion/list", ta, params={
        "userId": uid1, "questionId": q1, "notes": "本人修改备注", "tags": "重做",
        "difficultyRating": 5, "errorTimes": 8, "correctTimes": 9, "reviewCount": 11,
        "isStarred": 0, "favoriteStatus": 1, "lastReviewTime": "2026-09-09",
        "nextReviewTime": "2026-09-25", "pageNum": 1, "pageSize": 100})
    rows = rows_of(res)
    T("Q40", M, "前端契约-后台列表页 12 个查询条件同时生效",
      "user1 对 q1 的收藏（各字段均已设置）",
      "按 index.vue 的 queryParams 原样请求 12 个条件",
      "命中 1 条且各字段与条件一致",
      code_of(res) == 200 and len(rows) == 1
      and rows[0].get("userId") == uid1 and rows[0].get("questionId") == q1
      and rows[0].get("difficultyRating") == 5 and rows[0].get("reviewCount") == 11,
      "命中 %d 条" % len(rows))

    # ==================== Q8 清理 ====================
    call("DELETE", "/questionBank/questionBank/%s" % bank_a, t1)
    call("DELETE", "/questionBank/questionBank/%s" % bank_c, t2)
    st, res = call("GET", "/questionBank/questionBank/%s" % bank_a, t1)
    T("Q41", M, "清理-删除临时题库（级联清理收藏数据）",
      "测试题库 bankA=%s" % bank_a, "DELETE 后查详情",
      "详情不可见", code_of(res) != 200, "详情code=%s" % code_of(res))

    q_ids = tuple([q for q in (q1, q2, q3, q4) if q])
    db_exec("DELETE FROM user_practice_log WHERE question_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM user_question_error WHERE question_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM user_question_marked WHERE question_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM user_question_favorite WHERE question_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM note_question WHERE question_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM question_sub WHERE main_id IN :i", {"i": q_ids})
    db_exec("DELETE FROM question_main WHERE id IN :i", {"i": q_ids})
    db_exec("DELETE FROM user_bank_favorite WHERE bank_id IN :i", {"i": (bank_a, bank_c)})
    db_exec("DELETE FROM question_bank WHERE id IN :i", {"i": (bank_a, bank_c)})

    left_fav = db_one("SELECT COUNT(*) c FROM user_question_favorite WHERE question_id IN :i", {"i": q_ids})
    left_q = db_one("SELECT COUNT(*) c FROM question_main WHERE id IN :i", {"i": q_ids})
    left_bank = db_one("SELECT COUNT(*) c FROM question_bank WHERE id IN :i", {"i": (bank_a, bank_c)})
    T("Q42", M, "清理-无测试数据残留（不污染演示数据）",
      "清理完成", "直连库核对测试题目/题库/收藏",
      "题目=0 题库=0 收藏=0",
      all((x or {}).get("c") == 0 for x in (left_fav, left_q, left_bank)),
      "收藏残留=%s 题目残留=%s 题库残留=%s"
      % ((left_fav or {}).get("c"), (left_q or {}).get("c"), (left_bank or {}).get("c")))

    total = len(ROWS)
    passed = sum(1 for r in ROWS if r["status"] == "通过")
    print("=" * 78)
    print("用例总数 %d，通过 %d，失败 %d" % (total, passed, total - passed))
    print("=" * 78)
    for r in ROWS:
        if r["status"] == "失败":
            print("  [FAIL] %s %s" % (r["cid"], r["title"]))
            print("         预期: %s" % r["expected"])
            print("         实际: %s" % r["actual"])

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump({
            "prefix": P, "uid_admin": uid_admin, "uid1": uid1, "uid2": uid2, "rows": ROWS,
            "cleanup": {"bank_ids": [x for x in (bank_a, bank_c) if x]},
        }, f, ensure_ascii=False, indent=2)
    print("结果已写入", OUT)


if __name__ == "__main__":
    main()
