"""
RAG API 端到端测试（走真实 HTTP 链路，进程内 ASGI）

覆盖：
  1. 未登录 → 401
  2. 知识库 CRUD + 同名冲突
  3. 纯文本导入 → 后台流水线 → 轮询到 status=5（已向量化）
  4. 文件上传（multipart）→ 解析 → 处理
  5. 检索 / 问答
  6. **越权**：换个用户碰不到任何东西（这是 SpringBoot 版 `/rag/**` 的重灾区）
  7. 错误形状：`body.code` 是否按约定（404 / 400）
  8. 级联删除
  9. 清理本次全部测试数据

用法：
    ./.venv/Scripts/python.exe tests/e2e_rag_api.py
"""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from loguru import logger  # noqa: E402

logger.remove()

from httpx import ASGITransport, AsyncClient  # noqa: E402
from sqlalchemy import delete, func, select, text  # noqa: E402

from app.core.security import get_current_user  # noqa: E402
from app.database import async_session_factory, engine  # noqa: E402
from app.models.rag import (  # noqa: E402
    RagDocument,
    RagDocumentChunk,
    RagKnowledgeBase,
    RagQaRecord,
)
from app.services.rag import rag_write_service  # noqa: E402

PREFIX = "__e2e_rag__"
OTHER_USER_ID = 999999  # 只用来验越权，不落任何数据

SAMPLE_TEXT = """# 操作系统复习

## 进程与线程

进程是资源分配的基本单位，线程是 CPU 调度的基本单位。
同一进程内的线程共享地址空间，因此线程切换的开销远小于进程切换。

## 死锁

死锁的四个必要条件：互斥、占有并等待、非抢占、循环等待。
只要破坏其中任意一个条件，死锁就不会发生。
"""


class Checker:
    def __init__(self) -> None:
        self.passed = 0
        self.failed = 0

    def check(self, name: str, ok: bool, detail: str = "") -> None:
        if ok:
            self.passed += 1
            print(f"  [PASS] {name}" + (f" -> {detail}" if detail else ""))
        else:
            self.failed += 1
            print(f"  [FAIL] {name}" + (f" -> {detail}" if detail else ""))


async def _cleanup() -> int:
    """硬删本次测试数据。用精确前缀（非 LIKE —— 前缀里的 `_` 是通配符）。"""
    async with async_session_factory() as db:
        kb_ids = [
            r[0]
            for r in (
                await db.execute(
                    select(RagKnowledgeBase.id).where(
                        func.left(RagKnowledgeBase.name, len(PREFIX)) == PREFIX
                    )
                )
            ).all()
        ]
        if kb_ids:
            await db.execute(delete(RagQaRecord).where(RagQaRecord.kb_id.in_(kb_ids)))
            doc_ids = [
                r[0]
                for r in (
                    await db.execute(select(RagDocument.id).where(RagDocument.kb_id.in_(kb_ids)))
                ).all()
            ]
            if doc_ids:
                await db.execute(
                    delete(RagDocumentChunk).where(RagDocumentChunk.doc_id.in_(doc_ids))
                )
                await db.execute(delete(RagDocument).where(RagDocument.id.in_(doc_ids)))
            await db.execute(delete(RagKnowledgeBase).where(RagKnowledgeBase.id.in_(kb_ids)))
        # 越权用例会用假用户提问，可能留下 kb_id 为空的问答记录
        await db.execute(delete(RagQaRecord).where(RagQaRecord.user_id == OTHER_USER_ID))
        await db.commit()
        return len(kb_ids)


async def _wait_processed(client: AsyncClient, doc_id: int, timeout_s: float = 90) -> dict:
    """轮询文档状态直到处理结束（status 4 失败 / 5 已向量化）。"""
    last = {}
    for _ in range(int(timeout_s * 2)):
        r = await client.get(f"/rag/document/status/{doc_id}")
        last = (r.json() or {}).get("data") or {}
        if last.get("done"):
            return last
        await asyncio.sleep(0.5)
    return last


async def _call(client: AsyncClient, method: str, url: str, **kw) -> tuple[dict | None, str]:
    """发一次请求，把「应用自己抛异常」归一成一次失败的响应。

    有些缺陷（PG 拒收 NUL 这类）会让 `get_db` teardown 的 commit 再抛一次
    （Session 已需 rollback），ASGITransport 默认会把这个异常**原样抛出**，
    于是整个套件在第一个断言前就崩掉、后面的用例全不跑。
    用例真正要断言的是「这次请求成功了吗」，所以这里收成一个 `code=500` 的等价物。
    """
    try:
        r = await client.request(method, url, **kw)
    except Exception as e:  # 只在缺陷态出现
        return None, f"{type(e).__name__}: {str(e)[:90]}"
    try:
        return (r.json() or {}), ""
    except Exception:
        return {"code": r.status_code}, f"非 JSON 响应（HTTP {r.status_code}）"


async def run() -> int:
    c = Checker()
    await _cleanup()  # 开跑前先清，保证可重复执行

    from app.main import app

    async with async_session_factory() as db:
        uid = int(await db.scalar(text("SELECT user_id FROM sys_user ORDER BY user_id LIMIT 1")) or 0)
        # 需要一个**真实存在**的第二个用户：rag_knowledge_base.user_id 有指向 sys_user 的外键，
        # 用编造的 id 建库会直接被 FK 拒掉（读/删用假 id 反而没问题，因为不插入）
        other_uid = int(
            await db.scalar(
                text("SELECT user_id FROM sys_user WHERE user_id <> :u ORDER BY user_id LIMIT 1"),
                {"u": uid},
            )
            or 0
        )
    if not uid or not other_uid:
        print("用户不足（需要至少两个真实用户），无法测试")
        return 1

    def auth_as(user_id: int):
        app.dependency_overrides[get_current_user] = lambda: {
            "user_id": user_id,
            "userName": f"e2e-rag-{user_id}",
        }

    def auth_off():
        app.dependency_overrides.pop(get_current_user, None)

    client = AsyncClient(transport=ASGITransport(app=app), base_url="http://e2e", timeout=120)

    exit_code = 0
    try:
        # ---------- 1. 未登录 ----------
        print("\n=== 1. 未登录 ===")
        auth_off()
        r = await client.get("/rag/knowledgeBase/list")
        # ⚠️ 本仓库有 http_exception_handler 把 HTTPException 转成
        #    **HTTP 200 + body.code**（见 app/main.py 的 add_exception_handler）。
        #    所以这里断言的是 body.code，不是 HTTP 状态码 —— 按标准 HTTP 语义写会误判。
        c.check(
            "未登录访问知识库列表 → code=401",
            (r.json() or {}).get("code") == 401,
            f"body.code={(r.json() or {}).get('code')}",
        )

        auth_as(uid)

        # ---------- 2. 知识库 CRUD ----------
        print("\n=== 2. 知识库 CRUD ===")
        r = await client.post("/rag/knowledgeBase", json={"name": f"{PREFIX}操作系统", "description": "e2e"})
        body = r.json() or {}
        kb_id = (body.get("data") or {}).get("kbId")
        c.check("创建知识库", body.get("code") == 200 and bool(kb_id), f"code={body.get('code')} kbId={kb_id}")

        r = await client.post("/rag/knowledgeBase", json={"name": f"{PREFIX}操作系统"})
        c.check(
            "同名知识库被拒（code=400）",
            (r.json() or {}).get("code") == 400,
            f"code={(r.json() or {}).get('code')} msg={(r.json() or {}).get('msg')}",
        )

        r = await client.post("/rag/knowledgeBase", json={"name": "   "})
        c.check("空名称被拒（code=400）", (r.json() or {}).get("code") == 400)

        r = await client.get("/rag/knowledgeBase/list")
        names = [x["name"] for x in ((r.json() or {}).get("data") or [])]
        c.check("列表能看到刚建的知识库", f"{PREFIX}操作系统" in names, f"{len(names)} 个")

        r = await client.put(f"/rag/knowledgeBase/{kb_id}", json={"description": "改过的描述"})
        c.check("更新知识库", (r.json() or {}).get("code") == 200)

        r = await client.get(f"/rag/knowledgeBase/{kb_id}")
        c.check(
            "详情返回更新后的值",
            ((r.json() or {}).get("data") or {}).get("description") == "改过的描述",
        )

        r = await client.get("/rag/knowledgeBase/99999999")
        c.check("取不存在的知识库 → code=404", (r.json() or {}).get("code") == 404)

        # ---------- 3. 纯文本导入 → 后台流水线 ----------
        print("\n=== 3. 纯文本导入 → 后台向量化 ===")
        r = await client.post(
            "/rag/document/uploadText",
            json={"kbId": kb_id, "content": SAMPLE_TEXT, "fileName": f"{PREFIX}os.txt"},
        )
        body = r.json() or {}
        doc_id = (body.get("data") or {}).get("docId")
        c.check("文本导入成功", body.get("code") == 200 and bool(doc_id), f"docId={doc_id}")

        st = await _wait_processed(client, doc_id)
        c.check(
            "后台流水线跑到「已向量化」(status=5)",
            st.get("status") == 5,
            f"status={st.get('status')} {st.get('statusText')} err={st.get('errorMsg')}",
        )
        c.check("分块数 > 0", (st.get("chunkCount") or 0) > 0, f"{st.get('chunkCount')} 块")

        r = await client.get(f"/rag/document/list/{kb_id}")
        docs = (r.json() or {}).get("data") or []
        c.check("文档列表能看到它", any(d["docId"] == doc_id for d in docs), f"{len(docs)} 篇")

        r = await client.get(f"/rag/document/chunks/{doc_id}")
        chunks = (r.json() or {}).get("data") or []
        c.check("分块列表可读", len(chunks) > 0, f"{len(chunks)} 块")
        c.check("分块标记为已向量化", all(x.get("vectorized") for x in chunks))

        r = await client.get(f"/rag/document/content/{doc_id}")
        c.check(
            "能取回原文",
            "死锁" in ((r.json() or {}).get("data") or {}).get("content", ""),
        )

        # ---------- 4. 文件上传（multipart） ----------
        print("\n=== 4. 文件上传（multipart） ===")
        r = await client.post(
            "/rag/document/upload",
            params={"kbId": kb_id},
            files={"file": (f"{PREFIX}file.md", SAMPLE_TEXT.encode("utf-8"), "text/markdown")},
        )
        body = r.json() or {}
        fdoc_id = (body.get("data") or {}).get("docId")
        c.check("文件上传成功", body.get("code") == 200 and bool(fdoc_id), f"docId={fdoc_id}")
        if fdoc_id:
            st2 = await _wait_processed(client, fdoc_id)
            c.check(
                "上传的文件也向量化完成",
                st2.get("status") == 5,
                f"status={st2.get('status')} err={st2.get('errorMsg')}",
            )

        r = await client.post(
            "/rag/document/upload",
            params={"kbId": kb_id},
            files={"file": ("empty.txt", b"", "text/plain")},
        )
        c.check("空文件被拒（code=400）", (r.json() or {}).get("code") == 400)

        # ⚠️ PG 专有陷阱：PG 的 text 列不接受 NUL（0x00），MySQL 会静默存下来。
        #    含 NUL 的文本若不清洗，写入会抛
        #    `CharacterNotInRepertoireError: invalid byte sequence for encoding "UTF8": 0x00`。
        #    二进制/损坏文件被当文本解析时最容易带上它。
        r = await client.post(
            "/rag/document/upload",
            params={"kbId": kb_id},
            files={
                "file": (
                    f"{PREFIX}nul.txt",
                    "正常中文内容\x00混入空字节\x00结尾".encode("utf-8"),
                    "text/plain",
                )
            },
        )
        jn = r.json() or {}
        nul_doc_id = (jn.get("data") or {}).get("docId")
        c.check(
            "含 NUL 字节的文本能正常入库（PG 不接受 0x00）",
            jn.get("code") == 200 and bool(nul_doc_id),
            f"code={jn.get('code')} msg={(jn.get('msg') or '')[:60]}",
        )
        if nul_doc_id:
            r = await client.get(f"/rag/document/content/{nul_doc_id}")
            stored = ((r.json() or {}).get("data") or {}).get("content", "")
            c.check(
                "入库后内容里已无 NUL 字符",
                "\x00" not in stored and "正常中文内容" in stored,
                repr(stored[:40]),
            )
            await client.delete(f"/rag/document/{nul_doc_id}")

        # ---- 写库前一律 sanitize：JSON 能编码 NUL，而 PG 的 text 列拒收 0x00 ----
        # 审计发现三处写库点漏清（知识库改名 / 问答记录 / 记忆正文），漏了就是 500。
        print("\n=== 4b. 各写库点的 NUL 清理 ===")
        r = await client.post("/rag/knowledgeBase", json={"name": f"{PREFIX}NUL改名"})
        nul_kb_id = ((r.json() or {}).get("data") or {}).get("kbId")
        # ⚠️ 用 _call：缺陷态下这次请求会让应用抛异常（NUL → PG 拒收 → teardown 二次抛错），
        #    直接 await client.put 会把整个套件带崩，后面所有用例都不跑。
        body, exc = await _call(
            client,
            "PUT",
            f"/rag/knowledgeBase/{nul_kb_id}",
            json={"name": f"{PREFIX}NUL改名\u0000之后"},
        )
        c.check(
            "知识库改名带 NUL → 不 500（PG 拒收 0x00）",
            body is not None and body.get("code") == 200,
            f"code={(body or {}).get('code')} msg={((body or {}).get('msg') or '')[:60]} {exc}",
        )
        r = await client.get(f"/rag/knowledgeBase/{nul_kb_id}")
        nul_name = ((r.json() or {}).get("data") or {}).get("name", "")
        c.check(
            "改名已生效且 name 里没有 NUL 残留",
            nul_name == f"{PREFIX}NUL改名之后",
            repr(nul_name),
        )

        # 空知识库 → 问答走「不调模型」的确定性兜底，但仍会往 rag_qa_record 写一条 question
        body, exc = await _call(
            client,
            "POST",
            "/rag/ask",
            json={"query": "死锁\u0000的条件", "kbIds": [nul_kb_id]},
        )
        c.check(
            "问答 query 带 NUL → 不 500",
            body is not None and body.get("code") == 200,
            f"code={(body or {}).get('code')} msg={((body or {}).get('msg') or '')[:60]} {exc}",
        )
        async with async_session_factory() as db:
            questions = (
                await db.execute(
                    select(RagQaRecord.question).where(RagQaRecord.kb_id == nul_kb_id)
                )
            ).scalars().all()
        c.check(
            "落库的 question 里没有 NUL（写库点确实清了）",
            len(questions) >= 1 and all("\x00" not in (q or "") for q in questions),
            f"{len(questions)} 条: {[ (q or '')[:16] for q in questions ]}",
        )

        # 真实二进制文档走一遍解析器（.docx 是 zip 容器，与纯文本路径完全不同）
        import docx as _docx

        buf = io.BytesIO()
        _d = _docx.Document()
        _d.add_heading("网络协议复习", level=1)
        _d.add_paragraph("TCP 是面向连接的可靠传输协议，提供流量控制与拥塞控制。")
        _d.add_paragraph("UDP 是无连接的，不保证可靠交付，但开销更小。")
        _d.save(buf)

        r = await client.post(
            "/rag/document/upload",
            params={"kbId": kb_id},
            files={
                "file": (
                    f"{PREFIX}net.docx",
                    buf.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        jd = r.json() or {}
        docx_id = (jd.get("data") or {}).get("docId")
        c.check("docx 上传成功（走真实二进制解析）", jd.get("code") == 200 and bool(docx_id), f"code={jd.get('code')}")
        if docx_id:
            std = await _wait_processed(client, docx_id)
            c.check(
                "docx 也向量化完成",
                std.get("status") == 5,
                f"status={std.get('status')} err={std.get('errorMsg')}",
            )
            r = await client.get(f"/rag/document/content/{docx_id}")
            txt = ((r.json() or {}).get("data") or {}).get("content", "")
            c.check("docx 正文被正确抽取", "TCP" in txt and "流量控制" in txt, txt[:50])

        # ---------- 5. 检索与问答 ----------
        print("\n=== 5. 检索与问答 ===")
        r = await client.post("/rag/search", json={"query": "死锁的四个必要条件", "kbIds": [kb_id], "topK": 3})
        hits = ((r.json() or {}).get("data") or {}).get("hits") or []
        c.check("检索有结果", len(hits) > 0, f"{len(hits)} 条")
        c.check(
            "检索命中「死锁」那段",
            bool(hits) and "死锁" in hits[0]["content"],
            hits[0]["source"] if hits else "",
        )
        c.check("结果带出处", bool(hits) and bool(hits[0].get("source")))

        r = await client.post("/rag/ask", json={"query": "死锁的四个必要条件是什么", "kbIds": [kb_id]})
        ans = (r.json() or {}).get("data") or {}
        c.check("问答返回答案", bool(ans.get("answer")), (ans.get("answer") or "")[:60])
        c.check("问答带 citations", len(ans.get("citations") or []) > 0, f"{len(ans.get('citations') or [])} 条")

        r = await client.post("/rag/search", json={"query": "   ", "kbIds": [kb_id]})
        c.check("空查询被拒（code=400）", (r.json() or {}).get("code") == 400)

        # kbIds 省略 = 检索自己全部知识库
        r = await client.post("/rag/search", json={"query": "死锁"})
        c.check(
            "kbIds 省略时检索自己全部知识库",
            len(((r.json() or {}).get("data") or {}).get("hits") or []) > 0,
        )

        # ---------- 5b. 混合 kbIds（自己的 + 别人的）必须整体拒绝 ----------
        # 静默忽略越权 id 会掩盖越权尝试，所以这里要求整体 404
        auth_as(other_uid)
        r = await client.post("/rag/knowledgeBase", json={"name": f"{PREFIX}别人的库"})
        other_kb_id = ((r.json() or {}).get("data") or {}).get("kbId")
        auth_as(uid)
        if other_kb_id:
            r = await client.post(
                "/rag/search", json={"query": "死锁", "kbIds": [kb_id, other_kb_id]}
            )
            c.check(
                "kbIds 里混入别人的库 → 整体 404（不静默忽略）",
                (r.json() or {}).get("code") == 404,
                f"code={(r.json() or {}).get('code')}",
            )

        # ---------- 6. 越权 ----------
        print("\n=== 6. 越权（SpringBoot 版的重灾区）===")
        auth_as(OTHER_USER_ID)
        r = await client.get("/rag/knowledgeBase/list")
        c.check(
            "别人的知识库列表里看不到它",
            f"{PREFIX}操作系统" not in [x["name"] for x in ((r.json() or {}).get("data") or [])],
        )
        r = await client.get(f"/rag/knowledgeBase/{kb_id}")
        c.check("别人读知识库详情 → 404", (r.json() or {}).get("code") == 404, f"code={(r.json() or {}).get('code')}")
        r = await client.get(f"/rag/document/list/{kb_id}")
        c.check("别人列文档 → 404", (r.json() or {}).get("code") == 404)
        r = await client.get(f"/rag/document/content/{doc_id}")
        c.check("别人读文档原文 → 404", (r.json() or {}).get("code") == 404)
        r = await client.post("/rag/search", json={"query": "死锁", "kbIds": [kb_id]})
        c.check("别人检索该库 → 404", (r.json() or {}).get("code") == 404)
        r = await client.post("/rag/ask", json={"query": "死锁", "kbIds": [kb_id]})
        c.check("别人用该库问答 → 404", (r.json() or {}).get("code") == 404)
        r = await client.delete(f"/rag/knowledgeBase/{kb_id}")
        c.check("别人删不掉 → 404", (r.json() or {}).get("code") == 404)
        r = await client.delete(f"/rag/document/{doc_id}")
        c.check("别人删不掉文档 → 404", (r.json() or {}).get("code") == 404)
        r = await client.post(f"/rag/document/reprocess/{doc_id}")
        c.check("别人触发不了重跑 → 404", (r.json() or {}).get("code") == 404)

        # 越权尝试之后，数据必须还在
        auth_as(uid)
        r = await client.get(f"/rag/knowledgeBase/{kb_id}")
        c.check("越权尝试后数据仍在（没被误删）", (r.json() or {}).get("code") == 200)

        # ---------- 7. 补跑与重跑 ----------
        print("\n=== 7. 补跑 / 重跑 ===")
        r = await client.post(f"/rag/document/reprocess/{doc_id}")
        c.check("重跑成功", (r.json() or {}).get("code") == 200, f"{(r.json() or {}).get('data')}")
        r = await client.post("/rag/document/pending", params={"kbId": kb_id})
        c.check("补跑接口可用", (r.json() or {}).get("code") == 200, f"{(r.json() or {}).get('data')}")

        # 锁住 legacy 的一个真实缺陷：SpringBoot 的 `/parse/pending` 只做 parse+chunk
        # 就停了，文档永远停在 status=3、**检索不到**。这里人工造一份这样的半成品，
        # 验证补跑能把它救回来。
        r = await client.post(
            "/rag/document/uploadText",
            json={"kbId": kb_id, "content": SAMPLE_TEXT, "fileName": f"{PREFIX}half.txt"},
        )
        half_id = ((r.json() or {}).get("data") or {}).get("docId")
        if half_id:
            await _wait_processed(client, half_id)
            # 退回「已解析但未向量化」的状态，并清掉分块
            async with async_session_factory() as db:
                await db.execute(
                    delete(RagDocumentChunk).where(RagDocumentChunk.doc_id == half_id)
                )
                await db.execute(
                    text("UPDATE rag_document SET status = 3, chunk_count = 0 WHERE id = :i"),
                    {"i": half_id},
                )
                await db.commit()

            r = await client.get(f"/rag/document/status/{half_id}")
            c.check(
                "已造出 status=3 的半成品文档",
                ((r.json() or {}).get("data") or {}).get("status") == 3,
            )
            r = await client.post("/rag/document/pending", params={"kbId": kb_id})
            c.check(
                "补跑把它捞进队列",
                half_id in (((r.json() or {}).get("data") or {}).get("docIds") or []),
                f"{(r.json() or {}).get('data')}",
            )
            st_half = await _wait_processed(client, half_id)
            c.check(
                "status=3 的半成品被补跑成已向量化（legacy 会永远卡住）",
                st_half.get("status") == 5 and (st_half.get("chunkCount") or 0) > 0,
                f"status={st_half.get('status')} chunks={st_half.get('chunkCount')}",
            )

        # ---------- 7b. 补跑范围必须限当前用户 ----------
        # 锁住缺陷：`process_pending_documents` 曾不带 user_id 过滤 ——
        # 省略 kbId 时会拿调用方的 embedding 额度替**全部用户**补跑，
        # 并把别人的 docId 原样回给调用方。
        print("\n=== 7b. 补跑范围只限当前用户 ===")
        async with async_session_factory() as db:
            own_kb = await rag_write_service.create_knowledge_base(db, uid, f"{PREFIX}补跑自己的")
            own_doc = await rag_write_service.create_document(
                db,
                uid,
                own_kb.id,
                file_name=f"{PREFIX}own_pending.txt",
                raw_content="死锁",
                file_type="txt",
            )
            other_kb = await rag_write_service.create_knowledge_base(
                db, other_uid, f"{PREFIX}补跑别人的"
            )
            other_doc = await rag_write_service.create_document(
                db,
                other_uid,
                other_kb.id,
                file_name=f"{PREFIX}other_pending.txt",
                raw_content="死锁",
                file_type="txt",
            )
            own_doc_id, other_doc_id = own_doc.id, other_doc.id
            own_status, other_status = own_doc.status, other_doc.status
        c.check(
            "已造出两个不同用户的 pending 文档",
            own_status == 1 and other_status == 1,
            f"A doc{own_doc_id}(status={own_status}) / B doc{other_doc_id}(status={other_status})",
        )

        auth_as(other_uid)
        r = await client.post("/rag/document/pending")  # 不带 kbId
        pids = ((r.json() or {}).get("data") or {}).get("docIds") or []
        c.check(
            "B 补跑（不带 kbId）能捞到自己那篇",
            other_doc_id in pids,
            f"docIds={pids}",
        )
        c.check(
            "B 补跑（不带 kbId）不会替 A 跑、也不回 A 的 docId",
            own_doc_id not in pids,
            f"docIds={pids}（A 的 doc{own_doc_id} 是否被泄漏={own_doc_id in pids}）",
        )
        auth_as(uid)

        # ---------- 8. 级联删除 ----------
        print("\n=== 8. 级联删除 ===")
        r = await client.delete(f"/rag/document/{fdoc_id}")
        c.check("删除单个文档", (r.json() or {}).get("code") == 200)
        async with async_session_factory() as db:
            left = (
                await db.execute(
                    select(func.count()).select_from(RagDocumentChunk).where(
                        RagDocumentChunk.doc_id == fdoc_id
                    )
                )
            ).scalar() or 0
        c.check("该文档的分块被物理清掉", left == 0, f"剩 {left} 块")

        r = await client.delete(f"/rag/knowledgeBase/{kb_id}")
        c.check("删除知识库", (r.json() or {}).get("code") == 200, f"{(r.json() or {}).get('data')}")
        async with async_session_factory() as db:
            left_chunks = (
                await db.execute(
                    select(func.count()).select_from(RagDocumentChunk).where(
                        RagDocumentChunk.kb_id == kb_id
                    )
                )
            ).scalar() or 0
        c.check("知识库的分块被物理清掉（不留可检索残骸）", left_chunks == 0, f"剩 {left_chunks} 块")

        r = await client.get("/rag/knowledgeBase/list")
        c.check(
            "已删知识库不再出现在列表里（del_flag 生效）",
            f"{PREFIX}操作系统" not in [x["name"] for x in ((r.json() or {}).get("data") or [])],
        )
        r = await client.get(f"/rag/knowledgeBase/{kb_id}")
        c.check("已删知识库取详情 → 404", (r.json() or {}).get("code") == 404)

        # ---------- 8b. 已逻辑删知识库的文档不该还能按 docId 读到 ----------
        # `delete_knowledge_base` 会连带物理删文档与分块，但那是尽力而为 ——
        # 万一清理失败，只按 docId 取的接口不能还把「已删库」的正文/分块读出来。
        print("\n=== 8b. 已删知识库的文档不可读 ===")
        async with async_session_factory() as db:
            gone_kb = await rag_write_service.create_knowledge_base(db, uid, f"{PREFIX}待逻辑删")
            gone_doc = await rag_write_service.create_document(
                db,
                uid,
                gone_kb.id,
                file_name=f"{PREFIX}gone.txt",
                raw_content="残留正文",
                file_type="txt",
            )
            db.add(
                RagDocumentChunk(
                    kb_id=gone_kb.id,
                    doc_id=gone_doc.id,
                    chunk_index=0,
                    content="残留分块",
                    content_length=4,
                    token_count=4,
                    content_tokens="残留分块",
                )
            )
            await db.commit()
            gone_kb_id, gone_doc_id = gone_kb.id, gone_doc.id
            # 只把库标成逻辑删，**故意不清文档与分块** —— 模拟连带清理失败的残骸
            await db.execute(
                text("UPDATE rag_knowledge_base SET del_flag = '2' WHERE id = :i"),
                {"i": gone_kb_id},
            )
            await db.commit()
            left_docs = (
                await db.execute(
                    select(func.count())
                    .select_from(RagDocument)
                    .where(RagDocument.kb_id == gone_kb_id)
                )
            ).scalar() or 0
        c.check("已造出「库已逻辑删、文档仍在」的残骸", left_docs == 1, f"{left_docs} 篇残留")

        r = await client.get(f"/rag/document/content/{gone_doc_id}")
        c.check(
            "已删知识库的文档正文 → 404（不再读出残留正文）",
            (r.json() or {}).get("code") == 404,
            f"code={(r.json() or {}).get('code')} content={(((r.json() or {}).get('data') or {}).get('content'))!r}",
        )
        r = await client.get(f"/rag/document/chunks/{gone_doc_id}")
        c.check(
            "已删知识库的文档分块 → 404",
            (r.json() or {}).get("code") == 404,
            f"code={(r.json() or {}).get('code')}",
        )
        r = await client.get(f"/rag/document/status/{gone_doc_id}")
        c.check(
            "已删知识库的文档状态 → 404",
            (r.json() or {}).get("code") == 404,
            f"code={(r.json() or {}).get('code')}",
        )

    except Exception as e:
        c.failed += 1
        print(f"\n  [ERROR] 测试抛异常: {type(e).__name__}: {e}")
        import traceback

        traceback.print_exc()
        exit_code = 1
    finally:
        auth_off()
        await client.aclose()

    print("\n=== 9. 清理测试数据 ===")
    try:
        n = await _cleanup()
        print(f"  [清理] 删除 {n} 个知识库及其文档/分块/问答记录")
    except Exception as e:
        print(f"  [WARN] 清理失败: {e}")
        exit_code = 1

    await engine.dispose()
    print(f"\n{'='*56}\n结果：{c.passed} 通过 / {c.failed} 失败")
    return 1 if c.failed else exit_code


if __name__ == "__main__":
    raise SystemExit(asyncio.run(run()))
