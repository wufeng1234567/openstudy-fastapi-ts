# -*- coding: utf-8 -*-
"""
文档解析链路自检（离线，不连数据库、不起服务）

覆盖本轮修的两个 Bug（见 docs/openstudy-docs/v2的OS助手重构目标.md §6）：

  Bug② 文档内嵌图片抽不出来、送不出去（free + api 两端）
    · docx 正文插图 / **表格单元格内插图** 要抽得出来
    · PDF 按页抽图
    · 改名成 .doc 的 docx 要能解析（魔数分流）
    · api 链路 `chat_engine.extract_attachments`：正文里的 【图N】 与图片列表下标一一对应
    · Agent 链路 `_collect_attachments`：同上

  真老版 .doc（OLE2）：借 Word COM 转成 docx
    · 用本机 Word 现造一个「文字 - 图 - 文字」的 .doc，再交给 `parse_file`，
      要求文字与插图都拿得到、且图仍插在原文位置。
      —— 这一段需要 Windows + 装了 MS Word + 装了 pywin32，缺一样就自动跳过（不算失败）。

  Bug① 附件预览
    · 前端 `ChatAttachmentList` 依赖 `blocks` 才能在线预览 —— 这里校验
      `blocks` 里图片块的顺序与 `images` 一致（前端就靠它图文交错渲染）

用法：
    cd apps/fastapi-v2/openstudy-fastapi
    .venv/Scripts/python.exe tests/check_file_parser.py

退出码 0 = 全部通过；1 = 有断言失败。fixtures 会自动生成到 tests/_fixtures/。
"""

import base64
import gc
import io
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.file_parser import IMAGE_MARKER_TMPL, parse_file  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_fixtures")

_failures: list = []


def check(condition: bool, label: str) -> None:
    print(("  PASS  " if condition else "  FAIL  ") + label)
    if not condition:
        _failures.append(label)


# ------------------------------------------------------------
# fixtures（含图 docx / 含图 pdf / 改名成 .doc 的 docx）
# ------------------------------------------------------------
def _make_image(text: str, size=(320, 200), color: str = "red"):
    from PIL import Image, ImageDraw

    im = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(im)
    draw.rectangle([5, 5, size[0] - 5, size[1] - 5], outline=color, width=4)
    draw.text((20, size[1] // 2), text, fill="blue")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf


def build_fixtures() -> dict:
    """生成三份样例，返回 {逻辑名: (路径, 文件名)}"""
    from docx import Document
    from docx.shared import Inches

    os.makedirs(FIXTURES, exist_ok=True)

    # ---- docx：正文插图 ×2 + 表格单元格内插图 ×1 ----
    doc = Document()
    doc.add_paragraph("第一段：下面是正文插图。")
    doc.add_picture(_make_image("BODY IMAGE"), width=Inches(2.0))
    doc.add_paragraph("第二段：中间有文字。")
    doc.add_picture(_make_image("BODY IMAGE 2", color="green"), width=Inches(2.0))
    doc.add_paragraph("第三段：下面是表格，表格里也有一张图。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表格左格文字"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(
        _make_image("TABLE IMAGE", color="purple"), width=Inches(1.5)
    )
    doc.add_paragraph("第四段：结束。")
    docx_path = os.path.join(FIXTURES, "with_images.docx")
    doc.save(docx_path)

    # ---- pdf：两页，各嵌一张大图 ----
    from PIL import Image, ImageDraw

    pages = []
    for idx in range(2):
        im = Image.new("RGB", (900, 1200), "white")
        draw = ImageDraw.Draw(im)
        draw.rectangle([60, 60, 840, 500], outline="blue", width=6)
        draw.text((300, 250), f"PDF FIGURE {idx + 1}", fill="red")
        pages.append(im)
    pdf_path = os.path.join(FIXTURES, "with_images.pdf")
    pages[0].save(pdf_path, save_all=True, append_images=pages[1:])

    # ---- 把 docx 改名成 .doc，验证魔数分流 ----
    renamed_path = os.path.join(FIXTURES, "renamed_docx.doc")
    with open(docx_path, "rb") as src, open(renamed_path, "wb") as dst:
        dst.write(src.read())

    return {
        "docx": (docx_path, "with_images.docx"),
        "pdf": (pdf_path, "with_images.pdf"),
        "renamed_doc": (renamed_path, "renamed_docx.doc"),
    }


def as_data_url(path: str, filename: str) -> dict:
    mime = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pdf": "application/pdf",
        ".doc": "application/msword",
    }[os.path.splitext(filename)[1].lower()]
    with open(path, "rb") as f:
        blob = f.read()
    return {
        "name": filename,
        "type": mime,
        "base64": "data:{};base64,{}".format(mime, base64.b64encode(blob).decode()),
    }


def markers_of(text: str) -> list:
    return [int(m) for m in re.findall(r"【图(\d+)】", text)]


# ------------------------------------------------------------
# 各段自检
# ------------------------------------------------------------
def check_parse_file(fixtures: dict) -> None:
    print("\n[1] parse_file —— 解析与图片抽取")

    for key, expect_images in (("docx", 3), ("pdf", 2), ("renamed_doc", 3)):
        path, name = fixtures[key]
        with open(path, "rb") as f:
            parsed = parse_file(name, f.read())
        check(len(parsed["images"]) == expect_images,
              "{}：抽出 {} 张插图（期望 {}）".format(name, len(parsed["images"]), expect_images))
        check(not parsed.get("error"), "{}：无 error（实际 {}）".format(name, parsed.get("error")))
        check(parsed["totalTextLength"] == len(parsed["textContent"]),
              "{}：totalTextLength 与 textContent 长度一致".format(name))
        check(len(parsed["blocks"]) == len(parsed["images"]) + len(
            [b for b in parsed["blocks"] if b["type"] == "text"]),
              "{}：blocks 数量自洽".format(name))

    # 图必须落在「它原本所在的文字位置」之间，而不是全堆到开头或末尾
    path, name = fixtures["docx"]
    with open(path, "rb") as f:
        parsed = parse_file(name, f.read())
    kinds = [b["type"] for b in parsed["blocks"]]
    check(kinds[:2] == ["text", "image"],
          "docx：正文图文交替，首块是文字（实际 {}）".format(kinds[:4]))
    check("text" in kinds[kinds.index("image"):],
          "docx：图片之后还有文字块（说明图确实插在中间）")
    table_text_idx = next(
        (i for i, b in enumerate(parsed["blocks"])
         if b["type"] == "text" and "表格左格文字" in (b.get("content") or "")),
        -1,
    )
    check(table_text_idx >= 0, "docx：表格文本已输出")
    check(table_text_idx >= 0 and table_text_idx + 1 < len(parsed["blocks"])
          and parsed["blocks"][table_text_idx + 1]["type"] == "image",
          "docx：表格单元格内的插图紧跟在表格文本之后（位置正确）")
    # 用户 2026-10-11：表格必须输出成**合法 Markdown 表格**（含表头分隔行），
    # 否则前端 `marked` 认不出，会显示成一串 `a | b | c` 竖线文本。
    if table_text_idx >= 0:
        table_text = parsed["blocks"][table_text_idx]["content"] or ""
        check("| --- |" in table_text or "|:--" in table_text or "|:--" in table_text,
              "docx：表格输出为合法 Markdown 表格（含表头分隔行）")
        check(table_text.strip().startswith("|") and table_text.strip().endswith("|"),
              "docx：表格每行以 | 开头结尾（Markdown 表格形态）")

    # 真老版 .doc 之外的 OLE2 容器：认出是 .xls / .ppt 就直说，不要硬当 .doc 去转
    # （OLE2 目录项里的流名是 UTF-16LE 编码的，用 ASCII 拼是拼不出来的）
    ole2_head = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    xls_like = parse_file(
        "data.xls", ole2_head + "Workbook".encode("utf-16-le") + b"\x00" * 512
    )
    check(bool(xls_like.get("error")) and "xls" in xls_like.get("error", ""),
          "OLE2 + Workbook 流 → 认出是 .xls 并报错（不会误当 .doc 去转）")

    # 认不出流名时按 .doc 试一试（老 .doc 的流名理论上一定有，认不出多半是第三方造的文件）
    from app.services.doc_converter import ole2_flavor

    check(ole2_flavor(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512) == "",
          "流名认不出来时返回空串（交给调用方按 .doc 试）")


def check_api_chain(fixtures: dict) -> None:
    print("\n[2] api 聊天链路 —— chat_engine.extract_attachments")

    from app.services.ai import chat_engine

    files = [as_data_url(*fixtures["docx"]), as_data_url(*fixtures["pdf"])]
    text, images, warnings = chat_engine.extract_attachments(files, image_offset=2)

    check(len(images) == 5, "抽出 5 张文档插图（docx 3 + pdf 2，实际 {}）".format(len(images)))
    check(not warnings, "无告警（实际 {}）".format(warnings))
    check(markers_of(text) == list(range(2, 2 + len(images))),
          "正文 【图N】 编号连续且从 offset=2 起，与图片列表一一对应（实际 {}）".format(markers_of(text)))
    check("【附件：with_images.docx】" in text and "【附件：with_images.pdf】" in text,
          "正文按附件分段")


def check_agent_chain(fixtures: dict) -> None:
    print("\n[3] Agent 链路 —— loop._collect_attachments")

    from app.services.ai.agent.loop import AgentRequest, _collect_attachments

    req = AgentRequest(
        message="x",
        images=[],
        files=[as_data_url(*fixtures["docx"]), as_data_url(*fixtures["renamed_doc"])],
        user_id=1,
    )
    text, images, warnings, structured = _collect_attachments(req)

    check(len(images) == 6, "抽出 6 张插图（两份 docx 各 3，实际 {}）".format(len(images)))
    check(markers_of(text) == list(range(0, len(images))),
          "标记从 0 开始、与图片列表一一对应（实际 {}）".format(markers_of(text)))
    check(IMAGE_MARKER_TMPL == "【图{}】",
          "标记模板未被改写（改掉会导致 question_manage 的正则替换失效）")


def build_word_doc(png_path: str, doc_path: str) -> bool:
    """用本机 Word 现造一个「文字 - 图 - 文字」的老版 .doc（成功返回 True）

    只用于自检 —— 手工造一个真 OLE2 .doc 没有别的办法（Word 的专有二进制格式）。

    收尾复用 `doc_converter` 的清理逻辑：本机 Word 上 `Quit()` 不会真的结束进程，
    不显式清理的话跑一次自检就留一个隐藏的 WINWORD.EXE。
    """
    import pythoncom
    import win32com.client

    from app.services.doc_converter import _await_winword_exit_then_kill, _winword_pids

    pids_before = _winword_pids()
    pythoncom.CoInitialize()
    word = None
    doc = None
    owned_pids: set = set()
    try:
        word = win32com.client.DispatchEx("Word.Application")
        owned_pids = _winword_pids() - pids_before
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        doc = word.Documents.Add()
        doc.Content.Text = "第一段：图片之前的文字。"
        rng = doc.Content
        rng.Collapse(0)  # wdCollapseEnd
        # ⚠️ 必须把 Range 传进去，否则图片会插到文档开头
        doc.InlineShapes.AddPicture(png_path, False, True, rng)
        doc.Content.InsertAfter("\r第三段：图片之后的文字。")
        doc.SaveAs2(doc_path, FileFormat=0)  # wdFormatDocument → .doc
        doc.Close(False)
        doc = None
        return True
    except Exception as e:  # noqa: BLE001
        print("         （Word 造样例失败：{}）".format(e))
        return False
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:  # noqa: BLE001
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:  # noqa: BLE001
                pass
        del doc
        del word
        gc.collect()
        _await_winword_exit_then_kill(owned_pids)
        pythoncom.CoUninitialize()


def check_word_com(fixtures: dict) -> None:
    print("\n[4] 真老版 .doc（OLE2）—— 借 Word COM 转成 docx")

    from app.services.doc_converter import _winword_pids, is_available

    if not is_available():
        print("  SKIP  当前环境没有 Word COM（需要 Windows + MS Word + pywin32）")
        return

    pids_before = _winword_pids()

    png_path = os.path.join(FIXTURES, "word_fig.png")
    with open(png_path, "wb") as f:
        f.write(_make_image("WORD DOC FIGURE").getvalue())
    doc_path = os.path.join(FIXTURES, "word_real.doc")

    if not build_word_doc(png_path, doc_path):
        print("  SKIP  Word 造样例失败，跳过这一段")
        return

    with open(doc_path, "rb") as f:
        raw = f.read()
    check(raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
          "样例确实是 OLE2 容器（真 .doc，不是改名来的）")

    parsed = parse_file("word_real.doc", raw)
    kinds = [b["type"] for b in parsed["blocks"]]
    check(not parsed.get("error"),
          "真 .doc 已能解析（error={}）".format(parsed.get("error")))
    check(len(parsed["images"]) == 1, "抽出 1 张插图（实际 {}）".format(len(parsed["images"])))
    check(kinds == ["text", "image", "text"],
          "图文顺序仍是 文字→图→文字（实际 {}）".format(kinds))
    check("图片之前" in parsed["textContent"] and "图片之后" in parsed["textContent"],
          "图前后的两段文字都在")

    # 同一份文件再解析一次应命中缓存（不重复启动 Word）
    started = time.time()
    parse_file("word_real.doc", raw)
    elapsed = time.time() - started
    check(elapsed < 0.5, "同内容二次解析命中缓存，耗时 {:.3f}s".format(elapsed))

    # 转换完不能留下隐藏的 Word 进程（Quit 在这台机器的 Word 上是靠不住的）
    time.sleep(0.5)
    leaked = _winword_pids() - pids_before
    check(not leaked, "没有残留 WINWORD.EXE 进程（实际 {}）".format(sorted(leaked)))


def main() -> int:
    print("生成 fixtures →", FIXTURES)
    fixtures = build_fixtures()

    check_parse_file(fixtures)
    check_api_chain(fixtures)
    check_agent_chain(fixtures)
    check_word_com(fixtures)

    print()
    if _failures:
        print("✗ {} 项失败：".format(len(_failures)))
        for item in _failures:
            print("   -", item)
        return 1
    print("✓ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
